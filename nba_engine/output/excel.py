"""The day as a betting workbook: buys, every selection, the games' context, the record.

Sheets, in the order they get read: **Buys** (rows that cleared every gate,
tip-off order, strongest first within a game), **Selections** (every row of
record, refused rows included with their gates), **Games** (rest, travel, line
movement, the news alarm), **Injuries** (official report + ESPN), **Record**
(the graded ledger's ``scorecard.tables`` cuts with n and 95% intervals),
**Integrity** and **Legend**. Every table has a frozen, filtered header.

Written to bytes so the same object is saved and attached, and a test needs no
filesystem. Reads ledger rows; never writes one.
"""

from __future__ import annotations

from collections.abc import Sequence, Set
from dataclasses import asdict
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from nba_engine.audit.ledger import LedgerRow
from nba_engine.data.oddsapi import parse_utc
from nba_engine.output.card import (
    CONTEXT_NOTE,
    MARKET_LABEL,
    NO_MODEL_TEXT,
    PAPER_NOTE,
    SlateCard,
    board_order,
    params_version,
    row_kelly,
    strength,
    tip_key,
)
from nba_engine.output.context import ET

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="16324F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PCT = "0.0%"
MOVE_KEYS = ("ml", "spread", "total")
TIP_FORMAT = "yyyy-mm-dd h:mm AM/PM"

SELECTION_HEADER = (
    "Tip (ET)",
    "Matchup",
    "Market",
    "Side",
    "Player",
    "Line",
    "Book",
    "Price",
    "Opposite",
    "DK / BetMGM",
    "Fair prob",
    "Model prob",
    "Edge",
    "EV",
    "EV basis",
    "Kelly",
    "Tier",
    "Gates",
    "Params version",
    "Books",
    "Paired books",
    "Exec books",
    "Pass",
    "Price captured (UTC)",
    "Mode",
    "Event id",
)
_PCT_COLUMNS = {"Fair prob", "Model prob", "Edge", "EV", "Kelly"}

GAMES_HEADER = (
    "Tip (ET)",
    "Matchup",
    "Away rest",
    "Home rest",
    "Open home ML fair",
    "Now home ML fair",
    "Open home spread",
    "Now home spread",
    "Open total",
    "Now total",
    "Open board (ET)",
    "Now board (ET)",
    "News alarm",
    "Rows",
    "Buys",
    "Gates",
)

INJURY_HEADER = ("Tip (ET)", "Matchup", "Team", "Source", "Player", "Status", "Detail", "As of")

RECORD_HEADER = (
    "Cut",
    "Split",
    "Rows",
    "n",
    "Wins",
    "Losses",
    "Pushes",
    "Voids",
    "Win %",
    "Break-even %",
    "Base rate",
    "PPV",
    "NPV",
    "Units",
    "ROI",
    "ROI 95% lo",
    "ROI 95% hi",
    "CLV n",
    "Mean CLV",
    "CLV beat %",
    "CLV EV",
    "CLV EV 95% lo",
    "CLV EV 95% hi",
    "Pulled",
    "Moved",
    "Brier model",
    "Brier market",
)
_RECORD_PCT = {
    "Win %",
    "Break-even %",
    "Base rate",
    "PPV",
    "NPV",
    "ROI",
    "ROI 95% lo",
    "ROI 95% hi",
    "Mean CLV",
    "CLV beat %",
    "CLV EV",
    "CLV EV 95% lo",
    "CLV EV 95% hi",
}


def _tip(stamp: str) -> datetime | None:
    moment = parse_utc(stamp) if stamp else None
    return None if moment is None else moment.astimezone(ET).replace(tzinfo=None)


def _sheet(
    wb: Workbook,
    title: str,
    header: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    pct: Set[str] = frozenset(),
    empty: str = "",
) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(list(row))
    if not rows and empty:
        ws.append([empty])
    for i, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(9, min(len(name) + 4, 26))
        for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            if isinstance(cell.value, datetime):
                cell.number_format = TIP_FORMAT
            elif name in pct and isinstance(cell.value, (int, float)):
                cell.number_format = PCT
    if "Tip (ET)" in header:
        ws.column_dimensions["A"].width = 20
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(ws.max_row, 2)}"
    return ws


def selection_row(r: LedgerRow) -> list[Any]:
    return [
        _tip(r.tip_utc),
        r.matchup,
        MARKET_LABEL.get(r.market, r.market),
        r.side,
        r.entity,
        r.line,
        r.book,
        r.american,
        r.opposite_american,
        r.exec_prices,
        r.fair,
        r.model_prob,
        r.edge,
        r.ev,
        "" if r.ev is None else ("model" if r.model_prob is not None else "fair"),
        row_kelly(r),
        r.tier,
        r.gates.replace(";", ", "),
        params_version(r),
        r.books,
        r.paired_books,
        r.exec_books,
        r.pass_tag,
        r.priced_at,
        r.mode,
        r.event_id,
    ]


def _games_rows(card: SlateCard) -> list[list[Any]]:
    out = []
    for g in card.games:
        ctx = g.context
        mv = {m.key: m for m in ctx.moves}
        rest = {r.team: r.text().split(": ", 1)[-1] for r in ctx.rest}
        pair = {k: [mv[k].open, mv[k].now] if k in mv else [None, None] for k in MOVE_KEYS}

        out.append(
            [
                _tip(g.tip_utc),
                g.matchup,
                rest.get(g.away, ""),
                rest.get(g.home, ""),
                *pair["ml"],
                *pair["spread"],
                *pair["total"],
                _tip(ctx.opened_at),
                _tip(ctx.now_at),
                ctx.alarm,
                len(g.rows),
                len(g.buys()),
                ", ".join(f"{k} x{n}" for k, n in sorted(g.gates().items())),
            ]
        )
    return out


def _injury_rows(card: SlateCard) -> list[list[Any]]:
    return [
        [_tip(g.tip_utc), g.matchup, i.team, i.source, i.player, i.status, i.detail, i.as_of]
        for g in card.games
        for i in g.context.injuries
    ]


def _record_rows(card: SlateCard) -> list[list[Any]]:
    out = []
    for cut, table in card.record.items():
        for m in table:
            if not m.rows and cut != "buys":
                continue
            out.append(
                [
                    cut,
                    m.label,
                    m.rows,
                    m.n,
                    m.wins,
                    m.losses,
                    m.pushes,
                    m.voids,
                    m.win_pct,
                    m.required_win_pct,
                    m.base_rate,
                    m.ppv,
                    m.npv,
                    round(m.units, 4),
                    m.roi,
                    m.roi_lo,
                    m.roi_hi,
                    m.clv_n,
                    m.mean_clv,
                    m.clv_beat_pct,
                    m.mean_clv_ev,
                    m.clv_ev_lo,
                    m.clv_ev_hi,
                    m.pulled,
                    m.moved,
                    m.brier_model,
                    m.brier_market,
                ]
            )
    return out


def _legend(wb: Workbook, card: SlateCard) -> None:
    ws = wb.create_sheet("Legend")
    lines = [
        ("Card", card.title()),
        ("Mode", PAPER_NOTE),
        ("Passes of record", ", ".join(card.passes)),
        ("Params versions", "; ".join(card.versions())),
        ("Fair prob", "Consensus no-vig probability across every pairing book, Over bias removed."),
        ("Model prob", f"Engine probability (market-anchored); blank = {NO_MODEL_TEXT}."),
        ("EV", "Expected profit per unit at the price; EV basis says model or fair."),
        ("Kelly", "Full-Kelly stake fraction on the model probability, floored at 0."),
        ("Gates", "Reasons a row was refused; empty on a buy."),
        ("Book", "DraftKings or BetMGM only: the book a buy is priced at."),
        ("Context", CONTEXT_NOTE),
    ]
    for key, value in lines:
        ws.append([key, value])
        ws.cell(ws.max_row, 1).font = BOLD
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110


def build_workbook(card: SlateCard) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    buys = sorted(card.buys(), key=lambda r: (*tip_key(r), strength(r)))
    empty_buys = "No buys today: every row was refused (see Selections, Gates)."
    _sheet(
        wb,
        "Buys",
        SELECTION_HEADER,
        [selection_row(r) for r in buys],
        pct=_PCT_COLUMNS,
        empty=empty_buys,
    )
    _sheet(
        wb,
        "Selections",
        SELECTION_HEADER,
        [selection_row(r) for r in board_order(card.rows)],
        pct=_PCT_COLUMNS,
    )
    _sheet(
        wb, "Games", GAMES_HEADER, _games_rows(card), pct={"Open home ML fair", "Now home ML fair"}
    )
    _sheet(
        wb,
        "Injuries",
        INJURY_HEADER,
        _injury_rows(card),
        empty="No injury report or ESPN feed archived for this date.",
    )
    _sheet(
        wb,
        "Record",
        RECORD_HEADER,
        _record_rows(card),
        pct=_RECORD_PCT,
        empty="No graded ledger rows yet.",
    )
    integrity = wb.create_sheet("Integrity")
    integrity.append(["Check", "Value"])
    for cell in integrity[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    if card.integrity is None:
        integrity.append(["graded rows", 0])
    else:
        for key, value in asdict(card.integrity).items():
            integrity.append([key, str(value) if isinstance(value, dict) else value])
    integrity.column_dimensions["A"].width = 22
    integrity.column_dimensions["B"].width = 60
    integrity.freeze_panes = "A2"
    integrity.auto_filter.ref = f"A1:B{integrity.max_row}"
    _legend(wb, card)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["GAMES_HEADER", "INJURY_HEADER", "RECORD_HEADER", "SELECTION_HEADER", "build_workbook"]
