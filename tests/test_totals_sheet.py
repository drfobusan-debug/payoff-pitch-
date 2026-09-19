"""The hand totals sheet: + leans over, - leans under, and the row is the sum."""

from __future__ import annotations

from datetime import date as Date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from mlb_engine.data.parks import PARKS
from mlb_engine.data.vsin import Split, TotalSplit
from mlb_engine.filters.weather import WeatherConditions
from mlb_engine.output import totals_sheet
from mlb_engine.output.totals_audit import BANDS, sheet_bands
from mlb_engine.output.totals_sheet import (
    SheetRow,
    TeamSide,
    barrel_pts,
    book_pts,
    csw_pts,
    day_outs_under,
    kbb_pts,
    middle_relief,
    outs_under,
    sheet_is_current,
    sheet_legend_keys,
    short_start,
    siera_pts,
    weather_pts,
    woba_pts,
    wrc_pts,
    write_workbook,
    xera_pts,
)
from mlb_engine.recommendations import Recommendation, save_json
from mlb_engine.schemas import Pitcher


def test_a_league_average_arm_and_lineup_score_zero() -> None:
    """2026 medians: team wRC+ 99 / wOBA .315 / Barrel% 7.6; SP SIERA 4.26,
    xERA 4.31, CSW% 26.4, K-BB% 12.4; pens SIERA 3.85, CSW% 27.5, K-BB% 12.8."""
    assert wrc_pts(99) == woba_pts(0.315) == barrel_pts(7.6) == 0
    assert siera_pts(4.26) == xera_pts(4.31) == csw_pts(26.4) == kbb_pts(12.4) == 0
    assert siera_pts(3.85) == csw_pts(27.5) == kbb_pts(12.8) == 0


def test_the_bands_are_symmetric_around_the_centre() -> None:
    assert wrc_pts(110) == 1 and wrc_pts(90) == -1
    assert wrc_pts(135) == 3 and wrc_pts(65) == -3
    assert siera_pts(4.6) == 1 and siera_pts(3.5) == -1
    assert siera_pts(5.5) == 3 and siera_pts(2.5) == -3
    assert kbb_pts(8) == 1 and kbb_pts(17) == -1
    assert csw_pts(24) == 1 and csw_pts(30) == -1
    assert xera_pts(4.8) == 1 and xera_pts(3.4) == -1


def test_sharp_over_handle_is_positive_and_sharp_under_is_negative() -> None:
    over = TotalSplit(8.5, over=Split(76, 56), under=Split(24, 44))
    assert book_pts(over) == 2
    under = TotalSplit(8.5, over=Split(30, 45), under=Split(70, 55))
    assert book_pts(under) == -1
    flat = TotalSplit(8.5, over=Split(52, 50), under=Split(48, 50))
    assert book_pts(flat) == 0
    one_ticket = TotalSplit(8.5, over=Split(100, 100), under=Split(0, 0))
    assert book_pts(one_ticket) == 0
    assert book_pts(None) == 0


def test_a_roof_replaces_the_weather_score_with_minus_one() -> None:
    roofed = next(p for p in PARKS.values() if p.roof == "retractable")
    hot = WeatherConditions(95, 80, 20, 180, 20)
    pts, label = weather_pts(roofed, hot)
    assert pts == -1 and "roof" in label


def test_wind_is_signed_by_direction_and_a_cross_wind_is_nothing() -> None:
    open_park = next(p for p in PARKS.values() if p.roof == "open")
    blowing_in = WeatherConditions(70, 50, 12, 0, -12)
    blowing_out = WeatherConditions(70, 50, 12, 180, 12)
    cross = WeatherConditions(70, 50, 12, 90, 0)
    assert weather_pts(open_park, blowing_in)[0] == -2
    assert weather_pts(open_park, blowing_out)[0] == 2
    assert weather_pts(open_park, cross)[0] == 0
    assert weather_pts(open_park, WeatherConditions(75, 50, 3, 90, 0))[0] == 0


def test_the_workbook_row_sum_is_the_sum_of_its_signed_columns(tmp_path: Path) -> None:
    side = TeamSide(
        "AZ",
        "Zac Gallen (R)",
        off=3,
        sp=5,
        rp=1,
        kbb_sp=3,
        kbb_rp=1,
        bsr_pg=4.4,
        fatigue=1,
        fatigue_detail="pen",
        pen_detail="full pen",
    )
    row = SheetRow(
        "AZ @ KC",
        None,
        "8.5",
        side,
        side,
        circa=1,
        dk=2,
        weather=-1,
        weather_detail="wx",
        park=0,
        park_detail="Kauffman",
        bsr=0,
        ump=-1,
        ump_name="Ump",
        ump_status="posted",
        ump_detail="",
        engine_total=9.0,
        engine_p_over=0.5678,
    )
    path = write_workbook([row], Date(2026, 9, 9), tmp_path / "t.xlsx")
    ws = load_workbook(path)["Totals 2026-09-09"]
    header = [c.value for c in ws[1]]
    values = [c.value for c in ws[2]]
    engine_cols = {"Engine", "Eng vs line", "Eng O%"}
    pts = [
        v for h, v in zip(header, values, strict=True)
        if isinstance(v, int) and h != "SUM" and h not in engine_cols
    ]
    assert row.total_pts == 2 * (3 + 5 + 1 + 3 + 1 + 1) + 1 + 2 - 1 - 1
    assert values[header.index("SUM")] == sum(pts) == row.total_pts
    # the engine's read sits beside SUM and is never added to it
    assert header.index("Engine") == header.index("SUM") + 1
    assert values[header.index("Engine")] == 9.0 and values[header.index("Eng vs line")] == 0.5
    assert values[header.index("Eng O%")] == 0.568 and row.engine_delta == 0.5
    assert "Legend" in load_workbook(path).sheetnames
    assert sheet_bands(path) == BANDS and sheet_is_current(path)


def _rel(name: str, team: str, ip: float, siera: float, sv: int = 0, hld: int = 0, pid: int = 0) -> dict:
    return {
        "Name": f'<a href="p">{name}</a>', "Team": f'<a href="x">{team}</a>', "IP": ip, "SIERA": siera,
        "xERA": siera + 0.1, "xFIP": siera + 0.3, "C+SwStr%": 0.27, "K-BB%": 0.13,
        "SV": sv, "HLD": hld, "xMLBAMID": pid or hash(name) % 10_000,
    }


def test_middle_relief_leaves_out_the_closer_the_setup_arms_and_the_cups_of_coffee() -> None:
    rows = [
        _rel("Closer", "ATH", 60, 2.5, sv=30, pid=1),
        _rel("Setup1", "ATH", 55, 3.0, hld=25, pid=2),
        _rel("Setup2", "ATH", 50, 3.2, hld=20, pid=3),
        _rel("Bulk1", "ATH", 40, 4.8, pid=4),
        _rel("Bulk2", "ATH", 20, 5.4, pid=5),
        _rel("Cup", "ATH", 4, 1.0, pid=6),
        _rel("Traded", "- - -", 30, 6.0, pid=7),
        _rel("Ace", "PHI", 30, 3.3, pid=8),
    ]
    leverage, mid = middle_relief(rows)
    assert leverage["ATH"] == [1, 2, 3]
    ath = mid["ATH"]
    assert ath.arms == ["Bulk1", "Bulk2"] and ath.ip == 60
    assert ath.siera == pytest.approx((4.8 * 40 + 5.4 * 20) / 60)
    assert ath.row["xERA"] == pytest.approx(ath.siera + 0.1) and "SIERA 5.00" in ath.detail
    assert "- - -" not in mid and "PHI" in leverage
    # PHI's only arm is a leverage arm, so it has no middle relief read
    assert "PHI" not in mid


def _outs(pid: int, side: str, line: float, fair: float, market: str = "pitcher_outs") -> Recommendation:
    return Recommendation(
        Date(2026, 9, 15), 1, "SF @ STL", "pitcher", market, f"P{pid} Outs {side[0]}{line}", 0.5,
        line=line, fair_prob=fair, player_id=pid, side=side,
    )


def test_the_outs_prop_is_read_as_the_chance_the_starter_leaves_before_the_sixth() -> None:
    recs = [
        _outs(1, "under", 17.5, 0.42),
        _outs(1, "over", 15.5, 0.80),
        _outs(2, "under", 15.5, 0.55),
        _outs(2, "under", 18.5, 0.70),  # a rung at 6+ innings says nothing about a short start
        _outs(2, "under", 17.5, 0.90, market="pitcher_strikeouts"),
    ]
    outs = outs_under(recs)
    assert outs == {1: pytest.approx(0.42), 2: pytest.approx(0.55)}


def _card(pid: int, fair: float, lead: float | None) -> list[Recommendation]:
    r = _outs(pid, "under", 17.5, fair)
    r.hours_to_first_pitch = lead
    return [r]


def test_the_outs_prop_comes_off_the_card_the_audit_grades(tmp_path: Path) -> None:
    cfg = SimpleNamespace(audit_dir=tmp_path)
    day = Date(2026, 9, 14)
    local, pregame = tmp_path / "predictions_2026-09-14.json", tmp_path / "predictions_2026-09-14.pregame.json"
    assert day_outs_under(cfg, day) == {}

    # the pregame copy is the record when the local card is older (08:00 vs 11:00)
    save_json(_card(1, 0.40, lead=5.0), local)
    save_json(_card(1, 0.60, lead=2.0), pregame)
    assert day_outs_under(cfg, day) == {1: pytest.approx(0.60)}
    # ...or was re-priced after first pitch
    save_json(_card(1, 0.40, lead=-1.0), local)
    assert day_outs_under(cfg, day) == {1: pytest.approx(0.60)}
    # a later, still-pregame local card outranks it
    save_json(_card(1, 0.40, lead=1.0), local)
    assert day_outs_under(cfg, day) == {1: pytest.approx(0.40)}
    # a winner with no outs props defers to the other copy
    save_json([_outs(1, "under", 17.5, 0.40, market="pitcher_strikeouts")], local)
    assert day_outs_under(cfg, day) == {1: pytest.approx(0.60)}


def _total(matchup: str, side: str, line: float, model: float, fair: float | None) -> Recommendation:
    return Recommendation(
        Date(2026, 9, 19), 1, matchup, "game", "game_total", f"{side.title()} {line}", model,
        line=line, fair_prob=fair, side=side,
    )


def test_the_posted_total_falls_back_to_the_line_the_card_priced() -> None:
    recs = [
        _total("DET @ CWS", "over", 7.5, 0.64, 0.53),
        _total("DET @ CWS", "under", 7.5, 0.36, 0.47),
        _total("DET @ CWS", "over", 8.5, 0.50, 0.41),
        _total("DET @ CWS", "over", 8.0, 0.58, None),  # unpriced rung: no market read
        _total("SF @ LAD", "over", 8.5, 0.55, 0.50),
        _outs(1, "under", 17.5, 0.60),
    ]
    lines = totals_sheet.card_lines(recs)
    assert lines == {"DET @ CWS": 7.5, "SF @ LAD": 8.5}
    curves = totals_sheet.card_curves(recs)
    assert curves["DET @ CWS"][7.5] == (pytest.approx(0.64), pytest.approx(0.53))
    assert curves["DET @ CWS"][8.0] == (pytest.approx(0.58), None)
    assert sorted(curves["DET @ CWS"]) == [7.5, 8.0, 8.5]
    # VSIN wins when it has the game; the card fills the rest; neither leaves it blank
    circa = TotalSplit(8.0, over=Split(50, 50), under=Split(50, 50))
    splits = {("SF @ LAD", "circa"): circa}
    assert totals_sheet._total_label(splits, "SF @ LAD", lines) == "8"
    assert totals_sheet._total_label(splits, "DET @ CWS", lines) == "7.5"
    assert totals_sheet._total_label(splits, "NYY @ AZ", lines) == ""
    assert totals_sheet._total_label({}, "DET @ CWS") == ""


def test_the_card_lines_come_off_the_card_the_audit_grades(tmp_path: Path) -> None:
    cfg = SimpleNamespace(audit_dir=tmp_path)
    day = Date(2026, 9, 19)
    local, pregame = tmp_path / "predictions_2026-09-19.json", tmp_path / "predictions_2026-09-19.pregame.json"
    assert totals_sheet.day_card_lines(totals_sheet.day_cards(cfg, day)) == {}
    a, b = _total("DET @ CWS", "over", 7.5, 0.6, 0.5), _total("DET @ CWS", "over", 8.5, 0.5, 0.5)
    a.hours_to_first_pitch, b.hours_to_first_pitch = 5.0, 2.0
    save_json([a], local)
    save_json([b], pregame)
    cards = totals_sheet.day_cards(cfg, day)
    assert totals_sheet.day_card_lines(cards) == {"DET @ CWS": 8.5}
    assert list(totals_sheet.day_card_curves(cards)["DET @ CWS"]) == [8.5]
    # a card with no game totals defers to the other copy
    save_json([_outs(1, "under", 17.5, 0.6)], pregame)
    assert totals_sheet.day_card_lines(totals_sheet.day_cards(cfg, day)) == {"DET @ CWS": 7.5}


def test_a_stat_one_arm_lacks_is_averaged_over_the_arms_that_have_it() -> None:
    rows = [_rel(f"L{i}", "ATH", 40, 3.0, hld=10, pid=10 + i) for i in (1, 2, 3)]
    rows += [_rel("A", "ATH", 30, 4.0, pid=1), _rel("B", "ATH", 30, 4.0, pid=2)]
    rows[4]["K-BB%"] = None
    rows[4]["xERA"] = None
    rows[3]["xERA"] = None
    _, mid = middle_relief(rows)
    ath = mid["ATH"]
    assert ath.row["K-BB%"] == pytest.approx(0.13) and ath.row["xERA"] is None
    assert ath.stat("xERA") is None and ath.stat("K-BB%", 100) == 13.0


def test_a_short_start_is_the_market_first_then_the_season_line() -> None:
    webb = Pitcher(mlbam_id=1, name="Logan Webb")
    assert short_start(webb, {"IP": 100.0, "GS": 20.0}, 0.62).startswith("market 62%")
    assert short_start(webb, {"IP": 100.0, "GS": 20.0}, 0.45) == ""
    assert short_start(webb, {"IP": 100.0, "GS": 20.0}, None) == "5.0 IP/GS this season"
    assert short_start(webb, {"IP": 130.0, "Start-IP": 126.0, "GS": 20.0}, None) == ""
    assert short_start(webb, {"IP": 20.0, "GS": 0.0}, None) == "no starts this season"
    assert short_start(None, None, None) == "TBD starter"


def test_the_middle_relief_ranking_is_written_as_its_own_sheet(tmp_path: Path) -> None:
    rows = [_rel(f"L{i}", team, 40, 3.0, hld=10, pid=10 * i + k) for k, team in enumerate(("SFG", "PHI")) for i in (1, 2, 3)]
    rows += [_rel("A", "SFG", 30, 4.6, pid=101), _rel("B", "PHI", 30, 3.3, pid=102)]
    _, mid = middle_relief(rows)
    path = write_workbook([], Date(2026, 9, 15), tmp_path / "t.xlsx", mid)
    ws = load_workbook(path)["Middle relief"]
    ranked = [(r[1], r[2], r[10]) for r in ws.iter_rows(min_row=2, values_only=True)]
    assert ranked == [("PHI", -2, "B"), ("SFG", 2, "A")]


def test_a_sheet_on_the_current_bands_is_kept_and_an_older_one_is_rebuilt(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    day = Date(2026, 9, 10)
    out = tmp_path / f"totals_sheet_{day.isoformat()}.xlsx"
    monkeypatch.setattr(totals_sheet, "output_path", lambda cfg, d: out)
    monkeypatch.setattr(totals_sheet, "MLBStatsClient", lambda: pytest.fail("rebuilt a current sheet"))
    assert not sheet_is_current(out)
    # Rewrite the Legend without the stamp or the Centre rule: an older band set.
    write_workbook([], day, out)
    wb = load_workbook(out)
    for r in wb["Legend"].iter_rows(min_row=2):
        if r[0].value in ("Bands", "Centre"):
            r[0].value, r[1].value = None, None
    wb.save(out)
    assert not sheet_is_current(out)
    # A sheet from before a column was added is rebuilt even on the same bands.
    write_workbook([], day, out)
    wb = load_workbook(out)
    wb[f"Totals {day.isoformat()}"].delete_cols(1)
    wb.save(out)
    assert not sheet_is_current(out)
    write_workbook([], day, out)
    assert sheet_is_current(out)
    assert totals_sheet.build_totals_sheet(None, day, if_stale=True) == out  # type: ignore[arg-type]


def test_a_sheet_with_a_game_missing_its_line_is_rebuilt_on_the_next_pass(tmp_path: Path) -> None:
    day = Date(2026, 9, 16)
    out = tmp_path / f"totals_sheet_{day.isoformat()}.xlsx"
    side = TeamSide(
        "AZ", "Zac Gallen (R)", off=0, sp=0, rp=0, kbb_sp=0, kbb_rp=0, bsr_pg=4.4,
        fatigue=0, fatigue_detail="", pen_detail="",
    )
    def row(total: str) -> SheetRow:
        return SheetRow(
            "AZ @ KC", None, total, side, side, circa=0, dk=0, weather=0, weather_detail="", park=0,
            park_detail="", bsr=0, ump=0, ump_name="", ump_status="", ump_detail="",
        )

    write_workbook([row("7.0"), row("")], day, out)
    assert not sheet_is_current(out)
    write_workbook([row("7.0"), row("8.5")], day, out)
    assert sheet_is_current(out)


def test_a_watch_row_is_written_bold_italic_and_the_legend_carries_the_study(tmp_path: Path) -> None:
    def row(game: str, total: str, sp: int) -> SheetRow:
        side = TeamSide("AZ", "X (R)", off=0, sp=sp, rp=0, kbb_sp=0, kbb_rp=0, bsr_pg=None, fatigue=0, fatigue_detail="")
        return SheetRow(game, None, total, side, side, 0, 0, 0, "", 0, "", 0, 0, "", "", "")

    watched, high_line, weak = row("A @ B", "8.5", 4), row("C @ D", "9.0", 4), row("E @ F", "8.0", 1)
    assert (watched.flag, high_line.flag, weak.flag) == ("over", "", "")
    path = write_workbook([watched, high_line, weak], Date(2026, 9, 17), tmp_path / "t.xlsx")
    wb = load_workbook(path)
    ws = wb["Totals 2026-09-17"]
    header = [c.value for c in ws[1]]
    sum_col = header.index("SUM") + 1
    assert ws.cell(row=2, column=1).font.italic and ws.cell(row=2, column=sum_col).font.bold
    assert ws.cell(row=2, column=sum_col).font.italic
    assert not ws.cell(row=3, column=1).font.italic and not ws.cell(row=3, column=sum_col).font.italic
    assert not ws.cell(row=4, column=1).font.italic
    keys = sheet_legend_keys(path)
    assert "Factor study" in keys and "Over band" in keys and "Watch" in keys
    assert sheet_is_current(path)


def test_a_sheet_whose_legend_predates_the_study_is_rebuilt(tmp_path: Path) -> None:
    day = Date(2026, 9, 17)
    out = tmp_path / "t.xlsx"
    write_workbook([], day, out)
    assert sheet_is_current(out)
    wb = load_workbook(out)
    legend = wb["Legend"]
    legend.delete_rows(legend.max_row)
    wb.save(out)
    assert not sheet_is_current(out)


def test_a_game_vsin_has_not_posted_reads_its_total_off_the_card_board() -> None:
    """9/18 morning: VSIN listed the 12:35 game alone, so fourteen rows had no total
    while the slate pass had already priced every game's. VSIN's number still wins
    where it has one; the board fills the rest; a game neither holds stays blank."""
    splits = {("CHC @ CIN", "draftkings"): TotalSplit(8.5), ("CHC @ CIN", "circa"): TotalSplit(9.0)}
    board = {"CHC @ CIN": 8.5, "KC @ PIT": 8.0}
    assert totals_sheet._total_label(splits, "CHC @ CIN", board) == "9 / dk 8.5"
    assert totals_sheet._total_label(splits, "KC @ PIT", board) == "8"
    assert totals_sheet._total_label(splits, "SF @ LAD", board) == ""
    assert totals_sheet._total_label(splits, "KC @ PIT") == ""


def test_the_sheet_asks_the_public_boards_only_for_games_still_blank(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """9/19 on the Mac: no VSIN line, no local card, no board -- fourteen blanks.
    The public boards fill them; a game VSIN or the card already priced is not asked for."""
    asked: list[set[str]] = []

    def fake(cfg: object, slate: object, missing: set[str]) -> dict[str, float]:
        asked.append(set(missing))
        return {m: 8.0 for m in missing}

    monkeypatch.setattr(totals_sheet, "posted_totals", fake)
    splits = {("DET @ CWS", "draftkings"): TotalSplit(8.0)}
    board = {"CHC @ CIN": 9.5}
    games = ["DET @ CWS", "CHC @ CIN", "KC @ PIT"]
    missing = {m for m in games if not totals_sheet._total_label(splits, m, board)}
    board = {**fake(None, None, missing), **board}
    assert asked == [{"KC @ PIT"}]
    assert totals_sheet._total_label(splits, "KC @ PIT", board) == "8"
    assert totals_sheet._total_label(splits, "CHC @ CIN", board) == "9.5"


def test_a_day_with_no_local_card_is_pulled_off_the_state_branch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from mlb_engine import state

    pulled: list[tuple[Path, str, tuple[str, ...] | None]] = []

    def fake_pull(data_dir: Path, branch: str = "x", dates: tuple[str, ...] | None = None) -> None:
        pulled.append((data_dir, branch, dates))
        return None

    monkeypatch.setattr(state, "auto_pull", fake_pull)
    cfg = SimpleNamespace(
        data_dir=tmp_path, audit_dir=tmp_path / "audit", state_sync=True, state_branch="engine-state"
    )
    totals_sheet.pull_day_state(cfg, Date(2026, 9, 19))  # type: ignore[arg-type]
    assert pulled == [(tmp_path, "engine-state", ("2026-09-19",))]
    cfg.state_sync = False
    totals_sheet.pull_day_state(cfg, Date(2026, 9, 19))  # type: ignore[arg-type]
    assert len(pulled) == 1
