"""The card's star: a kind of bet with 100+ graded buys winning on all three tests."""

from __future__ import annotations

from datetime import date, timedelta

from openpyxl import load_workbook

from cfb_engine.audit.ledger import LedgerEntry
from cfb_engine.audit.proven import LEGEND, STAR, Proven
from cfb_engine.market.tiers import Tier
from cfb_engine.output.card import build_article
from cfb_engine.output.excel import write_workbook
from cfb_engine.recommendations import Recommendation

DAY = date(2026, 10, 10)


def _entry(i: int, market: str, result: str, drift: float | None = None) -> LedgerEntry:
    return LedgerEntry(
        date=(date(2026, 9, 1) + timedelta(days=i % 30)).isoformat(),
        matchup=f"A{i} @ H{i}",
        category=market,
        market=market,
        selection="H -3.5",
        line=-3.5,
        book="b",
        odds=-110.0,
        tier=Tier.STRONG.value,
        model_prob=0.55,
        ev=0.05,
        result=result,
        pnl=100 / 110 if result == "win" else -1.0,
        drift=drift,
    )


def _ledger(market: str, n: int, win_every: int, drift: float | None = None) -> list[LedgerEntry]:
    return [_entry(i, market, "loss" if i % win_every == 0 else "win", drift) for i in range(n)]


def _rec(market: str = "game_ats", tier: Tier = Tier.STRONG, drift: float | None = None):
    return Recommendation(
        game_date=DAY,
        game_id="g1",
        matchup="Away @ Home",
        market=market,
        selection="Home -3.5",
        model_prob=0.55,
        line=-3.5,
        market_american=-110,
        fair_prob=0.5,
        edge=0.05,
        ev=0.05,
        tier=tier,
        drift=drift,
        side="cover",
        home_abbrev="Home",
        away_abbrev="Away",
    )


def test_a_winning_market_needs_the_full_hundred() -> None:
    assert Proven.from_ledger(_ledger("game_ats", 120, 3)).markets == {"game_ats"}
    assert Proven.from_ledger(_ledger("game_ats", 99, 3)).markets == frozenset()


def test_a_losing_or_flat_market_gets_no_star() -> None:
    assert Proven.from_ledger(_ledger("game_total", 120, 2)).markets == frozenset()


def test_only_buys_in_the_proven_market_are_starred() -> None:
    proven = Proven.from_ledger(_ledger("game_ats", 120, 3))
    assert proven.starred(_rec())
    assert not proven.starred(_rec(market="game_total"))
    assert not proven.starred(_rec(tier=Tier.PASS))


def test_a_promoted_upgrade_stars_buys_the_line_moved_toward() -> None:
    proven = Proven.from_ledger(_ledger("game_total", 120, 3, drift=0.03))
    assert {u.name for u in proven.upgrades} >= {"drift_upgrade_agrees_2pct"}
    assert proven.starred(_rec(market="game_ml", drift=0.03))
    assert not proven.starred(_rec(market="game_ml", drift=0.0))


def test_the_card_and_workbook_mark_the_star_and_explain_it(tmp_path) -> None:
    proven = Proven.from_ledger(_ledger("game_ats", 120, 3))
    recs = [_rec(), _rec(market="game_total")]
    html, _ = build_article(DAY, recs, proven=proven)
    assert f"{STAR} Home -3.5" in html and LEGEND in html
    plain, _ = build_article(DAY, recs)
    assert STAR not in plain

    ws = load_workbook(write_workbook(recs, tmp_path / "c.xlsx", DAY, proven))["All"]
    cells = [c.value for c in ws["D"]]
    assert f"{STAR} Home -3.5" in cells and "Home -3.5" in cells
    starred = next(c for c in ws["D"] if c.value == f"{STAR} Home -3.5")
    assert starred.font.bold
    assert LEGEND in [c.value for c in ws["A"]]
