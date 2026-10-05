"""The graded-only totals rule: line moved toward the model's side before the bet."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from cfb_engine.audit.ledger import LedgerEntry
from cfb_engine.audit.probation import PROMOTE, WATCHING, probation_rows, rule_probation
from cfb_engine.market import lineagree
from cfb_engine.market.tiers import Tier
from cfb_engine.output.card import build_article
from cfb_engine.output.excel import LINE_AGREES_TAB, write_workbook
from cfb_engine.recommendations import Recommendation

DAY = date(2026, 10, 10)


def _rec(
    side: str,
    *,
    model: float,
    fair: float,
    open_drift: float | None = None,
    drift: float | None = None,
    market: str = "game_total",
) -> Recommendation:
    return Recommendation(
        game_date=DAY,
        game_id="g1",
        matchup="Away @ Home",
        market=market,
        selection=f"{side.title()} 52.5",
        model_prob=model,
        line=52.5,
        market_american=-110,
        fair_prob=fair,
        tier=Tier.PASS,
        open_drift=open_drift,
        drift=drift,
        side=side,
        home_abbrev="Home",
        away_abbrev="Away",
    )


def test_only_the_models_totals_side_with_the_line_coming_to_it_qualifies() -> None:
    assert lineagree.line_agrees("game_total", 0.56, 0.50, 0.012, None)
    assert not lineagree.line_agrees("game_total", 0.56, 0.50, 0.005, None)
    assert not lineagree.line_agrees("game_total", 0.44, 0.50, 0.03, None)
    assert not lineagree.line_agrees("game_ats", 0.56, 0.50, 0.03, None)
    assert not lineagree.line_agrees("game_total", 0.56, None, 0.03, None)
    assert not lineagree.line_agrees("game_total", 0.56, 0.50, None, None)


def test_the_opener_move_is_read_first_and_the_days_board_backs_it_up() -> None:
    assert not lineagree.line_agrees("game_total", 0.56, 0.50, -0.01, 0.03)
    assert lineagree.line_agrees("game_total", 0.56, 0.50, None, 0.03)


def test_the_recommendation_labels_its_cell_with_the_move() -> None:
    rec = _rec("over", model=0.57, fair=0.50, open_drift=0.023)
    assert rec.line_agrees
    assert rec.as_row()[lineagree.LABEL] == "YES (+2.3 pp)"
    assert _rec("under", model=0.43, fair=0.50, open_drift=0.023).as_row()[lineagree.LABEL] == ""


def _entry(result: str, day: int, *, tier: str = Tier.PASS.value, drift: float = 0.02) -> LedgerEntry:
    return LedgerEntry(
        date=f"2026-09-{day:02d}" if day <= 30 else f"2026-10-{day - 30:02d}",
        matchup=f"G{day} @ H{day}",
        category="Totals",
        market="game_total",
        selection="Over 52.5",
        line=52.5,
        book="b",
        odds=100.0,
        tier=tier,
        model_prob=0.56,
        ev=0.05,
        result=result,
        pnl=1.0 if result == "win" else -1.0,
        fair_prob=0.50,
        drift=drift,
    )


def _pattern(p: str, **kw: object) -> list[LedgerEntry]:
    return [_entry("win" if ch == "w" else "loss", i + 1, **kw) for i, ch in enumerate(p)]  # type: ignore[arg-type]


def test_the_rule_grades_every_row_it_picks_not_only_buys() -> None:
    rows = _pattern("w" * 20 + "l" * 6 + "w" * 20 + "l" * 6)
    (verdict,) = rule_probation(rows, min_n=50)
    assert verdict.name == lineagree.RULE_NAME
    assert verdict.kind == "rule"
    assert verdict.n == 52
    assert verdict.status == PROMOTE
    assert verdict.finding.startswith(f"{lineagree.RULE_NAME}: 52 games it would bet winning")


def test_the_rule_ignores_moves_away_and_waits_for_volume() -> None:
    rows = _pattern("w" * 30, drift=-0.02) + _pattern("w" * 10)
    (verdict,) = rule_probation(rows)
    assert (verdict.status, verdict.n) == (WATCHING, 10)
    assert lineagree.RULE_NAME in {p.name for p in probation_rows(rows)}


def test_the_workbook_shows_the_rule_on_its_own_labelled_tab(tmp_path: Path) -> None:
    agree = _rec("over", model=0.57, fair=0.50, open_drift=0.02)
    other = _rec("under", model=0.43, fair=0.50, open_drift=-0.02)
    wb = load_workbook(write_workbook([agree, other], tmp_path / "card.xlsx", DAY))
    ws = wb[LINE_AGREES_TAB]
    header = [c.value for c in ws[1]]
    assert lineagree.LABEL in header
    col = header.index(lineagree.LABEL)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1 and rows[0][col] == "YES (+2.0 pp)"
    all_rows = list(wb["All"].iter_rows(min_row=2, values_only=True))
    assert sorted(r[col] or "" for r in all_rows) == ["", "YES (+2.0 pp)"]


def test_the_card_lists_the_rule_apart_from_the_best_bets_with_its_record() -> None:
    agree = _rec("over", model=0.57, fair=0.50, open_drift=0.02)
    (record,) = rule_probation(_pattern("wl" * 5))
    html, _ = build_article(DAY, [agree], record)
    assert lineagree.LABEL in html
    block = html[html.index("class='lineagree'") :]
    assert "Over 52.5" in block and "market moved +2.0 pp our way" in block
    assert "Graded so far: 10 games" in block and "WATCHING" in block
    bare, _ = build_article(DAY, [_rec("over", model=0.57, fair=0.50, open_drift=0.0)])
    assert "No total has moved toward the model's side yet today." in bare
    assert "No graded record yet." in bare
