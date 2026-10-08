"""The low-total Under candidate screen: graded on the buys it would skip, never applied."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from cfb_engine.audit.ledger import LedgerEntry
from cfb_engine.audit.probation import (
    CANDIDATE_SCREENS,
    SHIP,
    WATCHING,
    candidate_probation,
    probation_rows,
)
from cfb_engine.market import lowtotal
from cfb_engine.market.tiers import Tier
from cfb_engine.output.card import build_article
from cfb_engine.output.excel import write_workbook
from cfb_engine.recommendations import Recommendation

DAY = date(2026, 10, 10)
(SCREEN,) = (c for c in CANDIDATE_SCREENS if c.name == lowtotal.RULE_NAME)


def _rec(
    selection: str, line: float, tier: Tier = Tier.STRONG, market: str = "game_total"
) -> Recommendation:
    return Recommendation(
        game_date=DAY,
        game_id="g1",
        matchup="Away @ Home",
        market=market,
        selection=selection,
        model_prob=0.55,
        line=line,
        market_american=-110,
        fair_prob=0.50,
        tier=tier,
        home_abbrev="Home",
        away_abbrev="Away",
    )


def test_only_an_under_below_fifty_is_refused() -> None:
    assert lowtotal.refuses("game_total", "Under 49.5", 49.5)
    assert not lowtotal.refuses("game_total", "Under 50", 50.0)
    assert not lowtotal.refuses("game_total", "Over 44.5", 44.5)
    assert not lowtotal.refuses("game_ats", "Under 3.5", 3.5)
    assert not lowtotal.refuses("game_total", "Under", None)


def test_the_card_flags_buys_only_and_leaves_the_tier_alone() -> None:
    buy = _rec("Under 46.5", 46.5)
    assert buy.low_total_under and buy.tier == Tier.STRONG
    assert buy.as_row()[lowtotal.LABEL] == "WOULD SKIP"
    assert not _rec("Under 46.5", 46.5, Tier.PASS).low_total_under
    assert _rec("Under 52.5", 52.5).as_row()[lowtotal.LABEL] == ""


def _entry(result: str, day: int, line: float = 47.5, tier: str = Tier.STRONG.value) -> LedgerEntry:
    return LedgerEntry(
        date=f"2026-09-{day:02d}" if day <= 30 else f"2026-10-{day - 30:02d}",
        matchup=f"G{day} @ H{day}",
        category="Totals",
        market="game_total",
        selection=f"Under {line:g}",
        line=line,
        book="b",
        odds=100.0,
        tier=tier,
        model_prob=0.55,
        ev=0.05,
        result=result,
        pnl=1.0 if result == "win" else -1.0,
        fair_prob=0.50,
    )


def test_the_screen_is_graded_on_the_low_total_under_buys_it_would_skip() -> None:
    rows = [_entry("loss" if i % 4 else "win", i + 1) for i in range(60)]
    rows += [_entry("win", i + 1, line=55.5) for i in range(20)]
    rows += [_entry("loss", i + 1, tier=Tier.PASS.value) for i in range(20)]
    (verdict,) = candidate_probation(rows, (SCREEN,), min_n=50)
    assert (verdict.n, verdict.status) == (60, SHIP)
    (thin,) = candidate_probation(rows[:10], (SCREEN,))
    assert thin.status == WATCHING
    assert lowtotal.RULE_NAME in {p.name for p in probation_rows(rows)}


def test_the_pdf_and_workbook_show_the_flag(tmp_path: Path) -> None:
    recs = [_rec("Under 46.5", 46.5), _rec("Over 46.5", 46.5, Tier.PASS)]
    html, _ = build_article(DAY, recs)
    assert lowtotal.LABEL in html and "Under 46.5" in html
    ws = load_workbook(write_workbook(recs, tmp_path / "card.xlsx", DAY))["All"]
    header = [c.value for c in ws[1]]
    col = header.index(lowtotal.LABEL)
    sel = header.index("Selection")
    flags = {r[sel]: r[col] for r in ws.iter_rows(min_row=2, values_only=True) if r[sel]}
    assert flags == {"Under 46.5": "WOULD SKIP", "Over 46.5": None}
