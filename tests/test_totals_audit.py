"""The totals sheet's receipt: rows filed at write time, graded off the finals, tallied honestly."""

from __future__ import annotations

from datetime import date as Date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from mlb_engine.output import totals_audit
from mlb_engine.output.totals_audit import (
    LedgerRow,
    grade,
    merge,
    read_ledger,
    summarize,
    summary_text,
    verdict,
    wilson,
    write_ledger,
    write_workbook,
)

D = "2026-09-09"


def _rows() -> list[LedgerRow]:
    return [
        LedgerRow(D, "AZ @ KC", 1, 8.5, 30),
        LedgerRow(D, "WSH @ SD", 2, 8.5, 25),
        LedgerRow(D, "MIN @ DET", 3, 8.5, 22),
        LedgerRow(D, "TEX @ SEA", 4, 8.0, 12),
        LedgerRow(D, "CLE @ BAL", 5, 8.5, 3),
        LedgerRow(D, "TB @ ATL", 6, 9.0, 0),
        LedgerRow(D, "HOU @ PHI", 7, 8.0, -4),
    ]


FINALS = {
    1: ("AZ @ KC", 2, 5),
    2: ("WSH @ SD", 2, 9),
    3: ("MIN @ DET", 2, 7),
    4: ("TEX @ SEA", 2, 3),
    5: ("CLE @ BAL", 5, 9),
    6: ("TB @ ATL", 4, 5),
    7: ("HOU @ PHI", 7, 11),
}


def test_grading_signs_the_result_against_the_line_and_the_sum() -> None:
    rows = _rows()
    assert grade(rows, FINALS) == 7
    by = {r.game: r for r in rows}
    assert by["AZ @ KC"].result == "under" and by["AZ @ KC"].hit is False
    assert by["WSH @ SD"].result == "over" and by["WSH @ SD"].hit is True
    assert by["TB @ ATL"].result == "push" and by["TB @ ATL"].hit is None
    assert by["HOU @ PHI"].result == "over" and by["HOU @ PHI"].hit is False
    assert grade(rows, FINALS) == 0  # never regraded


def test_a_label_match_is_refused_for_a_doubleheader_without_a_game_pk() -> None:
    twice = {11: ("NYY @ BOS", 3, 4), 12: ("NYY @ BOS", 8, 9)}
    assert grade([LedgerRow(D, "NYY @ BOS", 0, 9.0, 5)], twice) == 0
    assert grade([LedgerRow(D, "NYY @ BOS", 12, 9.0, 5)], twice) == 1
    assert grade([LedgerRow(D, "NYY @ BOS", 0, 9.0, 5)], {11: ("NYY @ BOS", 3, 4)}) == 1


def test_summary_counts_sign_rank_and_the_slates_own_over_rate() -> None:
    rows = _rows()
    grade(rows, FINALS)
    s = summarize(rows)
    # leans: 5 positive + 1 negative graded (TB is 0 -> no lean, and a push)
    assert (s.sign.hits, s.sign.misses, s.sign.pushes) == (3, 3, 0)
    assert (s.rank_over.hits, s.rank_over.misses) == (2, 1)  # KC under, SD/DET over
    assert (s.rank_under.hits, s.rank_under.misses, s.rank_under.pushes) == (0, 2, 1)  # BAL, PHI over; ATL push
    assert (s.over_rate.hits, s.over_rate.misses, s.over_rate.pushes) == (4, 2, 1)
    assert s.days == 1 and s.games == 7
    text = summary_text(Date(2026, 9, 9), rows, s)
    assert "sign 3-3 (50%)" in text and "top-3 overs 2-1 (67%)" in text


def test_magnitude_bands_tally_the_sign_call_regardless_of_direction() -> None:
    rows = _rows()
    grade(rows, FINALS)
    m = summarize(rows).by_mag
    # +3 hit and -4 miss share the 1..4 band; +30 miss, +25/+22 hits fill >= 15; 0 is no call
    assert (m["1..4"].hits, m["1..4"].misses) == (1, 1)
    assert (m["5..9"].hits, m["5..9"].misses) == (0, 0)
    assert (m["10..14"].hits, m["10..14"].misses) == (0, 1)
    assert (m[">= 15"].hits, m[">= 15"].misses) == (2, 1)
    assert sum(t.n for t in m.values()) == 6
    text = summary_text(Date(2026, 9, 9), rows, summarize(rows))
    assert "|SUM|  >= 15: 2-1 (67%) [21-94%] unproven (n<30)" in text
    assert "|SUM|   5..9: 0-0 unproven (n<30)" in text


def test_verdict_needs_the_whole_interval_past_break_even() -> None:
    lo, hi = wilson(60, 100)
    assert 0.50 < lo < 0.524 < 0.60 < hi < 0.70
    assert verdict(totals_audit.Tally(60, 40)) == "unproven"
    assert verdict(totals_audit.Tally(600, 400)) == "edge"
    assert verdict(totals_audit.Tally(400, 600)) == "no edge"
    assert verdict(totals_audit.Tally(20, 0)) == "unproven (n<30)"  # a hot streak is not a verdict
    assert verdict(totals_audit.Tally()) == "unproven (n<30)"


def test_ledger_round_trips_and_merge_never_duplicates(tmp_path: Path) -> None:
    rows = _rows()
    grade(rows, FINALS)
    path = tmp_path / "totals_ledger.csv"
    write_ledger(path, rows)
    back = read_ledger(path)
    assert back == sorted(rows, key=lambda r: (r.date, -r.sum_pts, r.game))
    again = merge(back, _rows())
    assert len(again) == 7 and all(r.graded for r in again)
    fresh = merge(back, [LedgerRow("2026-09-10", "AZ @ KC", 9, 8.0, 4)])
    assert len(fresh) == 8


def test_a_rewritten_sheet_replaces_its_ungraded_rows_but_never_graded_ones(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "totals_ledger.csv"
    graded = _rows()
    grade(graded, FINALS)
    today = "2026-09-10"
    write_ledger(path, graded + [LedgerRow(today, "AZ @ KC", 9, 8.0, 12)])
    monkeypatch.setattr(totals_audit, "ledger_path", lambda cfg: path)
    monkeypatch.setattr(
        totals_audit, "rows_from_sheet", lambda sheet, day, pks: [LedgerRow(today, "AZ @ KC", 9, 8.0, 3)]
    )
    totals_audit.record_sheet(None, Date(2026, 9, 10), tmp_path / "x.xlsx", {})  # type: ignore[arg-type]
    back = read_ledger(path)
    assert [r.sum_pts for r in back if r.date == today] == [3]
    assert len([r for r in back if r.date == D]) == 7
    monkeypatch.setattr(totals_audit, "rows_from_sheet", lambda sheet, day, pks: [LedgerRow(D, "AZ @ KC", 1, 8.5, 0)])
    totals_audit.record_sheet(None, Date(2026, 9, 9), tmp_path / "x.xlsx", {})  # type: ignore[arg-type]
    assert {r.game: r.sum_pts for r in read_ledger(path) if r.date == D}["AZ @ KC"] == 30


def test_workbook_has_yesterday_summary_and_ledger(tmp_path: Path) -> None:
    rows = _rows()
    grade(rows, FINALS)
    wb = load_workbook(write_workbook(tmp_path / "a.xlsx", Date(2026, 9, 9), rows, rows))
    assert wb.sheetnames == ["Graded 2026-09-09", "Summary", "Ledger"]
    graded = list(wb["Graded 2026-09-09"].iter_rows(values_only=True))
    assert graded[0][:4] == ("Date", "Game", "Line", "SUM")
    assert wb["Summary"]["A3"].value.startswith("Sign of SUM") and wb["Summary"]["B3"].value == "3-3"
    assert graded[1][1] == "AZ @ KC" and graded[1][-1] == "miss"
