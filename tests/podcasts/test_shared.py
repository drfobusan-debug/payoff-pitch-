"""The league-neutral podcast core: league classification, cache scope, the one
shared read, and the cross-league handicapper audit."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from engine_common.podcasts import audit, run
from engine_common.podcasts import extract as ex
from engine_common.podcasts.episodes import Episode
from engine_common.podcasts.feed import FeedItem
from engine_common.podcasts.picks import LEDGER_FIELDS, ledger_path
from engine_common.podcasts.shows import BY_KEY, CBB, CFB, MLB, NBA, NFL, SHOWS, for_league
from engine_common.podcasts.state import fill_in_file
from engine_common.podcasts.transcript import Line, write_transcript

PUB = "2026-09-10T15:00:00+00:00"
ITEM = FeedItem("g1", "NFL Week 2 and college picks", PUB, "https://x/a.mp3", "1:00:00")
MULTI = BY_KEY["wisekracks"]
LINES = [
    Line(60, "I love the Bills plus two and a half this week, my number is plus one."),
    Line(120, "And in college I'll lean Iowa plus fourteen and a half."),
]


def _raw(**kw: object) -> dict[str, object]:
    base: dict[str, object] = dict(
        league=NFL,
        kind="official",
        host=None,
        team="Bills",
        opponent="Chiefs",
        market="game_ats",
        side=None,
        line=2.5,
        price=-110,
        units=None,
        stamp="01:00",
        quote="I love the Bills plus two and a half this week",
        edge="my number is plus one",
        fair_line=1,
        reason="",
        description="Bills +2.5",
    )
    base.update(kw)
    return base


def test_every_supplied_nfl_show_is_registered_once() -> None:
    keys = [s.key for s in SHOWS]
    assert len(keys) == len(set(keys))
    nfl = {s.key for s in for_league(NFL)}
    assert {"wisekracks", "schwartz", "btp", "action_nfl", "numbers_game"} <= nfl
    assert all(s.leagues == (NFL,) for s in for_league(NFL))


def test_the_mlb_and_nba_shows_are_read_for_their_leagues() -> None:
    assert {"action_mlb", "daily_diamond"} <= {s.key for s in for_league(MLB)}
    nba = {s.key for s in for_league(NBA)}
    assert {"hardwood", "buckets", "nba_gambling", "duncd_on", "bet_the_board"} <= nba
    assert set(BY_KEY["bet_the_board"].leagues) == {NFL, CFB, NBA, CBB}
    assert BY_KEY["bet_the_board"].wants("Bet The Board: NFL Week 4 Picks", "")


def test_check_keeps_the_league_edge_and_only_spoken_numbers() -> None:
    ep = Episode(MULTI, ITEM)
    pick, why = ex.check(_raw(), ep, LINES)
    assert pick is not None, why
    assert (pick.league, pick.line, pick.fair_line, pick.edge) == (
        NFL,
        2.5,
        1.0,
        "my number is plus one",
    )
    assert pick.price is None  # -110 was never said


def test_check_refuses_leagues_the_show_is_not_read_for_and_misquotes() -> None:
    nfl_only = Episode(MULTI.for_league(NFL), ITEM)
    cfb = _raw(
        league=CFB,
        team="Iowa",
        line=14.5,
        stamp="02:00",
        quote="I'll lean Iowa plus fourteen and a half",
    )
    assert ex.check(cfb, nfl_only, LINES) == (None, "league not read from this show")
    assert (
        ex.check(_raw(league="xfl"), Episode(MULTI, ITEM), LINES)[1]
        == "league not read from this show"
    )
    assert ex.check(_raw(quote="the Raiders are a lock"), nfl_only, LINES)[1].startswith("quote")


def test_overlapping_windows_merge_and_official_beats_lean() -> None:
    ep = Episode(MULTI, ITEM)
    a, _ = ex.check(_raw(line=None), ep, LINES)
    b, _ = ex.check(_raw(), ep, LINES)
    lean, _ = ex.check(_raw(kind="lean"), ep, LINES)
    assert a and b and lean
    [kept] = ex._merge([a, b, lean])
    assert (kept.kind, kept.line) == ("official", 2.5)


def test_cache_is_reused_only_for_the_same_league_scope(tmp_path: Path) -> None:
    calls: list[str] = []

    def model(text: str, ep: Episode, **_: object) -> list[dict[str, object]]:
        calls.append(ep.show.key)
        return [_raw()]

    with patch.object(ex, "call_model", model):
        ex.extract(Episode(MULTI, ITEM), LINES, tmp_path, api_key="k")
        ex.extract(Episode(MULTI, ITEM), LINES, tmp_path, api_key="k")
        assert len(calls) == 1
        ex.extract(Episode(MULTI.for_league(NFL), ITEM), LINES, tmp_path, api_key="k")
        assert len(calls) == 2


def test_one_read_transcribes_once_and_extracts_for_every_league(tmp_path: Path) -> None:
    since = datetime(2026, 9, 1, tzinfo=timezone.utc)
    until = datetime(2026, 9, 20, tzinfo=timezone.utc)
    legacy = tmp_path / "old"
    old = legacy / "podcasts" / MULTI.key / f"{Episode(MULTI, ITEM).key}.txt"
    old.parent.mkdir(parents=True)
    write_transcript(old, LINES)
    seen: list[tuple[str, ...]] = []

    def model(text: str, ep: Episode, **_: object) -> list[dict[str, object]]:
        seen.append(ep.show.leagues)
        return [_raw()]

    store = tmp_path / "store"
    with (
        patch.object(run, "fetch_episodes", lambda s, a, b: [Episode(s, ITEM)]),
        patch.object(ex, "call_model", model),
    ):
        rep = run.read_feeds(
            store, since, until, shows=[MULTI], league=NFL, legacy=[legacy], api_key="k", log=False
        )
    assert (rep.episodes, rep.transcribed, rep.by_league[NFL]) == (1, 1, 1)
    assert seen == [MULTI.leagues]  # narrowed to choose episodes, read for all
    assert (store / "podcasts" / MULTI.key / old.name).exists()


def test_without_a_key_nothing_is_extracted(tmp_path: Path) -> None:
    with (
        patch.object(run, "fetch_episodes", lambda s, a, b: [Episode(s, ITEM)]),
        patch.object(run, "ensure_transcript", lambda *a, **k: None),
    ):
        rep = run.read_feeds(
            tmp_path,
            datetime.now(timezone.utc),
            datetime.now(timezone.utc),
            shows=[MULTI],
            api_key="",
            log=False,
        )
    assert rep.episodes == 1 and not rep.by_league and rep.missing


def test_a_dead_feed_is_reported_and_the_rest_still_read(tmp_path: Path) -> None:
    def fetch(show: object, a: object, b: object) -> list[Episode]:
        raise OSError("dns")

    with patch.object(run, "fetch_episodes", fetch):
        rep = run.read_feeds(
            tmp_path,
            datetime.now(timezone.utc),
            datetime.now(timezone.utc),
            shows=[MULTI],
            api_key="",
            log=False,
        )
    assert rep.failed and rep.episodes == 0


def _row(league: str, pick_id: str, result: str, **kw: str) -> dict[str, str]:
    base = {f: "" for f in LEDGER_FIELDS}
    base.update(
        date="2026-09-13",
        league=league,
        pick_id=pick_id,
        show="wisekracks",
        show_name="WiseKracks (WagerTalk)",
        kind="official",
        market="game_ats",
        matchup="BUF @ KC",
        selection="BUF +2.5",
        line="2.5",
        price="-110",
        price_source="stated",
        stake="1",
        result=result,
        units="0.909" if result == "win" else "-1.000",
        published=PUB,
    )
    base.update(kw)
    return base


def test_audit_splits_league_host_kind_and_market() -> None:
    rows = [
        _row(NFL, "a", "win", host="Bill Krackomberger"),
        _row(NFL, "b", "loss", matchup="NYJ @ NYG", selection="NYG -1.5"),
        _row(NFL, "c", "win", kind="lean", market="game_total", selection="Under 47.5"),
        _row(CFB, "d", "loss", matchup="Iowa @ OSU"),
        _row(NFL, "e", ""),
    ]
    got = {(a.league, a.host, a.kind, a.market): a.rec for a in audit.audit(rows)}
    assert got[(NFL, "", "official", "all")].wlp == "1-1"
    assert got[(NFL, "Bill Krackomberger", "official", "all")].wlp == "1-0"
    assert got[(NFL, "", "lean", "all")].wlp == "1-0"
    assert got[(CFB, "", "official", "all")].wlp == "0-1"
    assert got[(NFL, "", "official", "all")].underpowered


def test_a_show_counts_one_shared_opinion_once() -> None:
    rows = [
        _row(NFL, "a", "win", host="Bill Krackomberger"),
        _row(NFL, "b", "win", host="Jon Orlando"),
    ]
    got = {(a.host, a.market): a.rec.wlp for a in audit.audit(rows) if a.kind == "official"}
    assert got[("", "all")] == "1-0"
    assert got[("Bill Krackomberger", "all")] == got[("Jon Orlando", "all")] == "1-0"


def test_workbook_and_ledgers_read_every_league(tmp_path: Path) -> None:
    from datetime import date

    from engine_common.podcasts.picks import update_ledger

    update_ledger(ledger_path(tmp_path, NFL), [_row(NFL, "a", "win")], [date(2026, 9, 13)])
    update_ledger(ledger_path(tmp_path, CFB), [_row(CFB, "b", "loss")], [date(2026, 9, 13)])
    rows = audit.all_ledgers(tmp_path)
    assert sorted(r["league"] for r in rows) == [CFB, NFL]
    wb = load_workbook(io.BytesIO(audit.workbook(rows)))
    assert wb.sheetnames == ["Handicapper audit", "Picks", "Notes"]
    leagues = {c.value for c in wb["Handicapper audit"]["A"][1:]}
    assert leagues == {"NFL", "CFB"}
    assert "NFL" in audit.table(rows, league=NFL)


def test_state_pull_fills_in_without_overwriting(tmp_path: Path) -> None:
    remote, local = tmp_path / "r.json", tmp_path / "l" / "r.json"
    remote.write_text("remote")
    assert fill_in_file(remote, local) and local.read_text() == "remote"
    remote.write_text("changed")
    assert not fill_in_file(remote, local) and local.read_text() == "remote"
