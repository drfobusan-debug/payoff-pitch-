import smtplib
from dataclasses import asdict, replace
from datetime import date, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from nba_engine import cli
from nba_engine.audit import ledger
from nba_engine.config import load_config
from nba_engine.data import boxes
from nba_engine.data.capture import QuoteRow
from nba_engine.market.board import selections
from nba_engine.output import card as card_mod
from nba_engine.output import context, teamstats
from nba_engine.output import email as email_mod
from nba_engine.output.card import NO_MODEL_TEXT, PAPER_NOTE, build_card, render_html
from nba_engine.output.excel import SELECTION_HEADER, build_workbook
from nba_engine.schemas import GameResult, TeamBox

DAY = date(2025, 12, 10)
LATE, EARLY = "BOS @ NYK", "PHX @ DET"  # alphabetical order is the reverse of tip order
TIPS = {"late": "2025-12-11T03:00:00Z", "early": "2025-12-11T00:00:00Z"}
EVENTS = {"late": LATE, "early": EARLY}


def q(eid, market, side, line, book, am, opp, at="2025-12-10T22:00:00Z"):
    return QuoteRow(at, DAY.isoformat(), EVENTS[eid], eid, market, side, "", line, book, am, opp)


def board(at="2025-12-10T22:00:00Z"):
    out = []
    for eid in EVENTS:
        away, _, home = EVENTS[eid].partition(" @ ")
        for book, fav, dog in (
            ("draftkings", -150, 130),
            ("betmgm", -145, 125),
            ("fanduel", -155, 135),
        ):
            out += [
                q(eid, "game_ml", home, None, book, fav, dog, at),
                q(eid, "game_ml", away, None, book, dog, fav, at),
                q(eid, "game_ats", home, -3.5, book, -110, -110, at),
                q(eid, "game_ats", away, 3.5, book, -110, -110, at),
                q(eid, "game_total", "over", 221.5, book, -108, -112, at),
                q(eid, "game_total", "under", 221.5, book, -112, -108, at),
            ]
    return out


def priced(tag="t1", at="2025-12-10T22:00:00Z"):
    pre = ledger.pregame(board(at), TIPS)
    return ledger.price(
        selections(pre),
        ledger.quote_index(pre),
        tips=TIPS,
        shift={},
        versions="over_bias=v1",
        pass_tag=tag,
    )


def as_buy(r, model_prob=0.66):
    return replace(r, model_prob=model_prob, ev=0.05, tier="Strong buy", gates="")


@pytest.fixture
def scratch(monkeypatch, tmp_path):
    monkeypatch.setenv("NBAE_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NBAE_OUTPUT_DIR", str(tmp_path / "out"))
    monkeypatch.setenv("NBAE_STATE_SYNC", "0")
    monkeypatch.delenv("NBAE_PRESEASON", raising=False)
    for var in (
        "GMAIL_APP_PASSWORD",
        "GMAIL_USER",
        "EMAIL_ADDRESS",
        "NBAE_EMAIL_TO",
        "MLBE_EMAIL_TO",
    ):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.setattr(cli, "render_pdf", lambda html: b"%PDF-1.4 fake")
    rows = priced()
    path = ledger.pass_path(tmp_path, DAY, rows[0].priced_at, "t1")
    ledger.write_once(rows, path)
    return tmp_path, path


def run_card(*extra):
    return cli.main(["card", "--date", DAY.isoformat(), "--no-sync", "--offline", *extra])


# -- order -------------------------------------------------------------------
def test_games_in_tip_order_with_et_tip_in_each_header():
    card = build_card(priced(), day=DAY)
    assert [g.matchup for g in card.games] == [EARLY, LATE]
    page = render_html(card)
    assert page.index("Phoenix Suns at Detroit Pistons") < page.index(
        "Boston Celtics at New York Knicks"
    )
    assert "7:00 PM ET" in page and "10:00 PM ET" in page


def test_best_bets_in_tip_order_strongest_first_within_a_game():
    rows = priced()
    late_ml = next(r for r in rows if r.event_id == "late" and r.market == "game_ml")
    late_tot = next(r for r in rows if r.event_id == "late" and r.market == "game_total")
    early = next(r for r in rows if r.event_id == "early" and r.market == "game_ats")
    picks = {id(late_ml): as_buy(late_ml), id(late_tot): replace(as_buy(late_tot), ev=0.09)}
    picks[id(early)] = replace(as_buy(early), ev=0.01)
    card = build_card([picks.get(id(r), r) for r in rows], day=DAY)
    assert [(b.event_id, b.market) for b in card.buys()] == [
        ("early", "game_ats"),
        ("late", "game_total"),
        ("late", "game_ml"),
    ]


# -- card --------------------------------------------------------------------
def test_card_from_fixture_ledger_says_market_only_and_shows_gates():
    page = render_html(build_card(priced(), day=DAY))
    assert NO_MODEL_TEXT in page
    assert PAPER_NOTE in page
    assert "no_model" in page
    assert "DK -150" in page and "MGM" in page  # execution prices beside the fair
    assert "Moneyline" in page and "Spread" in page and "Total" in page


def test_kelly_only_on_the_model_probability():
    assert card_mod.kelly(0.6, 100) == pytest.approx(0.2)
    assert card_mod.kelly(0.4, 100) == 0.0
    r = priced()[0]
    assert card_mod.row_kelly(r) is None  # market only: no Kelly invented
    assert card_mod.row_kelly(as_buy(r)) is not None


def test_cli_card_writes_artifacts_from_the_ledger(scratch, capsys):
    root, _ = scratch
    assert run_card() == 0
    out = root / "out"
    names = sorted(p.name for p in out.iterdir())
    assert names == [f"PayoffPitch_NBA_{DAY}.{ext}" for ext in ("html", "pdf", "xlsx")]
    assert "2 games" in capsys.readouterr().out


def test_no_ledger_pass_is_an_error(scratch):
    root, path = scratch
    path.unlink()
    assert run_card() == 1


def test_pdf_failure_costs_only_the_pdf(scratch, monkeypatch):
    root, _ = scratch

    def boom(html):
        raise OSError("no pango")

    monkeypatch.setattr(cli, "render_pdf", boom)
    assert run_card() == 0
    out = root / "out"
    assert (out / f"PayoffPitch_NBA_{DAY}.xlsx").exists()
    assert not (out / f"PayoffPitch_NBA_{DAY}.pdf").exists()


def test_pricing_neutrality_card_never_changes_ledger_rows(scratch):
    root, path = scratch
    before_bytes = path.read_bytes()
    before = [asdict(r) for r in ledger.passes(root, DAY)]
    assert run_card() == 0
    assert run_card() == 0
    assert path.read_bytes() == before_bytes
    assert [asdict(r) for r in ledger.passes(root, DAY)] == before
    assert sorted(p.name for p in path.parent.iterdir()) == [path.name]


def test_preseason_card_reads_its_own_tree(monkeypatch, tmp_path):
    monkeypatch.setenv("NBAE_PRESEASON", "1")
    monkeypatch.setenv("NBAE_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NBAE_STATE_SYNC", "0")
    monkeypatch.delenv("NBAE_OUTPUT_DIR", raising=False)
    monkeypatch.delenv("GMAIL_APP_PASSWORD", raising=False)
    monkeypatch.setattr(cli, "render_pdf", lambda html: b"%PDF")
    rows = priced()
    ledger.write_once(rows, ledger.pass_path(tmp_path / "preseason", DAY, rows[0].priced_at, "t"))
    assert run_card() == 0
    out = tmp_path / "preseason" / "output"
    assert (out / f"PayoffPitch_NBA_Preseason_{DAY}.xlsx").exists()
    assert "NBA Preseason Slate" in (out / f"PayoffPitch_NBA_Preseason_{DAY}.html").read_text()


# -- workbook ----------------------------------------------------------------
def test_workbook_sheets_columns_frozen_and_filtered():
    rows = priced()
    rows[0] = as_buy(rows[0])
    wb = load_workbook(BytesIO(build_workbook(build_card(rows, day=DAY))))
    assert wb.sheetnames == [
        "Buys",
        "Selections",
        "Games",
        "Injuries",
        "Record",
        "Integrity",
        "Legend",
    ]
    sel = wb["Selections"]
    header = [c.value for c in sel[1]]
    assert header == list(SELECTION_HEADER)
    for col in (
        "Matchup",
        "Tip (ET)",
        "Market",
        "Side",
        "Line",
        "Book",
        "Price",
        "Fair prob",
        "Model prob",
        "EV",
        "Kelly",
        "Tier",
        "Gates",
        "Params version",
    ):
        assert col in header
    assert sel.max_row == len(rows) + 1
    assert sel.freeze_panes == "A2" and sel.auto_filter.ref.startswith("A1:")
    tips = [r[0] for r in sel.iter_rows(min_row=2, values_only=True)]
    assert tips == sorted(tips)
    buys = wb["Buys"]
    assert buys.max_row == 2 and buys["Q2"].value == "Strong buy"  # one buy
    assert buys.freeze_panes == "A2"


def test_record_sheet_carries_n_and_intervals():
    rows = [
        replace(r, outcome="win" if i % 2 else "loss", pnl=1.0 if i % 2 else -1.0)
        for i, r in enumerate(priced())
    ]
    wb = load_workbook(BytesIO(build_workbook(build_card(rows, day=DAY, graded=rows, draws=20))))
    rec = wb["Record"]
    header = [c.value for c in rec[1]]
    assert {"n", "ROI", "ROI 95% lo", "ROI 95% hi", "CLV EV 95% lo"} <= set(header)
    cuts = {r[0] for r in rec.iter_rows(min_row=2, values_only=True)}
    assert {"buys", "board"} <= cuts


# -- email -------------------------------------------------------------------
class FakeSMTP:
    sent: list = []

    def __init__(self, host, port, context=None, timeout=None):
        self.host, self.port = host, port

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def login(self, user, password):
        self.user, self.password = user, password

    def send_message(self, msg):
        FakeSMTP.sent.append((self, msg))


@pytest.fixture
def fake_smtp(monkeypatch):
    FakeSMTP.sent = []
    monkeypatch.setattr(email_mod.smtplib, "SMTP_SSL", FakeSMTP)
    monkeypatch.setattr(smtplib, "SMTP", None)  # any plain-SMTP path would blow up
    return FakeSMTP.sent


def test_email_without_config_keeps_files_and_returns_zero(scratch, fake_smtp):
    root, _ = scratch
    assert run_card("--email") == 0
    assert (root / "out" / f"PayoffPitch_NBA_{DAY}.xlsx").exists()
    assert fake_smtp == []


def test_email_sends_pdf_and_workbook_via_fake_smtp(scratch, fake_smtp, monkeypatch):
    monkeypatch.setenv("GMAIL_USER", "me@example.com")
    monkeypatch.setenv("GMAIL_APP_PASSWORD", "abcd efgh ijkl mnop")
    assert run_card("--email", "--to", "you@example.com") == 0
    ((server, msg),) = fake_smtp
    assert server.host == "smtp.gmail.com" and server.port == 465
    assert server.password == "abcdefghijklmnop"
    assert msg["To"] == "you@example.com"
    names = [p.get_filename() for p in msg.iter_attachments()]
    assert names == [f"PayoffPitch_NBA_{DAY}.pdf", f"PayoffPitch_NBA_{DAY}.xlsx"]


def test_size_guard_drops_attachments_measured_after_base64(monkeypatch, fake_smtp):
    monkeypatch.setenv("GMAIL_APP_PASSWORD", "pw")
    monkeypatch.setenv("GMAIL_USER", "me@example.com")
    big, small = b"x" * 60_000, b"y" * 1_000
    sent = email_mod.send_package(
        load_config(),
        subject="s",
        html_body="<html><body>card</body></html>",
        text_body="card",
        attachments=[("card.pdf", small), ("book.xlsx", big)],
        max_bytes=70_000,  # 60 kB raw is ~82 kB once base64'd
    )
    assert sent.attached == ("card.pdf",) and sent.left_out == ("book.xlsx",)
    assert sent.size <= 70_000
    ((_, msg),) = fake_smtp
    assert "book.xlsx" in msg.get_body(("plain",)).get_content()
    budget = email_mod.attachment_budget(email_mod.GMAIL_LIMIT_BYTES)
    assert budget / email_mod.GMAIL_LIMIT_BYTES == pytest.approx(1 / 1.37, abs=0.01)


def test_email_needs_a_recipient(monkeypatch):
    monkeypatch.setenv("GMAIL_APP_PASSWORD", "pw")
    for var in ("GMAIL_USER", "EMAIL_ADDRESS", "NBAE_EMAIL_TO", "MLBE_EMAIL_TO"):
        monkeypatch.delenv(var, raising=False)
    with pytest.raises(email_mod.EmailNotConfigured):
        email_mod.send_package(load_config(), subject="s", html_body="", text_body="")


# -- context -----------------------------------------------------------------
def test_rest_and_travel_from_prior_finals_and_unknown_when_a_day_is_missing():
    games = [context.Matchup("late", "BOS", "NYK", TIPS["late"])]
    prior = [GameResult("1", date(2025, 12, 9), "NYK", "BOS", "STATUS_FINAL")]
    days = {DAY - timedelta(days=k) for k in range(1, context.LOOKBACK_DAYS + 1)}
    rest = context.rest(DAY, games, prior, days)
    bos, nyk = rest[("late", "BOS")], rest[("late", "NYK")]
    assert bos.known and bos.b2b and nyk.b2b
    assert "back-to-back" in bos.text() and "mi travel" in nyk.text()
    # A team idle in the window can't be called rested if a day in it went unread.
    idle = [context.Matchup("x", "PHX", "DET", TIPS["early"])]
    assert not context.rest(DAY, idle, prior, days - {DAY - timedelta(days=2)})[("x", "PHX")].known
    assert context.rest(DAY, idle, prior, days)[("x", "PHX")].known


def test_line_move_open_to_now_from_archived_boards(scratch):
    root, _ = scratch
    from nba_engine.data import capture

    early = board("2025-12-10T16:00:00Z")
    late = [
        replace(r, american=-200, opposite_american=170)
        if r.market == "game_ml" and r.side == "DET"
        else replace(r, american=170, opposite_american=-200)
        if r.market == "game_ml"
        else r
        for r in board("2025-12-10T22:00:00Z")
    ]
    for rows in (early, late):
        capture.write_snapshot(rows, root, DAY, label="board", captured_at=rows[0].captured_at)
    games = context.games_of(priced())
    moves, opened, now = context.moves(root, DAY, games)["early"]
    ml = next(m for m in moves if m.key == "ml")
    assert ml.now > ml.open
    assert opened < now


def test_late_card_is_skipped_once_every_game_has_tipped(scratch, capsys):
    root, _ = scratch
    assert run_card("--late") == 0
    assert "no late tip left" in capsys.readouterr().out
    assert not (root / "out").exists()


def test_late_card_sends_while_a_game_is_untipped(scratch, capsys, monkeypatch):
    class Frozen(cli.datetime):
        @classmethod
        def now(cls, tz=None):
            return cli.datetime(2025, 12, 10, 12, 0, tzinfo=tz)

    monkeypatch.setattr(cli, "datetime", Frozen)
    assert run_card("--late") == 0
    assert "sending the late package" in capsys.readouterr().out


# -- exec books one line over, integrity and bands on the card -----------------


def test_main_line_without_dk_mgm_points_at_the_line_they_did_post():
    """Consensus main at -3.5 (three books); DK/MGM sit on -4: the card says so instead
    of 'neither posted it'."""
    eid, home, away = "late", "NYK", "BOS"
    rows = []
    for book in ("fanduel", "bovada", "pointsbetus"):
        rows += [
            q(eid, "game_ats", home, -3.5, book, -110, -110),
            q(eid, "game_ats", away, 3.5, book, -110, -110),
        ]
    for book in ("draftkings", "betmgm"):
        rows += [
            q(eid, "game_ats", home, -4.0, book, -105, -115),
            q(eid, "game_ats", away, 4.0, book, -115, -105),
        ]
    pre = ledger.pregame(rows, TIPS)
    priced_rows = ledger.price(
        selections(pre), ledger.quote_index(pre), tips=TIPS, shift={}, versions="v", pass_tag="t"
    )
    card = build_card(priced_rows, day=DAY)
    (game,) = card.games
    main = game.main("game_ats")
    assert [r.line for r in main] == [3.5, -3.5]
    sib = card_mod.exec_sibling(game, main[0])
    assert sib is not None and sib.line == 4.0
    page = render_html(card)
    assert "not at this number; at +4:" in page and "neither DK nor BetMGM" not in page


def test_integrity_line_and_schedule_bands_render_only_with_a_graded_ledger():
    from engine_common.bands import BandRead

    rows = priced()
    bare = render_html(build_card(rows, day=DAY))
    assert "Ledger integrity" not in bare and "Schedule bands" not in bare
    graded = [
        replace(r, outcome="win", pnl=0.9, close_status="close", close_captured_at=r.priced_at)
        for r in rows
    ]
    band = BandRead("rest_host_ats", [1.0, -1.0, 0.9], [1.0, -1.0, 0.9])
    page = render_html(build_card(rows, day=DAY, graded=graded, draws=10, bands=[band]))
    assert "Ledger integrity:" in page and "no duplicates" in page
    assert (
        "Schedule bands" in page and "rest_host_ats" in page and "probation (3/100 games)" in page
    )


# -- season ratings ----------------------------------------------------------
def _box(team, fgm=40, fga=88, fg3m=12, ftm=16, fta=20, oreb=10, dreb=34, tov=13):
    return TeamBox(team, fgm, fga, fg3m, 35, ftm, fta, oreb, dreb, tov, 18)


def _final(eid, day, away, home, ap, hp, *, season_type=2, neutral=False, abox=None, hbox=None):
    return GameResult(
        eid,
        day,
        away,
        home,
        "STATUS_FINAL",
        (ap, 0, 0, 0),
        (hp, 0, 0, 0),
        away_box=abox or _box(away),
        home_box=hbox or _box(home),
        neutral=neutral,
        season_type=season_type,
    )


def test_season_table_splits_home_and_away_per_100_and_skips_preseason_and_neutral():
    d = date(2025, 11, 1)
    games = [
        _final("1", d, "BOS", "NYK", 100, 120),
        _final("2", d, "NYK", "BOS", 110, 90),
        _final("3", d, "NYK", "BOS", 150, 50, season_type=1),  # preseason: left out
        _final("4", d, "BOS", "NYK", 100, 100, neutral=True),  # neutral: no venue split
    ]
    t = teamstats.season_table(games)
    poss = games[0].possessions
    nyk = t["NYK"]
    assert nyk.games == 3 and nyk.home.games == 1 and nyk.away.games == 1
    assert nyk.home.off.value == pytest.approx(100 * 120 / poss)
    assert nyk.home.dfn.value == pytest.approx(100 * 100 / poss)
    assert nyk.home.off.rank == 1 and t["BOS"].home.dfn.rank == 2
    assert t["BOS"].away.net.value == pytest.approx(-100 * 20 / poss)


def test_four_factors_rank_defense_on_what_opponents_did():
    d = date(2025, 11, 1)
    good = _box("NYK", fgm=45, fg3m=15, tov=8, oreb=14)
    bad = _box("BOS", fgm=35, fg3m=8, tov=18, oreb=6)
    t = teamstats.season_table([_final("1", d, "BOS", "NYK", 95, 120, abox=bad, hbox=good)])
    nyk, bos = t["NYK"], t["BOS"]
    assert nyk.offense.efg.value == pytest.approx((45 + 7.5) / 88)
    assert nyk.offense.efg.rank == 1 and nyk.offense.tov.rank == 1
    assert nyk.defense.efg.value == bos.offense.efg.value and nyk.defense.efg.rank == 1
    assert nyk.offense.orb.value == pytest.approx(14 / (14 + 34))


def test_card_prints_ratings_and_four_factors_or_says_none_are_archived(scratch):
    root, _ = scratch
    rows = ledger.of_record(ledger.passes(root, DAY))
    played = [_final("1", DAY - timedelta(days=3), "BOS", "NYK", 100, 110)]
    boxes.write_results(root, DAY - timedelta(days=3), played)
    page = render_html(build_card(rows, day=DAY, context=context.gather(root, DAY, rows)))
    assert "Per 100 poss" in page and "Four factors" in page
    assert "BOS O vs NYK D" in page and "NYK O vs BOS D" in page and "class='tn'" in page
    bare = render_html(build_card(rows, day=DAY))
    assert "no regular-season box scores archived yet" in bare and "Four factors" not in bare
