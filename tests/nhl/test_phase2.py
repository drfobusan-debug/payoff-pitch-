"""Phase 2: joint sim, market extraction, pricing gates, ledger, grading, outputs."""

from __future__ import annotations

from datetime import date as Date
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pytest

from engine_common.isotonic import IsotonicMap, pav
from engine_common.odds import american_to_prob, prob_to_american
from nhl_engine.audit import grade, scorecard
from nhl_engine.audit.ledger import LedgerRow, grade_rows, load_rows, row_from, write_once
from nhl_engine.calibration import Calibrator, ConfidenceShrink
from nhl_engine.config import Config, GateParams, SimParams
from nhl_engine.data.book_rules import BookRule, BookRules
from nhl_engine.data.capture import QuoteRow
from nhl_engine.features.starters import Starter
from nhl_engine.market.board import Selection, matchups, selections
from nhl_engine.market.pricing import Tier, ev_per_unit, stamp
from nhl_engine.models import markets
from nhl_engine.models.goals import game_rates, goalie_factor, side_rates
from nhl_engine.models.periods import SimResult, ot_home_prob, simulate
from nhl_engine.pipeline import TeamInputs, price_game
from nhl_engine.schemas import GameResult, PeriodScore

LEAGUE = {
    "xgf60_5v5": 2.55,
    "xga60_5v5": 2.55,
    "fin60_5v5": 0.0,
    "pp_xgf60": 7.4,
    "pp_xga60": 0.9,
    "pk_xga60": 7.4,
    "sh_xgf60": 0.9,
    "pen_taken60": 3.4,
    "pen_drawn60": 3.4,
}
STRONG = {**LEAGUE, "xgf60_5v5": 3.0, "xga60_5v5": 2.2, "pp_xgf60": 8.5, "pk_xga60": 6.5}
WEAK = {**LEAGUE, "xgf60_5v5": 2.2, "xga60_5v5": 3.0, "pp_xgf60": 6.3, "pk_xga60": 8.6}
P = SimParams(draws=4000)


def _sim(home=LEAGUE, away=LEAGUE, hg=0.0, ag=0.0, seed=7, params=P) -> SimResult:
    r = game_rates(home, away, LEAGUE, home_goalie_gsax60=hg, away_goalie_gsax60=ag, params=params)
    return simulate(r, params, seed=seed)


# -- rates --------------------------------------------------------------------------


def test_goalie_factor_only_from_gsax():
    assert goalie_factor(0.0) == 1.0
    assert goalie_factor(0.5) < 1.0 < goalie_factor(-0.5)
    assert goalie_factor(100.0) == 0.5


def test_side_rates_team_specific_shorthanded():
    hot_sh = {**LEAGUE, "sh_xgf60": 1.6}
    leaky_pp = {**LEAGUE, "pp_xga60": 1.5}
    base = side_rates(
        LEAGUE, LEAGUE, LEAGUE, is_home=False, opp_goalie_gsax60=0, own_goalie_gsax60=0, params=P
    )
    up = side_rates(
        hot_sh, leaky_pp, LEAGUE, is_home=False, opp_goalie_gsax60=0, own_goalie_gsax60=0, params=P
    )
    assert up.sh > base.sh
    assert up.g5 == pytest.approx(base.g5)


def test_home_share_and_goalie_applied():
    r = game_rates(LEAGUE, LEAGUE, LEAGUE, home_goalie_gsax60=0.0, away_goalie_gsax60=0.0, params=P)
    assert r.home.g5 > r.away.g5
    assert r.home.g5 / (r.home.g5 + r.away.g5) == pytest.approx(P.home_xg_share, abs=1e-6)
    elite_home = game_rates(
        LEAGUE, LEAGUE, LEAGUE, home_goalie_gsax60=0.6, away_goalie_gsax60=0.0, params=P
    )
    assert elite_home.away.g5 < r.away.g5  # away scores on the elite home goalie
    assert elite_home.home.g5 == pytest.approx(r.home.g5)


# -- simulation ---------------------------------------------------------------------


def test_seeded_reproducible_and_shapes():
    a, b = _sim(seed=3), _sim(seed=3)
    assert np.array_equal(a.ph, b.ph) and np.array_equal(a.so_h, b.so_h)
    assert a.ph.shape == (P.draws, 3)
    c = _sim(seed=4)
    assert not np.array_equal(a.ph, c.ph)


def test_regulation_ot_so_accounting():
    s = _sim()
    # OT/SO only after a regulation tie, exactly one OT-or-SO goal, SO credited as one goal
    assert not np.any((s.ot_h + s.ot_a + s.so_h + s.so_a > 0) & ~s.went_ot)
    decided = s.ot_h + s.ot_a + s.so_h + s.so_a
    assert np.all(decided[s.went_ot] == 1)
    assert np.all(s.final_h != s.final_a)
    assert 0.15 < s.went_ot.mean() < 0.30
    assert 0.03 < s.went_so.mean() < 0.12
    assert 5.4 < (s.final_h + s.final_a).mean() < 6.8


def test_period_multipliers_visible():
    s = _sim(seed=11, params=SimParams(draws=20000))
    p1 = (s.ph[:, 0] + s.pa[:, 0]).mean()
    p2 = (s.ph[:, 1] + s.pa[:, 1]).mean()
    p3 = (s.ph[:, 2] + s.pa[:, 2]).mean()
    assert p1 < p2 < p3  # P2 study multiplier; P3 adds empty-net goals


def test_empty_net_lifts_third_period_and_can_be_switched_off():
    on = SimParams(draws=20000)
    off = SimParams(draws=20000, pull_prob={1: 0.0, 2: 0.0, 3: 0.0})
    s_on, s_off = _sim(seed=5, params=on), _sim(seed=5, params=off)
    p3_on = (s_on.ph[:, 2] + s_on.pa[:, 2]).mean()
    p3_off = (s_off.ph[:, 2] + s_off.pa[:, 2]).mean()
    assert p3_on > p3_off + 0.1
    # a 2-goal regulation margin is more common with the net empty
    m_on = np.abs(s_on.reg_h - s_on.reg_a)
    m_off = np.abs(s_off.reg_h - s_off.reg_a)
    assert (m_on >= 2).mean() > (m_off >= 2).mean()


def test_special_teams_move_totals():
    hot = {**LEAGUE, "pen_taken60": 6.0, "pen_drawn60": 6.0}
    quiet = {**LEAGUE, "pen_taken60": 1.5, "pen_drawn60": 1.5}
    s_hot = _sim(hot, hot, seed=2, params=SimParams(draws=20000))
    s_quiet = _sim(quiet, quiet, seed=2, params=SimParams(draws=20000))
    assert (s_hot.final_h + s_hot.final_a).mean() > (s_quiet.final_h + s_quiet.final_a).mean()


def test_ot_block_shrunk_and_goalie_aware():
    even = game_rates(LEAGUE, LEAGUE, LEAGUE, home_goalie_gsax60=0, away_goalie_gsax60=0, params=P)
    strong = game_rates(STRONG, WEAK, LEAGUE, home_goalie_gsax60=0, away_goalie_gsax60=0, params=P)
    p_even, p_strong = ot_home_prob(even, P), ot_home_prob(strong, P)
    assert 0.52 < p_even < 0.58
    assert p_strong > p_even
    assert p_strong < 0.75  # the 3v3 link is far flatter than the 5v5 gap
    hot_goalie = game_rates(
        LEAGUE, LEAGUE, LEAGUE, home_goalie_gsax60=0.8, away_goalie_gsax60=0, params=P
    )
    assert ot_home_prob(hot_goalie, P) > p_even


def test_shootout_near_coin_flip():
    s = _sim(STRONG, WEAK, seed=9, params=SimParams(draws=40000))
    so = s.went_so
    assert so.sum() > 500
    assert abs(s.so_h[so].mean() - 0.5) < 0.05


def test_strength_moves_moneyline():
    s = _sim(STRONG, WEAK, seed=1)
    assert markets.moneyline(s, True).win > 0.62


# -- market extraction --------------------------------------------------------------


def test_markets_read_from_one_joint_draw():
    s = _sim(STRONG, WEAK, seed=8)
    ml = markets.moneyline(s, True).win
    ml3 = markets.moneyline3(s, "H", "H", "A")
    draw = markets.moneyline3(s, "draw", "H", "A")
    assert ml3.win + draw.win + markets.moneyline3(s, "A", "H", "A").win == pytest.approx(1.0)
    assert ml > ml3.win  # OT/SO wins count for the 2-way only
    # -1.5 cover + 1.5 dog cover partition the space (no pushes at .5)
    pl_h = markets.price(
        s,
        market="game_pl",
        side="H",
        entity="",
        line=-1.5,
        ot_rule="incl_ot_so",
        home="H",
        away="A",
    )
    pl_a = markets.price(
        s, market="game_pl", side="A", entity="", line=1.5, ot_rule="incl_ot_so", home="H", away="A"
    )
    assert pl_h.win + pl_a.win == pytest.approx(1.0)
    # a SO win never covers -1.5
    so_cover = ((s.final_h - s.final_a) > 1.5) & s.went_so
    assert not so_cover.any()
    over = markets.price(
        s,
        market="game_total",
        side="over",
        entity="",
        line=6.5,
        ot_rule="incl_ot_so",
        home="H",
        away="A",
    )
    under = markets.price(
        s,
        market="game_total",
        side="under",
        entity="",
        line=6.5,
        ot_rule="incl_ot_so",
        home="H",
        away="A",
    )
    assert over.win + under.win == pytest.approx(1.0)


def test_totals_settlement_rules_order():
    s = _sim(seed=6)
    kw = dict(market="team_total", side="over", entity="H", line=2.5, home="H", away="A")
    reg = markets.price(s, ot_rule="reg_only", **kw).win
    ot = markets.price(s, ot_rule="incl_ot", **kw).win
    full = markets.price(s, ot_rule="incl_ot_so", **kw).win
    assert reg <= ot <= full
    assert full > reg


def test_period_markets_and_push():
    s = _sim(seed=12)
    p1 = markets.price(
        s, market="p1_ml", side="H", entity="", line=None, ot_rule="reg_only", home="H", away="A"
    )
    assert p1.push > 0.25  # first-period draws are common
    p1_3 = markets.price(
        s,
        market="p1_ml3",
        side="draw",
        entity="",
        line=None,
        ot_rule="reg_only",
        home="H",
        away="A",
    )
    assert p1_3.win == pytest.approx(p1.push)
    tot = markets.price(
        s,
        market="p2_total",
        side="over",
        entity="",
        line=1.5,
        ot_rule="reg_only",
        home="H",
        away="A",
    )
    assert 0.3 < tot.win < 0.8
    with pytest.raises(markets.UnpricedMarket):
        markets.price(
            s,
            market="sk_sog",
            side="over",
            entity="x",
            line=2.5,
            ot_rule="incl_ot",
            home="H",
            away="A",
        )


# -- board + pricing ----------------------------------------------------------------


def _q(
    market,
    side,
    american,
    opp,
    *,
    line=None,
    entity="",
    book="dk",
    at="2026-10-01T15:00:00Z",
    matchup="A @ H",
):
    return QuoteRow(
        at, "2026-10-01", matchup, "ev1", market, side, entity, line, book, american, opp
    )


def test_board_consensus_devig_and_one_way():
    rows = [
        _q("game_ml", "H", -130, 110),
        _q("game_ml", "A", 110, -130),
        _q("game_ml", "H", -125, 105, book="fd"),
        _q("game_ml", "A", 105, -125, book="fd"),
        _q("game_ml", "H", -140, 115, at="2026-10-01T14:00:00Z"),  # superseded snapshot
        _q("team_total", "over", -110, None, line=2.5, entity="H"),
        _q("game_ml3", "H", 130, None),
        _q("game_ml3", "A", 180, None),
        _q("game_ml3", "draw", 300, None),
    ]
    sels = {s.key: s for s in selections(rows)}
    h = sels[("A @ H", "game_ml", "H", "", None)]
    assert h.books == 2 and h.best_american == -125 and h.best_book == "fd"
    assert 0.53 < h.consensus < 0.57 and h.consensus < american_to_prob(-130)
    tt = sels[("A @ H", "team_total", "over", "H", 2.5)]
    assert tt.one_way
    ml3 = sels[("A @ H", "game_ml3", "draw", "", None)]
    assert not ml3.one_way and 0.20 < ml3.consensus < 0.25
    assert matchups(rows) == {"A @ H": ("A", "H", "ev1")}


def _sel(
    market="game_ml",
    side="H",
    american=-120,
    consensus=0.52,
    line=None,
    entity="",
    two_sided=2,
    at="2026-10-01T15:00:00Z",
):
    return Selection(
        "A @ H", market, side, entity, line, consensus, 2, two_sided, american, "dk", at
    )


NOW = datetime(2026, 10, 1, 15, 30, tzinfo=timezone.utc)


def test_ev_and_american_round_trip():
    assert prob_to_american(american_to_prob(-150)) == pytest.approx(-150)
    assert prob_to_american(american_to_prob(240)) == pytest.approx(240)
    assert ev_per_unit(markets.Prob(0.5), 100) == pytest.approx(0.0)
    assert ev_per_unit(markets.Prob(0.55), 100) == pytest.approx(0.10)
    assert ev_per_unit(markets.Prob(0.55, push=0.5), 100) == pytest.approx(0.05)


def test_gates_are_stamped_not_filtered():
    thr = GateParams()
    ok = stamp(
        _sel(),
        markets.Prob(0.56),
        ot_rule="incl_ot_so",
        thr=thr,
        quote_age_minutes=5,
        goalie_confirmed=True,
        settlement_gate=None,
    )
    assert ok.pass_gate and ok.tier == Tier.STRONG and ok.is_buy
    mod = stamp(
        _sel(american=100, consensus=0.50),
        markets.Prob(0.53),
        ot_rule="incl_ot_so",
        thr=thr,
        quote_age_minutes=5,
        goalie_confirmed=True,
        settlement_gate=None,
    )
    assert mod.tier == Tier.MODERATE
    small = stamp(
        _sel(),
        markets.Prob(0.53),
        ot_rule="incl_ot_so",
        thr=thr,
        quote_age_minutes=5,
        goalie_confirmed=True,
        settlement_gate=None,
    )
    assert "min_edge" in small.gates and small.tier == Tier.PASS
    big = stamp(
        _sel(),
        markets.Prob(0.60),
        ot_rule="incl_ot_so",
        thr=thr,
        quote_age_minutes=5,
        goalie_confirmed=True,
        settlement_gate=None,
    )
    assert "max_edge" in big.gates and big.tier == Tier.PASS
    many = stamp(
        _sel(market="team_total", side="over", entity="H", line=2.5, american=170, two_sided=0),
        markets.Prob(0.56),
        ot_rule="incl_ot_so",
        thr=thr,
        quote_age_minutes=200,
        goalie_confirmed=False,
        settlement_gate="settlement_unverified",
    )
    assert set(many.gates) >= {
        "max_buy_odds",
        "stale_quote",
        "one_way_quote",
        "settlement_unverified",
        "goalie_unconfirmed",
        "probation",
    }
    assert not many.is_buy and many.tier != Tier.PASS  # tier kept, gates say why not
    period = stamp(
        _sel(market="p1_total", side="over", line=1.5),
        markets.Prob(0.56),
        ot_rule="reg_only",
        thr=thr,
        quote_age_minutes=1,
        goalie_confirmed=False,
        settlement_gate=None,
    )
    assert period.gates == ["probation"]  # goalie gate does not apply to period totals
    assert any("edge" in r for r in ok.reasons)


def test_goalie_gate_vetoes_ml_pl_team_total_only():
    thr = GateParams()
    for mk, gated in [
        ("game_ml", True),
        ("game_pl", True),
        ("team_total", True),
        ("game_total", False),
        ("p1_ml", True),
        ("p2_total", False),
    ]:
        r = stamp(
            _sel(market=mk, line=1.5 if "pl" in mk or "total" in mk else None, entity="H"),
            markets.Prob(0.56),
            ot_rule="incl_ot_so",
            thr=thr,
            quote_age_minutes=1,
            goalie_confirmed=False,
            settlement_gate=None,
        )
        assert ("goalie_unconfirmed" in r.gates) is gated, mk


# -- calibration --------------------------------------------------------------------


def test_pav_and_isotonic_map():
    assert pav([0.3, 0.2, 0.5], [1, 1, 1]) == [0.25, 0.25, 0.5]
    m = IsotonicMap.fit([(0.3, 0), (0.35, 1), (0.6, 0), (0.65, 1), (0.9, 1), (0.92, 1)], n_bins=5)
    assert m.apply(0.1) <= m.apply(0.5) <= m.apply(0.95)
    assert IsotonicMap().is_identity and IsotonicMap().apply(0.4) == 0.4


def test_calibrator_identity_until_samples_and_shrink():
    cal = Calibrator()
    p, reasons = cal.apply("game_ml", markets.Prob(0.55))
    assert p.win == 0.55 and reasons[0].startswith("calibration: none")
    p2, reasons2 = cal.apply("game_ml", markets.Prob(0.72))
    assert p2.win == pytest.approx(0.62 + 0.55 * 0.10) and any("shrink" in r for r in reasons2)
    assert ConfidenceShrink().apply(0.5) == 0.5
    rows = [
        LedgerRow(
            "d",
            "m",
            "H",
            "A",
            "game_ml",
            "H",
            "",
            None,
            "incl_ot_so",
            "dk",
            -110,
            1,
            0.5,
            0.6,
            0.0,
            0.1,
            0.05,
            "Pass",
            outcome="win" if i % 2 else "loss",
        )
        for i in range(10)
    ]
    assert Calibrator.fit(rows, min_samples=5).maps["game_ml"].apply(0.6) == pytest.approx(0.5)
    assert "game_ml" not in Calibrator.fit(rows, min_samples=50).maps


def test_calibrator_save_load_basis(tmp_path: Path):
    rows = [
        LedgerRow(
            "d",
            "m",
            "H",
            "A",
            "game_ml",
            "H",
            "",
            None,
            "incl_ot_so",
            "dk",
            -110,
            1,
            0.5,
            0.6,
            0.0,
            0.1,
            0.05,
            "Pass",
            outcome="win",
        )
        for _ in range(6)
    ]
    cal = Calibrator.fit(rows, min_samples=5)
    cal.save(tmp_path / "c.json")
    back = Calibrator.load(tmp_path / "c.json")
    assert back.counts == {"game_ml": 6} and "game_ml" in back.maps
    (tmp_path / "c.json").write_text('{"basis": "old", "markets": {}}')
    assert Calibrator.load(tmp_path / "c.json").maps == {}


# -- ledger, grading, CLV ------------------------------------------------------------


def _res(periods, decided="REG", so_winner=None) -> GameResult:
    ps = [PeriodScore(i + 1, "REG", a, h) for i, (a, h) in enumerate(periods[:3])]
    if decided == "OT":
        a, h = periods[3]
        ps.append(PeriodScore(4, "OT", a, h))
    if decided == "SO":
        ps.append(PeriodScore(4, "OT", 0, 0))
        ps.append(PeriodScore(5, "SO", 1 if so_winner == "A" else 0, 1 if so_winner == "H" else 0))
    return GameResult(1, Date(2026, 10, 1), "A", "H", "OFF", tuple(ps), decided)


def test_grade_regulation_ot_so_and_periods():
    so = _res([(1, 1), (1, 0), (1, 2)], "SO", so_winner="H")  # 3-3 after 60, H wins SO -> 4-3
    assert so.final_home == 4 and so.reg_home == 3

    def g(**kw: object) -> str | None:
        args: dict[str, object] = {
            "market": "game_ml",
            "side": "H",
            "entity": "",
            "line": None,
            "ot_rule": "incl_ot_so",
            **kw,
        }
        return grade.settle(res=so, **args)  # type: ignore[arg-type]

    assert g() == "win" and g(side="A") == "loss"
    assert g(market="game_ml3", side="draw") == "win" and g(market="game_ml3") == "loss"
    assert (
        g(market="game_pl", line=-1.5) == "loss"
        and g(market="game_pl", side="A", line=1.5) == "win"
    )
    assert g(market="game_total", side="over", line=6.5) == "win"
    assert g(market="game_total", side="over", line=6.5, ot_rule="reg_only") == "loss"
    assert g(market="game_total", side="over", line=6.5, ot_rule="incl_ot") == "loss"
    assert g(market="team_total", side="over", entity="H", line=3.5) == "win"
    assert g(market="team_total", side="over", entity="H", line=3.5, ot_rule="reg_only") == "loss"
    assert grade.dual_rule(market="team_total", side="over", entity="H", line=3.5, res=so)
    assert not grade.dual_rule(market="team_total", side="over", entity="H", line=1.5, res=so)
    # periods grade on official period scores
    assert g(market="p1_ml") == "push" and g(market="p1_ml3", side="draw") == "win"
    assert g(market="p2_ml", side="A") == "win" and g(market="p3_pl", line=-0.5) == "win"
    assert g(market="p3_total", side="under", line=2.5) == "loss"
    assert g(market="p1_team_total", side="over", entity="A", line=0.5) == "win"
    ot = _res([(0, 0), (1, 1), (1, 1), (1, 0)], "OT")
    assert (
        grade.settle(market="game_pl", side="H", entity="", line=1.5, ot_rule="incl_ot_so", res=ot)
        == "win"
    )
    assert (
        grade.settle(
            market="game_total", side="over", line=4.5, entity="", ot_rule="incl_ot", res=ot
        )
        == "win"
    )
    live = GameResult(1, Date(2026, 10, 1), "A", "H", "LIVE")
    assert (
        grade.settle(
            market="game_ml", side="H", entity="", line=None, ot_rule="incl_ot_so", res=live
        )
        is None
    )
    assert (
        grade.pnl("win", -150) == pytest.approx(2 / 3)
        and grade.pnl("win", 120) == 1.2
        and grade.pnl("push", 120) == 0
    )


def test_write_once_and_grade_rows_with_clv(tmp_path: Path):
    sel = _sel(american=-110, consensus=0.52)
    p = stamp(
        sel,
        markets.Prob(0.56),
        ot_rule="incl_ot_so",
        thr=GateParams(),
        quote_age_minutes=1,
        goalie_confirmed=True,
        settlement_gate=None,
    )
    row = row_from(
        p,
        slate_date=Date(2026, 10, 1),
        home="H",
        away="A",
        home_goalie="g1",
        away_goalie="g2",
        goalie_status="confirmed/confirmed",
        lineup_source="team_rate",
        priced_at="t",
        card_tag="initial",
    )
    path = tmp_path / "predictions.json"
    _, wrote = write_once([row], path)
    assert wrote
    _, wrote2 = write_once([row], path)
    assert not wrote2
    alt, wrote3 = write_once([row], path, force=True)
    assert wrote3 and alt.name == "predictions.v2.json"
    loaded = load_rows(path)
    assert len(loaded) == 1 and loaded[0].is_buy and loaded[0].gates == []
    close = [
        _q("game_ml", "H", -140, 120, at="2026-10-01T23:00:00Z"),
        _q("game_ml", "A", 120, -140, at="2026-10-01T23:00:00Z"),
    ]
    res = {"A @ H": _res([(0, 1), (0, 1), (1, 0)])}
    graded = grade_rows(loaded, res, close, graded_at="g")
    r = graded[0]
    assert r.outcome == "win" and r.pnl == pytest.approx(100 / 110)
    assert r.close_american == -140 and r.clv is not None and r.clv > 0  # market moved our way
    sc = scorecard.scorecard(graded)
    d = sc.by_market["game_ml"].as_dict()
    assert d["n"] == 1 and d["wins"] == 1 and sc.probation()["game_ml"]["needed"] == 99
    assert "game_ml" in sc.render()


# -- pipeline + outputs -------------------------------------------------------------


def _inputs(code: str, rates: dict[str, float], status: str) -> TeamInputs:
    return TeamInputs(
        code,
        dict(rates),
        {k: 0.5 for k in rates},
        10,
        "team_rate [reported, not scored]",
        Starter(code, 1, f"{code} G", status, "test", None),
    )


def test_price_game_and_outputs(tmp_path: Path):
    from nhl_engine import outputs
    from nhl_engine.pipeline import SlateCard

    rows = [
        _q("game_ml", "H", -150, 130),
        _q("game_ml", "A", 130, -150),
        _q("game_pl", "H", 150, -180, line=-1.5),
        _q("game_pl", "A", -180, 150, line=1.5),
        _q("game_total", "over", -110, -110, line=6.5),
        _q("game_total", "under", -110, -110, line=6.5),
        _q("team_total", "over", -115, -105, line=2.5, entity="H"),
        _q("team_total", "over", -115, -105, line=2.5, entity="H", book="verified"),
        _q("p1_total", "over", -120, 100, line=1.5),
        _q("p1_ml3", "draw", 210, None),
        _q("sk_sog", "over", -110, -110, line=2.5, entity="Someone"),
    ]
    rules = BookRules([BookRule("verified", "team_total", "reg_only", "2026-09-30")])
    board = selections(rows)
    cfg = Config(sim=SimParams(draws=3000))
    card = price_game(
        board,
        home_in=_inputs("H", STRONG, "confirmed"),
        away_in=_inputs("A", WEAK, "projected"),
        league=LEAGUE,
        cfg=cfg,
        slate=Date(2026, 10, 1),
        rules=rules,
        calib=Calibrator(),
        now=NOW,
        priced_at="t",
        tag="initial",
        seed=1,
    )
    by = {(r.market, r.side, r.entity, r.book): r for r in card.rows}
    assert ("sk_sog", "over", "Someone", "dk") not in by  # props are Phase 3
    ml = by[("game_ml", "H", "", "dk")]
    assert ml.model_prob > 0.5 and "goalie_unconfirmed" in ml.gates  # away goalie only projected
    assert by[("game_total", "over", "", "dk")].model_prob + by[
        ("game_total", "under", "", "dk")
    ].model_prob == pytest.approx(1.0)
    tt = by[("team_total", "over", "H", "dk")]
    assert "settlement_unverified" in tt.gates and tt.ot_rule == "incl_ot_so"
    tt_v = by[("team_total", "over", "H", "verified")]
    assert "settlement_unverified" not in tt_v.gates and tt_v.ot_rule == "reg_only"
    assert "probation" in by[("p1_total", "over", "", "dk")].gates
    assert all(r.gates or r.pass_gate for r in card.rows)

    slate = SlateCard(Date(2026, 10, 1), 2026, "t", "initial", "v1", [card])
    paths = outputs.write_all(slate, tmp_path)
    txt = paths["txt"].read_text()
    assert "A @ H" in txt and "goalie_unconfirmed" in txt and "Period read" in txt
    assert "H G (confirmed" in txt
    assert paths["md"].read_text().startswith("# NHL brief")
    assert paths["xlsx"].exists()
    from openpyxl import load_workbook

    wb = load_workbook(paths["xlsx"])
    assert wb.sheetnames == ["Buys", "All", "Games"]
    assert wb["All"].max_row == len(card.rows) + 1
