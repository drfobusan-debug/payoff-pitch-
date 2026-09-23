"""The daily worksheet: scoring, weights, the ledger and its grading."""

from __future__ import annotations

import math
from datetime import date as Date
from pathlib import Path

from openpyxl import load_workbook

from mlb_engine.output import daily_worksheet as dw
from mlb_engine.output.daily_worksheet import (
    WEIGHTS,
    LedgerRow,
    Prices,
    Ranking,
    StarterTable,
    TeamLine,
    band_of,
    build_rows,
    decile_points,
    grade,
    implied,
    load_ledger,
    merge,
    rescale_batting,
    save_ledger,
    score,
    score_deciles,
    score_z,
    tally,
    write_workbook,
)
from mlb_engine.schemas import Game, Hand, Pitcher, Slate, TeamGameInfo, Venue

# --- scoring -----------------------------------------------------------------------


def test_score_top_bottom_middle_and_nan_last() -> None:
    vals = {"A": 1.0, "B": 2.0, "C": 3.0, "D": math.nan, "E": 5.0}
    rk = score(list(vals), [("m", 1, False, True)], lambda e, _l: vals[e], 1, lambda e: 0.0)
    assert rk.pts["A"]["m"] == 2  # best (lower is better)
    assert rk.pts["D"]["m"] == -2  # missing sorts worst
    assert rk.pts["B"]["m"] == rk.pts["C"]["m"] == rk.pts["E"]["m"] == 1
    assert rk.ranked[0] == "A" and rk.ranked[-1] == "D"
    assert math.isnan(rk.vals["D"]["m"])


def test_score_higher_is_better_and_percent_display() -> None:
    vals = {"A": 0.10, "B": 0.30, "C": 0.20}
    rk = score(list(vals), [("k", 1, True, False)], lambda e, _l: vals[e], 1, lambda e: 0.0)
    assert rk.pts["B"]["k"] == 2 and rk.pts["A"]["k"] == -2
    assert rk.vals["B"]["k"] == 30.0


def test_score_z_standardises_signs_clips_and_fills_nan_with_worst() -> None:
    vals = {"A": 1.0, "B": 2.0, "C": 3.0, "D": math.nan, "E": 100.0}
    rk = score_z(list(vals), [("m", 1, False, True)], lambda e, _l: vals[e], lambda e: 0.0)
    assert rk.pts["A"]["m"] >= rk.pts["B"]["m"] >= rk.pts["C"]["m"] > rk.pts["E"]["m"]  # lower is better
    assert rk.pts["E"]["m"] >= -dw.Z_CLIP and rk.pts["A"]["m"] <= dw.Z_CLIP
    assert rk.pts["D"]["m"] == rk.pts["E"]["m"]  # missing takes the table worst
    assert rk.ranked[0] == "A" and set(rk.ranked[-2:]) == {"D", "E"}
    assert rk.k == 0 and math.isnan(rk.vals["D"]["m"])


def test_score_z_higher_is_better_sums_metrics_and_handles_constant_column() -> None:
    vals = {"A": {"k": 0.10, "c": 5.0}, "B": {"k": 0.30, "c": 5.0}, "C": {"k": 0.20, "c": 5.0}}
    rk = score_z(list(vals), [("k", 1, True, False), ("c", 1, False, False)],
                 lambda e, label: vals[e][label], lambda e: 0.0)
    assert rk.pts["B"]["k"] > 0 > rk.pts["A"]["k"] and abs(rk.pts["C"]["k"]) < 1e-9
    assert all(rk.pts[e]["c"] == 0.0 for e in vals)  # zero spread -> no points
    assert rk.total("B") == round(rk.pts["B"]["k"] + rk.pts["B"]["c"], 1)
    assert rk.ranked == ["B", "C", "A"] and rk.vals["B"]["k"] == 30.0


def test_decile_points_ties_nan_and_direction() -> None:
    vals = {f"T{i}": float(i) for i in range(1, 31)}  # 30 teams, 1..30
    hi = decile_points(vals, lower=False)
    assert hi["T30"] == 10 and hi["T28"] == 10 and hi["T27"] == 9 and hi["T1"] == 1 and hi["T3"] == 1
    assert sorted(hi.values()) == sorted([d for d in range(1, 11) for _ in range(3)])
    lo = decile_points(vals, lower=True)
    assert lo["T1"] == 10 and lo["T30"] == 1 and all(lo[t] == 11 - hi[t] for t in vals)
    # NaN is the bottom decile; ties share one decile.
    vals["T30"] = math.nan
    vals["T28"] = vals["T29"]
    hi = decile_points(vals, lower=False)
    assert hi["T30"] == 1 and hi["T28"] == hi["T29"] == 10
    assert decile_points({"A": 1.0, "B": 1.0}, lower=False) == {"A": 10, "B": 10}


def test_score_deciles_is_uncompressed_and_ranked_by_sum() -> None:
    ents = [f"T{i}" for i in range(1, 31)]
    vals = {e: {"a": float(i), "b": 31.0 - i} for i, e in enumerate(ents, 1)}
    rk = score_deciles(ents, [("a", 1, False, False), ("b", 1, False, True)], lambda e, m: vals[e][m], lambda e: -vals[e]["a"])
    assert rk.deciles and rk.pts["T30"] == {"a": 10, "b": 10} and rk.pts["T1"] == {"a": 1, "b": 1}
    assert len({rk.raw_total(e) for e in ents}) == 10  # ten distinct sums, not three
    assert rk.ranked[0] == "T30" and rk.total("T30") == rescale_batting(20)


def test_rescale_batting_keeps_the_old_mean_and_range() -> None:
    assert rescale_batting(60) + rescale_batting(61) == 2 * dw.BAT_OLD_MEAN
    assert rescale_batting(11) == -6.05 and rescale_batting(110) == 23.65  # ~ old 5..22 plus the tails


# --- weights -----------------------------------------------------------------------


def test_weighted_total_uses_run_share_weights_and_skips_missing() -> None:
    t = TeamLine("SD", "X", "R", 10.5, 6, 8)
    assert t.raw_total() == 24.5
    assert t.weighted() == 10.5 + 6 * 1.5 + 8 * 3
    assert WEIGHTS == {"bat vs hand": 1.0, "bp": 1.5, "sp": 3.0}
    missing = TeamLine("SD", "", "", 10.5, 6, None)
    assert missing.weighted() == 10.5 + 9  # nothing counted for the absent starter, not zero
    assert missing.label() == "sd (TBD)"


# --- slate rows --------------------------------------------------------------------


def _rank(pts: dict[str, int]) -> Ranking:
    return Ranking({t: {"m": v} for t, v in pts.items()}, {t: {"m": 0.0} for t in pts},
                   {t: set() for t in pts}, sorted(pts, key=lambda t: -pts[t]), 3)


def _dec(sums: dict[str, int]) -> Ranking:
    """A decile batting table whose single metric carries the 11-metric sum."""
    return Ranking({t: {"m": v} for t, v in sums.items()}, {t: {"m": 0.0} for t in sums},
                   {t: set() for t in sums}, sorted(sums, key=lambda t: -sums[t]), 0, deciles=True)


def _team(abbrev: str, home: bool, sp: str | None, hand: Hand = Hand.R) -> TeamGameInfo:
    pp = Pitcher(mlbam_id=hash(sp) % 10000, name=sp, throws=hand) if sp else None
    return TeamGameInfo(team_id=1, name=abbrev, abbrev=abbrev, is_home=home, probable_pitcher=pp)


def _game(pk: int, away: TeamGameInfo, home: TeamGameInfo) -> Game:
    return Game(game_pk=pk, game_date=Date(2026, 9, 14), status="S",
                venue=Venue(venue_id=1, name="v"), home=home, away=away)


def _fixture() -> tuple[Slate, dict[str, Ranking], Ranking, StarterTable, dict[tuple[str, str], Prices]]:
    bats = {
        "Overall": _dec({"SD": 90, "COL": 30, "NYY": 70, "MIN": 70}),
        "vs LHP": _dec({"SD": 95, "COL": 25, "NYY": 60, "MIN": 60}),
        "vs RHP": _dec({"SD": 90, "COL": 35, "NYY": 65, "MIN": 70}),
    }
    pens = _rank({"SD": 14, "COL": -10, "NYY": 10, "MIN": 10})
    sps = StarterTable(_rank({"1": 15, "2": -7, "3": 9}), {1: "Ace", 2: "Scrub", 3: "Arm"},
                       {1: "SD", 2: "COL", 3: "NYY"}, {1: 100.0, 2: 50.0, 3: 90.0},
                       {"ace": 1, "scrub": 2, "arm": 3})
    slate = Slate(slate_date=Date(2026, 9, 14), games=[
        _game(1, _team("SD", False, "Ace", Hand.L), _team("COL", True, "Scrub")),
        _game(2, _team("NYY", False, "Arm"), _team("MIN", True, "Unknown Guy")),
    ])
    prices = {(g.matchup(), t.abbrev): Prices() for g in slate.games for t in (g.away, g.home)}
    prices[("SD @ COL", "SD")] = Prices(ml=-175, ml_book="lowvig", rl_line=-1.5, rl=+105, rl_book="dk",
                                        dk_ml_handle=70, dk_ml_bets=60)
    prices[("SD @ COL", "COL")] = Prices(ml=+160, ml_book="dk", rl_line=1.5, rl=-125, rl_book="dk")
    return slate, bats, pens, sps, prices


def test_build_rows_reads_the_split_by_opposing_hand_and_flags_missing_starter() -> None:
    slate, bats, pens, sps, prices = _fixture()
    rows = build_rows(slate, bats, pens, sps, prices)
    (_, a, h, r1), (_, a2, h2, r2) = rows
    # COL faces the lefty Ace -> vs LHP; SD faces the righty Scrub -> vs RHP; both rescaled from the decile sum.
    assert h.bat_vs_hand == rescale_batting(25) == -1.85 and a.bat_vs_hand == rescale_batting(90) == 17.65
    assert a.sp == 15 and h.sp == -7 and r1.complete
    assert a.weighted() == round(17.65 + 14 * 1.5 + 15 * 3, 1) and r1.away_bat == 17.65
    assert r1.gap == round(a.weighted() - h.weighted(), 1) and r1.fav == "SD"
    assert r1.away_ml == -175 and r1.home_ml == 160 and r1.away_rl_line == -1.5
    assert r1.fav_implied is not None and 0.6 < r1.fav_implied < 0.65
    assert r1.dk_ml_handle_away == 70 and r1.dk_ml_handle_home is None
    prices[("SD @ COL", "COL")].circa_rl_handle = 33
    prices[("SD @ COL", "COL")].circa_rl_bets = 44
    (_, _, _, r1b), _ = build_rows(slate, bats, pens, sps, prices)
    assert r1b.circa_rl_handle_home == 33 and r1b.circa_rl_bets_home == 44
    # MIN's probable is not in the starter table: no points, row marked incomplete.
    assert h2.sp is None and a2.sp == 9 and not r2.complete
    assert r2.home_sp_pts is None


def test_implied_strips_the_vig() -> None:
    assert implied(-110, -110) == 0.5
    p = implied(-175, 160)
    assert p is not None and abs(p - 0.6229) < 0.002
    assert implied(None, 160) is None


# --- ledger ------------------------------------------------------------------------


def _row(pk: int, gap: float, fav: str = "SD", result: str = "", complete: bool = True, day: str = "2026-09-14") -> LedgerRow:
    return LedgerRow(day, "SD @ COL", pk, "SD", "COL", "Ace", "Scrub", 9, 14, 15, -8, -10, -7,
                     80.0, 80.0 - gap, gap, fav, complete, away_ml=-175, home_ml=160,
                     away_rl_line=-1.5, fav_implied=0.62, result=result)


def test_ledger_round_trips_with_none_and_bool(tmp_path: Path) -> None:
    path = tmp_path / "l.csv"
    rows = [_row(1, 30.5), _row(2, -4.0, fav="COL", complete=False)]
    rows[1].away_ml = None
    save_ledger(path, rows)
    back = load_ledger(path)
    assert back == rows
    assert back[1].away_ml is None and back[1].complete is False and back[0].complete is True


def test_old_ledger_rows_with_bat6_columns_still_load_and_stay_graded(tmp_path: Path) -> None:
    path = tmp_path / "old.csv"
    path.write_text(
        "date,game,game_pk,away,home,away_sp,home_sp,away_bat,away_bat6,away_bp,away_sp_pts,home_bat,home_bat6,"
        "home_bp,home_sp_pts,away_w,home_w,gap,fav,complete,away_ml,home_ml,away_rl_line,away_rl,home_rl,"
        "fav_implied,weights,away_runs,home_runs,result,rl_result\n"
        "2026-09-01,SD @ COL,7,SD,COL,Ace,Scrub,9,8,14,15,-8,-6,-10,-7,83.0,-52.0,135.0,SD,True,-175,160,-1.5,105,"
        "-125,0.62,w1-1-1.5-3,5,3,fav,away\n"
    )
    (r,) = load_ledger(path)
    assert r.game_pk == 7 and r.away_bat == 9 and r.home_bat == -8 and r.complete
    assert r.graded and r.result == "fav" and r.rl_result == "away" and r.away_runs == 5
    assert r.weights == "w1-1-1.5-3" != dw.WEIGHTS_VERSION
    assert not hasattr(r, "away_bat6")
    # The old weights tag keeps the row out of the band tally; a fresh row is counted.
    assert all(t.n == 0 for t in tally([r]))
    assert {t.band: t.n for t in tally([r, _row(8, 30.5, result="fav")])}["all"] == 1
    # Re-saving writes the current header and the row round-trips.
    save_ledger(path, [r])
    assert "bat6" not in path.read_text().splitlines()[0] and load_ledger(path) == [r]


def test_merge_rewrites_ungraded_rows_and_keeps_graded_ones() -> None:
    graded = _row(1, 30.5, result="fav")
    ungraded = _row(2, 5.0)
    fresh = [_row(1, 99.0), _row(2, 7.0), _row(3, 1.0)]
    out = merge([graded, ungraded], fresh)
    by = {r.game_pk: r for r in out}
    assert by[1].gap == 30.5 and by[1].result == "fav"  # never touched
    assert by[2].gap == 7.0  # re-run replaces the ungraded price/gap
    assert by[3].gap == 1.0 and len(out) == 3


def test_grade_fills_finals_favourite_and_runline() -> None:
    rows = [_row(1, 30.5), _row(2, -10.0, fav="COL"), _row(3, 2.0)]
    n = grade(rows, {1: ("SD @ COL", 5, 3), 2: ("SD @ COL", 6, 1)})
    assert n == 2
    assert rows[0].result == "fav" and rows[0].rl_result == "away"  # 5-1.5 > 3
    assert rows[1].result == "dog" and rows[1].rl_result == "away"
    assert rows[2].result == "" and rows[2].away_runs is None
    # A second pass never regrades.
    assert grade(rows, {1: ("SD @ COL", 0, 9)}) == 0 and rows[0].result == "fav"
    # A level final is neither a hit nor a miss and stays out of the tally.
    assert grade(rows, {3: ("SD @ COL", 4, 4)}) == 1 and rows[2].result == "tie"
    assert all(t.n == 0 for t in dw.tally([rows[2]]))


def test_starter_override_replaces_the_listed_probable_on_that_day_only() -> None:
    cws = _team("CWS", False, "Sean Newcomb", Hand.L)
    cle = _team("CLE", True, "Gavin Williams")
    slate = dw.apply_starter_overrides(Slate(slate_date=Date(2026, 9, 14), games=[_game(1, cws, cle)]))
    pp = slate.games[0].away.probable_pitcher
    assert pp is not None and pp.name == "Sean Burke" and pp.throws == Hand.R
    later = _team("CWS", False, "Sean Newcomb", Hand.L)
    slate = dw.apply_starter_overrides(Slate(slate_date=Date(2026, 9, 15), games=[_game(2, later, cle)]))
    pp = slate.games[0].away.probable_pitcher
    assert pp is not None and pp.name == "Sean Newcomb"


def test_tally_by_gap_band_excludes_incomplete_and_compares_to_market() -> None:
    rows = [
        _row(1, 30.5, result="fav"), _row(2, 22.0, result="dog"),
        _row(3, 3.0, result="fav"), _row(4, 40.0, result="fav", complete=False), _row(5, 12.0),
    ]
    t = {b.band: b for b in tally(rows)}
    assert t["20-34"].n == 2 and t["20-34"].fav_wins == 1 and t["20-34"].win_rate == 0.5
    assert t["0-9"].n == 1 and t["35+"].n == 0 and t["10-19"].n == 0  # incomplete and ungraded skipped
    assert t["all"].n == 3 and t["all"].market_rate is not None and abs(t["all"].market_rate - 0.62) < 1e-9
    assert t["all"].edge is not None and abs(t["all"].edge - (2 / 3 - 0.62)) < 1e-9
    assert band_of(-9.9) == "0-9" and band_of(35) == "35+"


# --- workbook ----------------------------------------------------------------------


def test_workbook_layout_and_ranked_block(tmp_path: Path) -> None:
    slate, bats, pens, sps, prices = _fixture()
    rows = build_rows(slate, bats, pens, sps, prices)
    ledger = [r for _, _, _, r in rows] + [_row(9, 25.0, result="fav", day="2026-09-13")]
    out = write_workbook(tmp_path / "w.xlsx", Date(2026, 9, 14), Date(2026, 9, 13), rows, prices,
                         pens, bats, sps, ledger, Date(2026, 9, 13))
    wb = load_workbook(out)
    assert wb.sheetnames == ["Matchups", "Audit", "Ledger", "Bullpens", "Bat Overall", "Bat vs LHP",
                             "Bat vs RHP", "Starters"]
    ws = wb["Matchups"]
    hdr = [c.value for c in ws[1]]
    assert hdr[:7] == ["game", "team", "bat vs hand", "bp", "sp", "raw TOTAL", "wTOTAL"]
    assert "bat 6+" not in hdr and "sp k-bb%" not in hdr and "BES rp" not in hdr
    assert ws["A2"].value == "SD @ COL" and ws["B2"].value == "sd (Ace, L)"
    assert ws["C2"].value == 17.65 and ws["G2"].value == round(17.65 + 14 * 1.5 + 15 * 3, 1)  # rescaled bat in wTOTAL
    assert ws["H2"].value == "-175 lowvig" and ws["I2"].value == "-1.5 +105 dk" and ws["L2"].value == "70/60"
    assert ws["B4"].value == "away - home"
    assert ws["G4"].value == round(ws["G2"].value - ws["G3"].value, 1)
    legend = " ".join(str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str))
    assert "decile sum rescaled" in legend and "bat 6+" not in legend and "Innings 6+" not in legend
    assert "wTOTAL = bat x1.0 + bp x1.5 + sp x3.0" in legend
    bo = wb["Bat Overall"]
    assert [c.value for c in bo[1]][:4] == ["Rk", "Team", "PTS", "dec sum"]
    assert bo["B2"].value == "SD" and bo["C2"].value == 17.65 and bo["D2"].value == 90
    notes = " ".join(str(c.value) for row in bo.iter_rows() for c in row if isinstance(c.value, str))
    assert "decile among the 30" in notes and "d10 = best" in notes and "best 3 +2" not in notes
    assert ws["B8"].value == "away - home *"  # MIN's starter missing
    # Ranked block: biggest |gap| first.
    c0 = hdr.index("rank") + 1
    gaps = [ws.cell(row=r, column=c0 + 6).value for r in (2, 3)]
    assert gaps == sorted(gaps, reverse=True)
    assert ws.cell(row=2, column=c0 + 1).value == "SD @ COL"
    assert ws.cell(row=3, column=c0 + 1).value == "NYY @ MIN *"
    wa = wb["Audit"]
    assert [c.value for c in wa[1]][:3] == ["gap band", "games", "fav W"]
    assert wb["Ledger"]["A2"].value == "2026-09-14"


def test_pdf_pages_are_the_worksheet_first_then_the_tables_in_the_same_fills(tmp_path: Path) -> None:
    slate, bats, pens, sps, prices = _fixture()
    rows = build_rows(slate, bats, pens, sps, prices)
    ledger = [r for _, _, _, r in rows]
    out = write_workbook(tmp_path / "worksheet_2026-09-14.xlsx", Date(2026, 9, 14), Date(2026, 9, 13),
                         rows, prices, pens, bats, sps, ledger, None)
    page = dw.worksheet_html(out, Date(2026, 9, 14))
    titles = [t for _, t in dw.PDF_SHEETS]
    positions = [page.index(t) for t in titles]
    assert positions == sorted(positions) and positions[0] < page.index("Matchups ranked by weighted gap")
    assert "Audit - 2026" not in page and "Ledger - 2026" not in page
    assert "SD @ COL" in page and "background:#" in page
    pdf = dw.write_pdf(out, Date(2026, 9, 14))
    assert pdf == tmp_path / "worksheet_2026-09-14.pdf" and pdf.read_bytes()[:4] == b"%PDF"


def test_summary_text_reports_yesterday_and_bands() -> None:
    ledger = [_row(1, 30.5, result="fav", day="2026-09-13"), _row(2, 3.0, result="dog", day="2026-09-13")]
    text = dw.summary_text(ledger, Date(2026, 9, 13))
    assert text.startswith("Worksheet 2026-09-13: favoured side 1-1")
    assert "20-34: 1-0 (100% vs mkt 62%)" in text
    assert dw.summary_text([], None) == ""


def test_summary_text_lists_today_by_gap_with_incomplete_starred() -> None:
    ledger = [_row(1, 3.0), _row(2, -30.5, fav="COL", complete=False)]
    text = dw.summary_text(ledger, None, Date(2026, 9, 14))
    assert text == "Worksheet 2026-09-14 by weighted gap: SD @ COL COL 30.5*; SD @ COL SD 3.0  (* starter missing)"
