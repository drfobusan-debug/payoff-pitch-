"""The nightly audit fits Gmail's message cap: article first, ledger in parts."""

from __future__ import annotations

import base64
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from mlb_engine.audit.ledger import LedgerEntry
from mlb_engine.output import audit_insight
from mlb_engine.output.email import MAX_EMAIL_BYTES, attachment_budget
from mlb_engine.output.excel import (
    ledger_date_chunks,
    write_ledger_parts,
    write_ledger_workbook,
)


def _entry(day: str, i: int) -> LedgerEntry:
    return LedgerEntry(
        date=day,
        matchup=f"A @ B {i}",
        category="Batter",
        market="batter_hits",
        selection=f"Player {i} Over 0.5",
        line=0.5,
        book="dk",
        odds=-110.0,
        tier="PASS",
        model_prob=0.5,
        ev=None,
        result="win",
        pnl=0.909,
    )


def _ledger() -> list[LedgerEntry]:
    days = [f"2026-09-{d:02d}" for d in range(1, 21)]
    return [_entry(d, i) for d in days for i in range(40)]


def test_budget_leaves_the_encoded_attachment_under_the_cap() -> None:
    raw = attachment_budget()
    assert len(base64.encodebytes(b"x" * raw)) <= MAX_EMAIL_BYTES


def test_chunks_are_whole_dates_in_order_and_lose_nothing() -> None:
    entries = _ledger()
    chunks = ledger_date_chunks(entries, 3)
    assert len(chunks) == 3
    assert sum(len(c) for c in chunks) == len(entries)
    spans = [(c[0].date, c[-1].date) for c in chunks]
    assert all(a[1] < b[0] for a, b in zip(spans, spans[1:], strict=False))
    assert max(len(c) for c in chunks) - min(len(c) for c in chunks) <= 40


def test_a_ledger_that_fits_goes_whole(tmp_path: Path) -> None:
    full = write_ledger_workbook(_ledger(), [], [], tmp_path / "ledger.xlsx")
    assert write_ledger_parts(_ledger(), [], [], full, max_bytes=10**9) == [full]


def test_an_oversize_ledger_splits_into_parts_that_each_fit(tmp_path: Path) -> None:
    entries = _ledger()
    full = write_ledger_workbook(entries, [], [], tmp_path / "ledger.xlsx")
    (tmp_path / "ledger_part1of9_old_old.xlsx").write_bytes(b"stale")
    budget = full.stat().st_size // 2
    parts = write_ledger_parts(entries, [], [], full, max_bytes=budget)
    assert len(parts) >= 2
    assert not (tmp_path / "ledger_part1of9_old_old.xlsx").exists()
    rows = 0
    for p in parts:
        assert p.stat().st_size <= budget
        assert p.name.startswith("ledger_part")
        wb = load_workbook(p, read_only=True)
        assert {"Overall", "Daily", "Bets"} <= set(wb.sheetnames)
        rows += wb["Bets"].max_row - 1
    assert rows == len(entries)
    assert full.exists()


def test_article_and_each_ledger_part_go_in_their_own_email(monkeypatch) -> None:
    sent: list[tuple[str, list[str]]] = []

    def fake_send(cfg, *, subject, html_body, text_body, to, attachments):
        if "part 1 of 2" in subject:
            raise OSError("552 message too big")
        sent.append((subject, [name for name, _ in attachments]))
        return "a@b.com"

    monkeypatch.setattr(audit_insight, "send_card_email", fake_send)
    n = audit_insight.send_audit_emails(
        object(),
        date(2026, 9, 29),
        [("PayoffPitch_Audit_2026-09-29.pdf", b"%PDF"), ("PayoffPitch_Audit_2026-09-29.mp3", b"ID3")],
        [("ledger_part1of2.xlsx", b"PK"), ("ledger_part2of2.xlsx", b"PK")],
        to=None,
    )
    assert n == 2
    assert sent[0][1] == ["PayoffPitch_Audit_2026-09-29.pdf", "PayoffPitch_Audit_2026-09-29.mp3"]
    assert sent[1] == ("Payoff Pitch — Nightly Audit Ledger (2026-09-29) — part 2 of 2", ["ledger_part2of2.xlsx"])
