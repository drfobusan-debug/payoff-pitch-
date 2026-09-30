"""Card renderers: terminal card, markdown brief, Excel workbook (outputs plan §2).

Every row on the card carries its gate stamps; buys are the ``pass_gate`` rows
with a non-Pass tier. Period markets and any market outside
``GateParams.live_markets`` show as "Period read"/"probation" -- priced and
graded, never a buy.
"""

from __future__ import annotations

from collections import Counter
from datetime import date as Date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from nhl_engine.audit.ledger import LedgerRow
from nhl_engine.pipeline import GameCard, SlateCard

_ORDER = [
    "game_ml",
    "game_ml3",
    "game_pl",
    "game_pl_alt",
    "game_total",
    "game_total_alt",
    "team_total",
    "team_total_alt",
]


def _fmt_line(r: LedgerRow) -> str:
    if r.line is None:
        return ""
    return f"{r.line:+.1f}" if "pl" in r.market else f"{r.line:.1f}"


def _sel(r: LedgerRow) -> str:
    ent = f" {r.entity}" if r.entity and r.entity not in (r.side,) else ""
    return f"{r.side}{ent} {_fmt_line(r)}".strip()


def _american(x: float) -> str:
    return f"{x:+.0f}"


def _game_header(g: GameCard) -> list[str]:
    s = g.summary
    hi, ai = g.home_in, g.away_in
    return [
        f"{g.matchup}   {ai.code} {ai.games}gp  {hi.code} {hi.games}gp",
        f"  home win {s['home_win']:.3f}  reg {s['home_reg_win']:.3f}/{s['draw_reg']:.3f} draw"
        f"  exp {s['exp_away']:.2f}-{s['exp_home']:.2f} total {s['exp_total']:.2f}"
        f"  P1/P2/P3 {s['p1_exp_total']:.2f}/{s['p2_exp_total']:.2f}/{s['p3_exp_total']:.2f}",
        f"  goalies {ai.starter.name} ({ai.starter.status}, {ai.starter.gsax60:+.2f} GSAx/60)"
        f" vs {hi.starter.name} ({hi.starter.status}, {hi.starter.gsax60:+.2f})",
        f"  5v5 xGF/xGA {ai.rates.get('xgf60_5v5', 0):.2f}/{ai.rates.get('xga60_5v5', 0):.2f}"
        f" | {hi.rates.get('xgf60_5v5', 0):.2f}/{hi.rates.get('xga60_5v5', 0):.2f}"
        f"   PP {ai.rates.get('pp_xgf60', 0):.1f}/{hi.rates.get('pp_xgf60', 0):.1f}"
        f" PK {ai.rates.get('pk_xga60', 0):.1f}/{hi.rates.get('pk_xga60', 0):.1f}",
        f"  lineup: {ai.lineup_source} | {hi.lineup_source}",
    ]


def _row_line(r: LedgerRow) -> str:
    gates = "pass_gate" if r.pass_gate else ",".join(r.gates)
    return (
        f"    {r.market:<15s} {_sel(r):<22s} {_american(r.american):>6s} {r.book:<14s}"
        f" model {r.model_prob:.3f} mkt {r.consensus:.3f} edge {r.edge:+.3f} EV {r.ev:+.3f}"
        f"  {r.tier:<12s} {gates}"
    )


def render_card(card: SlateCard) -> str:
    lines = [
        f"NHL card {card.slate_date}  [{card.tag}]  priced {card.priced_at}  prior {card.prior_version}",
        "",
    ]
    buys = [r for r in card.rows if r.is_buy]
    lines.append(f"BUYS ({len(buys)})")
    for r in sorted(buys, key=lambda r: -r.edge):
        lines.append(f"  {r.matchup:<10s}" + _row_line(r).lstrip())
    if not buys:
        lines.append("  none")
    lines.append("")
    for g in card.games:
        lines.extend(_game_header(g))
        game_rows = sorted(
            [r for r in g.rows if not r.market.startswith("p")],
            key=lambda r: (
                _ORDER.index(r.market) if r.market in _ORDER else 99,
                r.line or 0,
                r.side,
            ),
        )
        for r in game_rows:
            lines.append(_row_line(r))
        period = [r for r in g.rows if r.market.startswith("p")]
        if period:
            lines.append(
                f"    Period read ({len(period)} rows, probation): "
                + ", ".join(
                    f"{r.market} {_sel(r)} {r.edge:+.3f}"
                    for r in sorted(period, key=lambda r: -abs(r.edge))[:6]
                )
            )
        lines.append("")
    if card.unpriced:
        lines.append("UNPRICED: " + "; ".join(card.unpriced))
    gate_counts = Counter(g for r in card.rows for g in r.gates)
    lines.append("gates: " + ", ".join(f"{k} {v}" for k, v in gate_counts.most_common()))
    return "\n".join(lines) + "\n"


def render_brief(card: SlateCard) -> str:
    buys = [r for r in card.rows if r.is_buy]
    out = [f"# NHL brief {card.slate_date} ({card.tag})", ""]
    out.append(
        f"{len(card.games)} games priced, {len(card.rows)} selections read from one joint sim per game, {len(buys)} buys."
    )
    out.append("")
    if buys:
        out.append("| game | market | selection | price | book | model | market | edge | tier |")
        out.append("|---|---|---|---|---|---|---|---|---|")
        for r in sorted(buys, key=lambda r: -r.edge):
            out.append(
                f"| {r.matchup} | {r.market} | {_sel(r)} | {_american(r.american)} | {r.book} |"
                f" {r.model_prob:.3f} | {r.consensus:.3f} | {r.edge:+.3f} | {r.tier} |"
            )
        out.append("")
    out.append("## Games")
    for g in card.games:
        s = g.summary
        out.append(
            f"- **{g.matchup}** home {s['home_win']:.1%}, total {s['exp_total']:.2f}; "
            f"{g.away_in.starter.name} ({g.away_in.starter.status}) vs {g.home_in.starter.name} ({g.home_in.starter.status})"
        )
    if card.unpriced:
        out.append("")
        out.append("Unpriced: " + "; ".join(card.unpriced))
    out.append("")
    out.append(
        "Period markets and alternates are priced for the ledger only (probation) until ~100 graded rows per market."
    )
    return "\n".join(out) + "\n"


_COLS = [
    "matchup",
    "market",
    "side",
    "entity",
    "line",
    "ot_rule",
    "book",
    "american",
    "books",
    "consensus",
    "model_prob",
    "push_prob",
    "edge",
    "ev",
    "tier",
    "pass_gate",
    "gates",
    "kelly",
    "away_goalie",
    "home_goalie",
    "goalie_status",
    "lineup_source",
    "priced_at",
    "card_tag",
]


def write_excel(card: SlateCard, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Buys"
    ws.append(_COLS)
    for r in sorted((r for r in card.rows if r.is_buy), key=lambda r: -r.edge):
        ws.append(_cells(r))
    all_ws = wb.create_sheet("All")
    all_ws.append(_COLS)
    for r in card.rows:
        all_ws.append(_cells(r))
    gm = wb.create_sheet("Games")
    keys = [
        "home_win",
        "home_reg_win",
        "draw_reg",
        "went_so",
        "exp_away",
        "exp_home",
        "exp_total",
        "sd_total",
        "p1_exp_total",
        "p2_exp_total",
        "p3_exp_total",
    ]
    gm.append(
        ["matchup", "away_goalie", "away_status", "home_goalie", "home_status", "lineup"] + keys
    )
    for g in card.games:
        gm.append(
            [
                g.matchup,
                g.away_in.starter.name,
                g.away_in.starter.status,
                g.home_in.starter.name,
                g.home_in.starter.status,
                f"{g.away_in.lineup_source} | {g.home_in.lineup_source}",
            ]
            + [round(g.summary[k], 4) for k in keys]
        )
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for i, _ in enumerate(sheet[1], start=1):
            sheet.column_dimensions[get_column_letter(i)].width = 14
        sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _cells(r: LedgerRow) -> list[object]:
    return [
        r.matchup,
        r.market,
        r.side,
        r.entity,
        r.line,
        r.ot_rule,
        r.book,
        r.american,
        r.books,
        round(r.consensus, 4),
        round(r.model_prob, 4),
        round(r.push_prob, 4),
        round(r.edge, 4),
        round(r.ev, 4),
        r.tier,
        r.pass_gate,
        ",".join(r.gates) or "pass_gate",
        round(r.kelly, 4),
        r.away_goalie,
        r.home_goalie,
        r.goalie_status,
        r.lineup_source,
        r.priced_at,
        r.card_tag,
    ]


def card_paths(out_dir: Path, slate: Date, tag: str) -> dict[str, Path]:
    stem = f"card_{slate.isoformat()}_{tag}"
    return {
        "txt": out_dir / f"{stem}.txt",
        "md": out_dir / f"brief_{slate.isoformat()}_{tag}.md",
        "xlsx": out_dir / f"{stem}.xlsx",
    }


def write_all(card: SlateCard, out_dir: Path) -> dict[str, Path]:
    paths = card_paths(out_dir, card.slate_date, card.tag)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths["txt"].write_text(render_card(card), encoding="utf-8")
    paths["md"].write_text(render_brief(card), encoding="utf-8")
    write_excel(card, paths["xlsx"])
    return paths


__all__ = ["card_paths", "render_brief", "render_card", "write_all", "write_excel"]
