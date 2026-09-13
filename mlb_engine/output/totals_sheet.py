"""The hand-written totals sheet, built for the day and written as a workbook.

Every column is a signed point score: positive leans Over, negative leans Under,
and the row's sum is the read. The columns are the hand-written ones -- offense
on wRC+/wOBA (vs the opposing starter's hand) and barrel rate; each arm on
SIERA, xERA, CSW% and K-BB%; DraftKings and Circa total handle-minus-bets signed
by the side the money is on; temperature, wind (signed by direction) and
humidity at first pitch with a roof scored -1; park factor; BaseRuns per game;
bullpen fatigue off the last two days' boxscores; and the home-plate umpire's
runs per game against the league.

Nothing here touches a price. The sheet is research beside the card: measured
against the close it explains the posted number rather than beating it (see the
Legend sheet), so it is delivered as a reading aid, not as a bet.
"""

from __future__ import annotations

import json
import logging
import math
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date as Date
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from mlb_engine.audit.ledger import load_ledger
from mlb_engine.config import Config
from mlb_engine.data import http
from mlb_engine.data.mlb_statsapi import BASE as STATSAPI
from mlb_engine.data.mlb_statsapi import MLBStatsClient
from mlb_engine.data.parks import Park, get_park
from mlb_engine.data.vsin import TotalSplit, VSINClient
from mlb_engine.filters.weather import WeatherConditions, WeatherProvider
from mlb_engine.output.totals_audit import (
    BANDS,
    OverCurve,
    engine_at,
    engine_ledger_path,
    engine_median,
    first_line,
    over_curves,
    record_sheet,
    sheet_bands,
)
from mlb_engine.schemas import Slate, TeamGameInfo

log = logging.getLogger(__name__)

_FG_URL = (
    "https://www.fangraphs.com/api/leaders/major-league/data?age=&pos=all&stats={stats}"
    "&lg=all&qual=0&season={season}&season1={season}&month={month}&hand=&team={team}"
    "&pageitems=3000&pagenum=1&ind=0&rost=0&players=0&type=8&sortdir=default&sortstat=WAR"
)
_FG_HEADERS = {"User-Agent": "curl/8.5.0", "Accept": "*/*"}
# FanGraphs' team codes where they differ from StatsAPI's.
_FG_ABBR = {"AZ": "ARI", "CWS": "CHW", "KC": "KCR", "SD": "SDP", "SF": "SFG", "TB": "TBR", "WSH": "WSN"}
_FG_MONTH_VS_L, _FG_MONTH_VS_R = 13, 14

SEASON_START = (3, 20)
FATIGUE_CAP = 4
PEN_PITCHES_TIRED = 120
UMP_MIN_GAMES = 10
UMP_BAND_RUNS = 1.0


# --- the bands, centred on the league ------------------------------------------
# Every scale has a 0 window around the league median (roughly its middle half),
# +/-1 for the next quartile out, +/-2 and +/-3 beyond. A league-average arm or
# lineup therefore adds nothing, and +5 and -5 are the same strength of lean in
# opposite directions. Windows come from the 2026 FanGraphs distributions: team
# wRC+ 95-103 / wOBA .311-.319 / Barrel% 7.3-8.1 between the quartiles; starter
# SIERA 3.84-4.57, xERA 3.75-4.88, CSW% 25.4-28.0, K-BB% 10.0-16.1; team pens
# SIERA 3.71-4.02, CSW% 26.6-27.9, K-BB% 11.3-14.1.


def wrc_pts(v: float) -> int:
    return 3 if v > 130 else 2 if v > 115 else 1 if v > 105 else 0 if v >= 95 else -1 if v >= 85 else -2 if v >= 70 else -3


def woba_pts(v: float) -> int:
    return (
        3 if v > 0.355 else 2 if v > 0.340 else 1 if v > 0.325 else 0 if v >= 0.305
        else -1 if v >= 0.290 else -2 if v >= 0.275 else -3
    )


def barrel_pts(v: float) -> int:
    """``v`` in percent."""
    return 2 if v > 10 else 1 if v > 8.5 else 0 if v >= 6.5 else -1 if v >= 5 else -2


def siera_pts(v: float) -> int:
    return (
        -3 if v < 2.75 else -2 if v < 3.25 else -1 if v < 3.75 else 0 if v <= 4.35
        else 1 if v <= 4.75 else 2 if v <= 5.25 else 3
    )


def csw_pts(v: float) -> int:
    """``v`` in percent."""
    return -2 if v > 31 else -1 if v > 29 else 0 if v >= 25 else 1 if v >= 23 else 2


def xera_pts(v: float) -> int:
    return -2 if v < 3.2 else -1 if v < 3.7 else 0 if v <= 4.5 else 1 if v <= 5.1 else 2


def kbb_pts(v: float) -> int:
    """``v`` in percent."""
    return -3 if v > 25 else -2 if v > 19 else -1 if v > 15 else 0 if v >= 10 else 1 if v >= 7 else 2 if v >= 4 else 3


def temp_pts(t: float) -> int:
    return 2 if t > 90 else 1 if t > 80 else 0 if t >= 60 else -1 if t >= 50 else -2


def wind_pts(w: float) -> int:
    """Speed points for a wind with a direction that matters; cross winds score 0."""
    return 3 if w > 15 else 2 if w >= 10 else 1 if w >= 5 else 0


def humidity_pts(h: float) -> int:
    return 1 if h > 65 else 0 if h >= 35 else -1


def park_pts(pf: float) -> int:
    return 1 if pf > 102 else -1 if pf < 98 else 0


def baseruns_pts(v: float) -> int:
    """Provisional (not hand-written): league mean ~4.47 +/- 0.2."""
    return 1 if v >= 4.65 else -1 if v <= 4.25 else 0


def book_pts(sp: TotalSplit | None) -> int:
    """Handle minus bets on the side the money is on, signed toward that side.

    >=20 -> 2, 11-19 -> 1, else 0; a 100/100 split is one ticket, not a market,
    and scores 0. Over is positive, Under negative; the two sides are one
    signal, never two.
    """
    if sp is None:
        return 0
    best = 0
    for sign, side in ((1, sp.over), (-1, sp.under)):
        if side.handle_pct is None or side.bets_pct is None:
            continue
        if side.handle_pct == 100 and side.bets_pct == 100:
            continue
        d = side.handle_pct - side.bets_pct
        pts = 2 if d >= 20 else 1 if d >= 11 else 0
        if pts > abs(best):
            best = sign * pts
    return best


def weather_pts(park: Park | None, cond: WeatherConditions | None) -> tuple[int, str]:
    if park is not None and park.roof in ("closed", "dome", "retractable"):
        return -1, f"roof ({park.roof})"
    if cond is None:
        return 0, "n/a"
    along = cond.wind_mph * math.cos(math.radians(45))
    blowing_in = cond.wind_mph > 0 and cond.out_to_cf_mph <= -along
    blowing_out = cond.wind_mph > 0 and cond.out_to_cf_mph >= along
    w = wind_pts(cond.wind_mph) if blowing_out else -wind_pts(cond.wind_mph) if blowing_in else 0
    pts = temp_pts(cond.temp_f) + w + humidity_pts(cond.humidity_pct)
    direction = "in" if blowing_in else "out" if blowing_out else "cross"
    return pts, f"{cond.temp_f:.0f}F, {cond.wind_mph:.0f} mph {direction}, {cond.humidity_pct:.0f}%"


def baseruns_per_game(r: dict[str, float], games: int) -> float | None:
    if not games:
        return None
    a = r["H"] + r["BB"] + r["HBP"] - r["HR"] - 0.5 * r["IBB"]
    tb = r["1B"] + 2 * r["2B"] + 3 * r["3B"] + 4 * r["HR"]
    b = 1.1 * (
        1.4 * tb - 0.6 * r["H"] - 3 * r["HR"] + 0.1 * (r["BB"] + r["HBP"] - r["IBB"])
        + 0.9 * (r["SB"] - r["CS"] - r["GDP"])
    )
    c = r["AB"] - r["H"] + r["CS"] + r["GDP"]
    if b + c <= 0:
        return None
    return (a * b / (b + c) + r["HR"]) / games


# --- sources -------------------------------------------------------------------


@dataclass
class FanGraphsTables:
    bat_all: dict[str, dict] = field(default_factory=dict)
    bat_vs_l: dict[str, dict] = field(default_factory=dict)
    bat_vs_r: dict[str, dict] = field(default_factory=dict)
    rel: dict[str, dict] = field(default_factory=dict)
    pit: dict[int, dict] = field(default_factory=dict)
    leverage: dict[str, list[int]] = field(default_factory=dict)


def _fg(stats: str, season: int, month: int = 0, team: str = "0,ts") -> list[dict]:
    url = _FG_URL.format(stats=stats, season=season, month=month, team=team.replace(",", "%2C"))
    resp = http.get(url, headers=_FG_HEADERS, timeout=60)
    resp.raise_for_status()
    return list(resp.json().get("data", []))


def fetch_fangraphs(season: int) -> FanGraphsTables:
    t = FanGraphsTables()
    t.bat_all = {r["TeamName"]: r for r in _fg("bat", season)}
    t.bat_vs_l = {r["TeamName"]: r for r in _fg("bat", season, _FG_MONTH_VS_L)}
    t.bat_vs_r = {r["TeamName"]: r for r in _fg("bat", season, _FG_MONTH_VS_R)}
    t.rel = {r["TeamName"]: r for r in _fg("rel", season)}
    t.pit = {r["xMLBAMID"]: r for r in _fg("pit", season, team="0") if r.get("xMLBAMID")}
    by_team: dict[str, list[tuple[float, int]]] = {}
    for r in _fg("rel", season, team="0"):
        if r.get("xMLBAMID"):
            by_team.setdefault(r["Team"], []).append(
                ((r.get("SV") or 0) + (r.get("HLD") or 0), r["xMLBAMID"])
            )
    t.leverage = {tm: [pid for _, pid in sorted(v, reverse=True)[:3]] for tm, v in by_team.items()}
    return t


def _fg_team(abbrev: str) -> str:
    return _FG_ABBR.get(abbrev, abbrev)


def games_played(season: int) -> dict[int, int]:
    resp = http.get(f"{STATSAPI}/standings", params={"leagueId": "103,104", "season": season}, timeout=30)
    resp.raise_for_status()
    return {
        tr["team"]["id"]: int(tr.get("gamesPlayed") or 0)
        for rec in resp.json().get("records", [])
        for tr in rec.get("teamRecords", [])
    }


def _schedule(day: Date, **hydrate: str) -> list[dict]:
    params: dict[str, str | int] = {"sportId": 1, "date": day.isoformat(), **hydrate}
    resp = http.get(f"{STATSAPI}/schedule", params=params, timeout=30)
    resp.raise_for_status()
    dates = resp.json().get("dates") or []
    return list(dates[0].get("games", [])) if dates else []


def _hp_umpire(g: dict) -> str | None:
    return next(
        (o["official"]["fullName"] for o in g.get("officials", []) if o.get("officialType") == "Home Plate"),
        None,
    )


def home_plate_umpires(day: Date) -> dict[int, tuple[str | None, str]]:
    """game_pk -> (umpire, "posted"|"projected"|"unknown").

    MLB posts the crew the morning of the game. Until then the crew rotates one
    base clockwise a day, so yesterday's first-base umpire for the same series is
    the likely plate umpire, and is labelled as a projection.
    """
    first_base: dict[frozenset[int], str | None] = {}
    for g in _schedule(day - timedelta(days=1), hydrate="officials"):
        offs = {o.get("officialType"): o["official"]["fullName"] for o in g.get("officials", [])}
        key = frozenset((g["teams"]["away"]["team"]["id"], g["teams"]["home"]["team"]["id"]))
        first_base[key] = offs.get("First Base")
    out: dict[int, tuple[str | None, str]] = {}
    for g in _schedule(day, hydrate="officials"):
        hp = _hp_umpire(g)
        if hp:
            out[g["gamePk"]] = (hp, "posted")
            continue
        key = frozenset((g["teams"]["away"]["team"]["id"], g["teams"]["home"]["team"]["id"]))
        fb = first_base.get(key)
        out[g["gamePk"]] = (fb, "projected" if fb else "unknown")
    return out


# --- boxscore cache: pitcher usage and umpire of record, one record per final ---


class BoxscoreCache:
    """Compact facts from every regular-season final, kept up to date incrementally."""

    def __init__(self, path: Path, season: int) -> None:
        self.path = path
        self.season = season
        self.records: dict[int, dict] = {}
        if path.exists():
            try:
                self.records = {int(r["pk"]): r for r in json.loads(path.read_text())}
            except (ValueError, KeyError, TypeError):
                self.records = {}

    def refresh(self, through: Date, workers: int = 8) -> None:
        start = Date(self.season, *SEASON_START)
        params: dict[str, str | int] = {
            "sportId": 1, "gameType": "R",
            "startDate": start.isoformat(), "endDate": through.isoformat(),
        }
        resp = http.get(f"{STATSAPI}/schedule", params=params, timeout=60)
        resp.raise_for_status()
        todo = [
            g
            for d in resp.json().get("dates", [])
            for g in d.get("games", [])
            if g["status"]["abstractGameState"] == "Final" and g["gamePk"] not in self.records
        ]
        if not todo:
            return
        log.info("totals sheet: fetching %d boxscores", len(todo))
        with ThreadPoolExecutor(workers) as ex:
            for rec in ex.map(_box_record, todo):
                if rec is not None:
                    self.records[rec["pk"]] = rec
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(list(self.records.values())))

    def reliever_usage(self, team_id: int, day: Date, back: int = 2) -> tuple[dict[int, list[int]], list[int]]:
        """{pitcher_id: [pitches day-1, day-2, ...]} and the pen's total per day."""
        per: dict[int, list[int]] = {}
        total = [0] * back
        for k in range(1, back + 1):
            d = (day - timedelta(days=k)).isoformat()
            for b in self.records.values():
                if b["date"] != d:
                    continue
                for side in ("away", "home"):
                    if b[side] != team_id:
                        continue
                    for a in b["teams"][side]["arms"]:
                        if a["starter"]:
                            continue
                        n = int(a["pitches"] or 0)
                        per.setdefault(a["id"], [0] * back)[k - 1] += n
                        total[k - 1] += n
        return per, total

    def umpire_table(self, before: Date) -> tuple[dict[str, tuple[int, float, float]], float]:
        """umpire -> (games, runs/game, K%) on finals strictly before ``before``; league r/g."""
        acc: dict[str, list[float]] = {}
        lg = [0.0, 0.0]
        cutoff = before.isoformat()
        for b in self.records.values():
            if b["date"] >= cutoff or not b.get("hp"):
                continue
            so = sum(b["teams"][s]["so_pitch"] or 0 for s in ("away", "home"))
            pa = sum(b["teams"][s]["pa"] or 0 for s in ("away", "home"))
            u = acc.setdefault(b["hp"], [0, 0.0, 0.0, 0.0])
            u[0] += 1
            u[1] += b["runs"]
            u[2] += so
            u[3] += pa
            lg[0] += 1
            lg[1] += b["runs"]
        table = {k: (int(v[0]), v[1] / v[0], v[2] / max(v[3], 1)) for k, v in acc.items() if v[0]}
        return table, lg[1] / max(lg[0], 1)


def _box_record(g: dict) -> dict | None:
    pk = g["gamePk"]
    try:
        resp = http.get(f"{STATSAPI}/game/{pk}/boxscore", timeout=60)
        resp.raise_for_status()
        b = resp.json()
    except Exception as exc:  # one missing box must not sink the sheet
        log.warning("boxscore %s failed: %s", pk, exc)
        return None
    rec: dict = {
        "pk": pk, "date": g["officialDate"], "hp": _hp_umpire(b),
        "away": g["teams"]["away"]["team"]["id"], "home": g["teams"]["home"]["team"]["id"],
        "runs": (g["teams"]["away"].get("score") or 0) + (g["teams"]["home"].get("score") or 0),
        "teams": {},
    }
    for side in ("away", "home"):
        t = b["teams"][side]
        pitchers = t.get("pitchers", [])
        arms = []
        for pid in pitchers:
            st = t["players"].get(f"ID{pid}", {}).get("stats", {}).get("pitching", {})
            if not st:
                continue
            arms.append({
                "id": pid,
                "pitches": st.get("numberOfPitches", st.get("pitchesThrown", 0)),
                "starter": pid == pitchers[0],
            })
        ps = t.get("teamStats", {}).get("pitching", {})
        bs = t.get("teamStats", {}).get("batting", {})
        rec["teams"][side] = {"arms": arms, "so_pitch": ps.get("strikeOuts"), "pa": bs.get("plateAppearances")}
    return rec


def fatigue_pts(box: BoxscoreCache, team: TeamGameInfo, leverage: list[int], day: Date) -> tuple[int, str]:
    """+1 per top-three leverage arm used yesterday (+2 if both of the last two
    days), +1 if the pen threw >= 120 pitches over those two days; capped."""
    per, total = box.reliever_usage(team.team_id, day)
    pts = 0
    detail: list[str] = []
    for pid in leverage:
        y, y2 = per.get(pid, [0, 0])
        if y and y2:
            pts += 2
            detail.append(f"{pid} b2b")
        elif y:
            pts += 1
            detail.append(f"{pid} yday")
    if sum(total) >= PEN_PITCHES_TIRED:
        pts += 1
        detail.append(f"pen {sum(total)}p/2d")
    return min(pts, FATIGUE_CAP), ", ".join(detail)


def umpire_pts(box: BoxscoreCache, name: str | None, day: Date) -> tuple[int, str]:
    table, lg = box.umpire_table(day)
    if not name or name not in table or table[name][0] < UMP_MIN_GAMES:
        return 0, ""
    g, rg, k = table[name]
    diff = rg - lg
    pts = 1 if diff >= UMP_BAND_RUNS else -1 if diff <= -UMP_BAND_RUNS else 0
    return pts, f"{rg:.1f} r/g ({diff:+.1f} vs lg), K {k:.1%}, n={g}"


# --- the row -------------------------------------------------------------------


@dataclass
class TeamSide:
    abbrev: str
    starter: str
    off: int
    sp: int
    rp: int
    kbb_sp: int
    kbb_rp: int
    bsr_pg: float | None
    fatigue: int
    fatigue_detail: str


@dataclass
class SheetRow:
    game: str
    first_pitch_utc: str | None
    total: str
    away: TeamSide
    home: TeamSide
    circa: int
    dk: int
    weather: int
    weather_detail: str
    park: int
    park_detail: str
    bsr: int
    ump: int
    ump_name: str
    ump_status: str
    ump_detail: str
    # The simulator's own median total and its over probability at the sheet's
    # line, from the engine ledger; None when the engine did not price the game.
    # Shown beside SUM and never added to it: the sheet is the hand method, and
    # the column is there so the ledger can grade the two against each other.
    engine_total: float | None = None
    engine_p_over: float | None = None

    @property
    def engine_delta(self) -> float | None:
        """Engine median minus the sheet's leading line; + when the engine sits over it."""
        line = first_line(self.total)
        if self.engine_total is None or line is None:
            return None
        return round(self.engine_total - line, 1)

    @property
    def total_pts(self) -> int:
        a, h = self.away, self.home
        return (
            a.off + h.off + a.sp + h.sp + a.rp + h.rp
            + a.kbb_sp + a.kbb_rp + h.kbb_sp + h.kbb_rp
            + self.circa + self.dk + self.weather + self.park + self.bsr
            + a.fatigue + h.fatigue + self.ump
        )


def _offense(fg: FanGraphsTables, team: TeamGameInfo, opp: TeamGameInfo) -> int:
    key = _fg_team(team.abbrev)
    hand = opp.probable_pitcher.throws if opp.probable_pitcher else None
    split = fg.bat_vs_l if hand == "L" else fg.bat_vs_r if hand == "R" else fg.bat_all
    b = split.get(key) or fg.bat_all.get(key)
    allb = fg.bat_all.get(key)
    pts = 0
    if b:
        pts += wrc_pts(b["wRC+"]) + woba_pts(b["wOBA"])
    if allb and allb.get("Barrel%") is not None:
        pts += barrel_pts(allb["Barrel%"] * 100)
    return pts


def _arm(row: dict | None) -> tuple[int, int]:
    """(SIERA + xERA + CSW points, K-BB points); an unknown arm scores 0."""
    if not row:
        return 0, 0
    pts = siera_pts(row["SIERA"]) + csw_pts(row["C+SwStr%"] * 100)
    if row.get("xERA") is not None:
        pts += xera_pts(row["xERA"])
    return pts, kbb_pts(row["K-BB%"] * 100)


def _side(
    fg: FanGraphsTables, box: BoxscoreCache, gp: dict[int, int],
    team: TeamGameInfo, opp: TeamGameInfo, day: Date,
) -> TeamSide:
    key = _fg_team(team.abbrev)
    pp = team.probable_pitcher
    sp, kbb_sp = _arm(fg.pit.get(pp.mlbam_id) if pp else None)
    rp, kbb_rp = _arm(fg.rel.get(key))
    allb = fg.bat_all.get(key)
    bsr = baseruns_per_game(allb, gp.get(team.team_id, 0)) if allb else None
    fat, fat_detail = fatigue_pts(box, team, fg.leverage.get(key, []), day)
    return TeamSide(
        abbrev=team.abbrev,
        starter=f"{pp.name} ({pp.throws or '?'})" if pp else "TBD",
        off=_offense(fg, team, opp),
        sp=sp, rp=rp, kbb_sp=kbb_sp, kbb_rp=kbb_rp,
        bsr_pg=bsr, fatigue=fat, fatigue_detail=fat_detail,
    )


def _total_label(splits: VSINClient.TotalSplits, matchup: str) -> str:
    dk = splits.get((matchup, "draftkings"))
    circa = splits.get((matchup, "circa"))
    if dk and circa and dk.line != circa.line:
        return f"{circa.line:g} / dk {dk.line:g}"
    src = dk or circa
    return f"{src.line:g}" if src else ""


def build_rows(
    day: Date, slate: Slate, fg: FanGraphsTables, box: BoxscoreCache, gp: dict[int, int],
    splits: VSINClient.TotalSplits, weather: WeatherProvider, umps: dict[int, tuple[str | None, str]],
    engine: dict[str, OverCurve] | None = None,
) -> list[SheetRow]:
    rows: list[SheetRow] = []
    for g in slate.games:
        away = _side(fg, box, gp, g.away, g.home, day)
        home = _side(fg, box, gp, g.home, g.away, day)
        park = get_park(g.venue.venue_id)
        cond = weather.fetch(park, g.game_datetime_utc).conditions if park else None
        wx, wx_detail = weather_pts(park, cond)
        pf = park.park_factor if park else 100.0
        bsr = sum(baseruns_pts(s.bsr_pg) for s in (away, home) if s.bsr_pg is not None)
        ump_name, status = umps.get(g.game_pk, (None, "unknown"))
        ump, ump_detail = umpire_pts(box, ump_name, day)
        matchup = g.matchup()
        total = _total_label(splits, matchup)
        curve = (engine or {}).get(matchup) or {}
        p_over, _ = engine_at(curve, first_line(total))
        rows.append(SheetRow(
            game=matchup,
            first_pitch_utc=g.game_datetime_utc,
            total=total,
            engine_total=engine_median(curve) if curve else None,
            engine_p_over=p_over,
            away=away, home=home,
            circa=book_pts(splits.get((matchup, "circa"))),
            dk=book_pts(splits.get((matchup, "draftkings"))),
            weather=wx, weather_detail=wx_detail,
            park=park_pts(pf), park_detail=f"{park.name} {pf:g}" if park else "unknown park",
            bsr=bsr,
            ump=ump, ump_name=ump_name or "", ump_status=status, ump_detail=ump_detail,
        ))
    rows.sort(key=lambda r: -r.total_pts)
    return rows


# --- workbook ------------------------------------------------------------------

_COLUMNS = [
    "Game", "Total", "SP A", "SP H", "Off A", "Off H", "SP A pts", "SP H pts", "RP A pts", "RP H pts",
    "K-BB SP A", "K-BB RP A", "K-BB SP H", "K-BB RP H", "Circa", "DK", "Weather", "Park", "BsR",
    "Pen A", "Pen H", "Ump", "SUM",
    "Engine", "Eng vs line", "Eng O%",
    "Weather detail", "Park", "BsR/G A", "BsR/G H", "Pen A detail", "Pen H detail",
    "HP umpire", "Ump status", "Ump detail",
]

_LEGEND = [
    ("Bands", BANDS),
    ("Sign", "+ leans Over, - leans Under; SUM is every points column added. Bigger magnitude = stronger lean."),
    ("Engine", "The simulator's median total for the game, read off the engine ledger's game_total rows (calibrated, run-environment corrected, before the market anchor). "
               "Eng vs line = Engine minus the leading Total; Eng O% = the engine's over probability at that line. Not part of SUM; the audit grades the sheet and the engine side by side and splits the record by whether they agreed."),
    ("Off A / Off H", "Away / home offense: wRC+ pts + wOBA pts (team split vs the opposing starter's hand) + Barrel% pts (season)."),
    ("Centre", "Every band's 0 window is the league's middle half (2026 FanGraphs), so an average arm or lineup adds nothing and +5 / -5 are equal leans in opposite directions."),
    ("wRC+", ">130 3 | 116-130 2 | 106-115 1 | 95-105 0 | 85-94 -1 | 70-84 -2 | <70 -3"),
    ("wOBA", ">.355 3 | .341-.355 2 | .326-.340 1 | .305-.325 0 | .290-.304 -1 | .275-.289 -2 | <.275 -3"),
    ("Barrel%", ">10 2 | 8.6-10 1 | 6.5-8.5 0 | 5-6.4 -1 | <5 -2"),
    ("SP / RP pts", "Starter and bullpen: SIERA pts + xERA pts + CSW% pts. Softer arm = positive, league-average arm = 0. TBD starter = 0."),
    ("SIERA", "<2.75 -3 | 2.75-3.24 -2 | 3.25-3.74 -1 | 3.75-4.35 0 | 4.36-4.75 1 | 4.76-5.25 2 | >5.25 3"),
    ("xERA", "<3.20 -2 | 3.20-3.69 -1 | 3.70-4.50 0 | 4.51-5.10 1 | >5.10 2"),
    ("CSW%", ">31 -2 | 29-31 -1 | 25-29 0 | 23-25 1 | <23 2"),
    ("K-BB%", ">25 -3 | 19-25 -2 | 15-19 -1 | 10-15 0 | 7-10 1 | 4-7 2 | <4 3 (starter and pen, each team)"),
    ("Circa / DK", "VSIN total handle% - bets% on the side the money is on: >=20 2, 11-19 1, else 0; a 100/100 split is one ticket = 0. Signed + Over / - Under. Missing = 0."),
    ("Weather", "First pitch, Open-Meteo: temp >90 2, >80 1, 60-79 0, 50-59 -1, <50 -2; humidity >65 1, 35-65 0, <35 -1; wind by direction (within 45 deg of the home plate-CF line): blowing out +1 (5-9 mph) / +2 (10-15) / +3 (>15), blowing in the same points negative, cross wind or <5 mph 0. Any roof (dome or retractable) = -1 total."),
    ("Park", "Engine park factor (runs): >102 1, <98 -1, else 0."),
    ("BsR", "PROVISIONAL bands (not hand-written): team BaseRuns/G >=4.65 1, <=4.25 -1, else 0; both teams summed."),
    ("Pen A / Pen H", "Bullpen fatigue: +1 per top-3 leverage arm (SV+HLD) used yesterday, +2 if used both of the last two days, +1 if the pen threw 120+ pitches over the last two days; capped at 4."),
    ("Ump", "Home-plate umpire's 2026 runs/game vs league on finals before today: >= +1.0 run 1, <= -1.0 -1, else 0; needs 10+ games. 'projected' = crew not posted yet; yesterday's 1B umpire assumed by rotation."),
    ("Backtest", "Aug 4 - Sep 8 2026, 452 games: SUM correlates 0.49 with the closing total and ~0.03 with the result vs the close; top-3 overs / bottom-3 unders per day hit 50% / 47%. Read it as what the market already knows."),
]


def write_workbook(rows: list[SheetRow], day: Date, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Totals {day.isoformat()}"
    ws.append(_COLUMNS)
    for r in rows:
        a, h = r.away, r.home
        ws.append([
            r.game, r.total, a.starter, h.starter, a.off, h.off, a.sp, h.sp, a.rp, h.rp,
            a.kbb_sp, a.kbb_rp, h.kbb_sp, h.kbb_rp, r.circa, r.dk, r.weather, r.park, r.bsr,
            a.fatigue, h.fatigue, r.ump, r.total_pts,
            r.engine_total, r.engine_delta, None if r.engine_p_over is None else round(r.engine_p_over, 3),
            r.weather_detail, r.park_detail,
            round(a.bsr_pg, 2) if a.bsr_pg is not None else None,
            round(h.bsr_pg, 2) if h.bsr_pg is not None else None,
            a.fatigue_detail, h.fatigue_detail, r.ump_name, r.ump_status, r.ump_detail,
        ])
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    sum_col = _COLUMNS.index("SUM") + 1
    over = PatternFill("solid", fgColor="C6EFCE")
    under = PatternFill("solid", fgColor="FFC7CE")
    for row in ws.iter_rows(min_row=2):
        cell = row[sum_col - 1]
        cell.font = bold
        if isinstance(cell.value, int) and cell.value > 0:
            cell.fill = over
        elif isinstance(cell.value, int) and cell.value < 0:
            cell.fill = under
    for i, name in enumerate(_COLUMNS, start=1):
        width = 14 if i <= 4 else 7 if i <= sum_col else 9 if i <= sum_col + 3 else 22
        ws.column_dimensions[get_column_letter(i)].width = max(width, min(len(name) + 2, 14))
    ws.freeze_panes = "C2"

    legend = wb.create_sheet("Legend")
    legend.append(["Column", "Rule"])
    for c in legend[1]:
        c.font = bold
    for k, v in _LEGEND:
        legend.append([k, v])
    legend.column_dimensions["A"].width = 16
    legend.column_dimensions["B"].width = 140
    for row in legend.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def output_path(cfg: Config, day: Date) -> Path:
    return cfg.output_dir / f"totals_sheet_{day.isoformat()}.xlsx"


def sheet_is_current(path: Path) -> bool:
    """True when a sheet already on disk was scored by the bands this code carries."""
    if not path.exists():
        return False
    try:
        return sheet_bands(path) == BANDS
    except Exception as exc:
        log.warning("totals sheet: could not read %s: %s", path.name, exc)
        return False


def build_totals_sheet(cfg: Config, day: Date, *, if_stale: bool = False) -> Path | None:
    """Score the day's slate and write ``totals_sheet_<day>.xlsx``; None when no games.

    With ``if_stale`` a sheet already written by these bands is left as it is,
    so a job can call this every pass and only rewrite when the bands moved on.
    """
    out = output_path(cfg, day)
    if if_stale and sheet_is_current(out):
        log.info("totals sheet: %s already on bands %s", out.name, BANDS)
        return out
    stats = MLBStatsClient()
    slate = stats.get_slate(day)
    if not slate.games:
        log.info("totals sheet: no games on %s", day)
        return None
    season = day.year
    fg = fetch_fangraphs(season)
    gp = games_played(season)
    box = BoxscoreCache(cfg.cache_dir / f"totals_boxscores_{season}.json", season)
    try:
        box.refresh(day - timedelta(days=1))
    except Exception as exc:  # stale usage beats no sheet
        log.warning("totals sheet: boxscore refresh failed: %s", exc)
    splits: VSINClient.TotalSplits = {}
    try:
        splits = VSINClient(cfg.creds).fetch_total_splits(slate)
    except Exception as exc:
        log.warning("totals sheet: VSIN splits unavailable: %s", exc)
    try:
        umps = home_plate_umpires(day)
    except Exception as exc:
        log.warning("totals sheet: umpire feed unavailable: %s", exc)
        umps = {}
    weather = WeatherProvider(cache_dir=cfg.weather_cache_dir)
    engine: dict[str, OverCurve] = {}
    try:
        engine = over_curves(load_ledger(engine_ledger_path(cfg)), day)
    except Exception as exc:
        log.warning("totals sheet: engine ledger unreadable: %s", exc)
    if not engine:
        log.info("totals sheet: no engine game totals for %s; Engine columns left blank", day)
    rows = build_rows(day, slate, fg, box, gp, splits, weather, umps, engine)
    path = write_workbook(rows, day, out)
    try:
        record_sheet(cfg, day, path, {g.matchup(): g.game_pk for g in slate.games})
    except Exception as exc:  # the sheet is the deliverable; the ledger row is the receipt
        log.warning("totals sheet: ledger not updated: %s", exc)
    return path

