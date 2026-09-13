"""Grade the hand totals sheet the morning after, and keep the running ledger.

The sheet says + for Over and - for Under, bigger for stronger. This module is
its receipt: every row the sheet showed is appended to ``totals_ledger.csv``
with its line and SUM, the next run fills in the final score, and the daily
package carries ``totals_audit_<day>.xlsx`` -- yesterday graded, the whole
ledger, and the cumulative hit rates the sheet has actually earned:

* **Sign**: did the Over hit when SUM > 0, the Under when SUM < 0?
* **Rank**: each day's top-3 sums as Overs and bottom-3 as Unders, counted only
  when the sum's sign agrees with the side (a +3 at the bottom of an all-Over
  day is not an Under).
* **Bucket**: Over rate by SUM band, so a strong lean can be told from a weak one.
* **Engine**: the simulator's own total for the same game, read off the engine
  ledger's ``game_total`` rows (its calibrated over probability at each posted
  line), graded the same way and split by whether it agreed with the sheet. The
  sheet's backtest says SUM tracks the closing number and not the result; what
  the engine's *disagreement* with that number is worth is the open question
  this column exists to answer.

Every row carries the version of the bands that scored it. The tallies only
count rows scored by the current bands: a sheet from an older band set is on a
different scale (the pre-centred bands scored a league-average arm +1), so its
rows stay in the ledger as history and are excluded from the rates.

Nothing here feeds a price. It is a record.
"""

from __future__ import annotations

import csv
import logging
import math
import re
from dataclasses import asdict, dataclass, field, fields
from datetime import date as Date
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from mlb_engine.audit.ledger import ENGINE, LedgerEntry, load_ledger
from mlb_engine.config import Config
from mlb_engine.data import http
from mlb_engine.data.mlb_statsapi import BASE as STATSAPI

log = logging.getLogger(__name__)

LEDGER_NAME = "totals_ledger.csv"
ENGINE_LEDGER_NAME = "ledger.csv"
GAME_TOTAL = "game_total"
OVER, UNDER, PUSH = "over", "under", "push"
RANK_N = 3
RANK_MIN_GAMES = 4  # fewer graded games than this and the day's ends are not ranked
# Version of the scoring bands; bump when a band table changes so sheets and
# ledger rows scored on the old scale can be told apart from the new.
BANDS = "centred-2026.09"
LEGACY = "legacy"

# Sheet rows that reference the sheet's own columns, not the audit's.
_GAME, _TOTAL, _SUM = "Game", "Total", "SUM"
_LEGEND_SHEET, _BANDS_KEY = "Legend", "Bands"


@dataclass
class LedgerRow:
    date: str
    game: str
    game_pk: int
    line: float | None
    sum_pts: int
    away_runs: int | None = None
    home_runs: int | None = None
    result: str = ""  # over | under | push | "" (ungraded)
    bands: str = LEGACY  # BANDS version that scored sum_pts; LEGACY predates versioning
    # The engine's read of the same game: its median total, its calibrated over
    # probability at the sheet's line, and the devigged market over at that line
    # when a book posted it. None when the engine ledger has no rows for the game.
    engine_total: float | None = None
    engine_p_over: float | None = None
    market_p_over: float | None = None

    @property
    def current(self) -> bool:
        return self.bands == BANDS

    @property
    def graded(self) -> bool:
        return self.result != ""

    @property
    def runs(self) -> int | None:
        if self.away_runs is None or self.home_runs is None:
            return None
        return self.away_runs + self.home_runs

    @property
    def lean(self) -> str:
        return OVER if self.sum_pts > 0 else UNDER if self.sum_pts < 0 else ""

    @property
    def hit(self) -> bool | None:
        """True when the sign of SUM matched the result; None on push, no lean or ungraded."""
        if not self.graded or self.result == PUSH or not self.lean:
            return None
        return self.lean == self.result

    @property
    def engine_lean(self) -> str:
        """Which side of the sheet's line the engine's own total sits on."""
        p = self.engine_p_over
        if p is None:
            return ""
        return OVER if p > 0.5 else UNDER if p < 0.5 else ""

    @property
    def engine_hit(self) -> bool | None:
        if not self.graded or self.result == PUSH or not self.engine_lean:
            return None
        return self.engine_lean == self.result

    @property
    def edge_lean(self) -> str:
        """The side the engine likes more than the market does, at the sheet's line."""
        p, m = self.engine_p_over, self.market_p_over
        if p is None or m is None:
            return ""
        return OVER if p > m else UNDER if p < m else ""

    @property
    def edge_hit(self) -> bool | None:
        if not self.graded or self.result == PUSH or not self.edge_lean:
            return None
        return self.edge_lean == self.result

    @property
    def agreement(self) -> str:
        """'agree' when sheet and engine lean the same way, 'split' when opposite, '' when either is silent."""
        if not self.lean or not self.engine_lean:
            return ""
        return "agree" if self.lean == self.engine_lean else "split"


_FIELDS = tuple(f.name for f in fields(LedgerRow))


# --- ledger ---------------------------------------------------------------------


def ledger_path(cfg: Config) -> Path:
    return cfg.audit_dir / LEDGER_NAME


def read_ledger(path: Path) -> list[LedgerRow]:
    if not path.exists():
        return []
    out: list[LedgerRow] = []
    with path.open(newline="") as fh:
        for r in csv.DictReader(fh):
            out.append(LedgerRow(
                date=r["date"], game=r["game"], game_pk=int(r["game_pk"] or 0),
                line=float(r["line"]) if r["line"] else None, sum_pts=int(r["sum_pts"]),
                away_runs=int(r["away_runs"]) if r["away_runs"] else None,
                home_runs=int(r["home_runs"]) if r["home_runs"] else None,
                result=r["result"],
                bands=r.get("bands") or LEGACY,
                engine_total=_opt_float(r.get("engine_total")),
                engine_p_over=_opt_float(r.get("engine_p_over")),
                market_p_over=_opt_float(r.get("market_p_over")),
            ))
    return out


def _opt_float(v: str | None) -> float | None:
    return float(v) if v else None


def write_ledger(path: Path, rows: list[LedgerRow]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    with tmp.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=_FIELDS)
        w.writeheader()
        for r in sorted(rows, key=lambda r: (r.date, -r.sum_pts, r.game)):
            d = asdict(r)
            w.writerow({k: "" if d[k] is None else d[k] for k in _FIELDS})
    tmp.replace(path)


def first_line(total: str) -> float | None:
    """'8.5 / dk 8' -> 8.5 (the Circa number leads; DK when it is the only one)."""
    m = re.search(r"\d+(\.\d+)?", str(total or ""))
    return float(m.group()) if m else None


def _bands_of(wb) -> str:
    """The band version stamped on a sheet's Legend; LEGACY when it has none."""
    if _LEGEND_SHEET not in wb.sheetnames:
        return LEGACY
    rules = {str(k): v for k, v, *_ in wb[_LEGEND_SHEET].iter_rows(values_only=True) if k is not None}
    if rules.get(_BANDS_KEY):
        return str(rules[_BANDS_KEY])
    # The first centred sheets shipped before the stamp; only they carry the Centre rule.
    return BANDS if "Centre" in rules else LEGACY


def sheet_bands(sheet: Path) -> str:
    return _bands_of(load_workbook(sheet, read_only=True))


def rows_from_sheet(sheet: Path, day: Date, game_pks: dict[str, int]) -> list[LedgerRow]:
    """One ungraded ledger row per game on the day's sheet."""
    wb = load_workbook(sheet, read_only=True)
    bands = _bands_of(wb)
    ws = wb[f"Totals {day.isoformat()}"]
    it = ws.iter_rows(values_only=True)
    header = [str(c) for c in next(it)]
    gi, ti, si = header.index(_GAME), header.index(_TOTAL), header.index(_SUM)
    out: list[LedgerRow] = []
    for r in it:
        if r[gi] is None or not isinstance(r[si], int):
            continue
        game = str(r[gi])
        out.append(LedgerRow(
            day.isoformat(), game, game_pks.get(game, 0), first_line(str(r[ti] or "")), r[si], bands=bands,
        ))
    return out


# --- the engine's total -------------------------------------------------------------


OverCurve = dict[float, tuple[float, float | None]]  # line -> (engine p_over, market p_over)


def over_curves(entries: list[LedgerEntry], day: Date) -> dict[str, OverCurve]:
    """Matchup -> the engine's calibrated over probability at every game-total line it posted that day.

    ``model_prob`` is the number the engine actually stood behind: calibrated,
    corrected into the league's run environment, before the market anchor. The
    market side is the devigged over where a book priced that line; an under-only
    quote is read as one minus its fair probability.
    """
    iso = day.isoformat()
    p: dict[str, dict[float, list[float]]] = {}
    m: dict[str, dict[float, list[float]]] = {}
    for e in entries:
        if e.source != ENGINE or e.market != GAME_TOTAL or e.date != iso or e.line is None:
            continue
        side = e.selection.split()[0].lower() if e.selection else ""
        if side not in (OVER, UNDER):
            continue
        over = side == OVER
        line = float(e.line)
        if over:
            p.setdefault(e.matchup, {}).setdefault(line, []).append(e.model_prob)
        if e.fair_prob is not None:
            m.setdefault(e.matchup, {}).setdefault(line, []).append(e.fair_prob if over else 1.0 - e.fair_prob)
    out: dict[str, OverCurve] = {}
    for matchup, lines in p.items():
        out[matchup] = {
            line: (sum(v) / len(v), (sum(m[matchup][line]) / len(m[matchup][line])) if m.get(matchup, {}).get(line) else None)
            for line, v in sorted(lines.items())
        }
    return out


def _interp(xs: list[float], ys: list[float], x: float) -> float | None:
    """Linear read of ``ys`` at ``x``; None outside the posted lines (no extrapolation)."""
    if not xs or x < xs[0] or x > xs[-1]:
        return None
    for i in range(len(xs) - 1):
        x0, x1 = xs[i], xs[i + 1]
        if x0 <= x <= x1:
            return ys[i] if x1 == x0 else ys[i] + (ys[i + 1] - ys[i]) * (x - x0) / (x1 - x0)
    return ys[-1]


def engine_median(curve: OverCurve) -> float | None:
    """The total at which the engine's over probability crosses one half.

    Read between the posted lines only. When every line sits on one side of .5
    the engine's total is off the board and the median is None -- extending a
    near-flat curve past its ends invented totals of 17 runs -- and the row
    still carries the engine's over probability at the sheet's line.
    """
    xs = sorted(curve)
    ys = [curve[x][0] for x in xs]
    for i in range(len(xs) - 1):
        y0, y1 = ys[i], ys[i + 1]
        if (y0 - 0.5) * (y1 - 0.5) <= 0 and y0 != y1:
            return round(xs[i] + (xs[i + 1] - xs[i]) * (0.5 - y0) / (y1 - y0), 2)
    return None


def engine_at(curve: OverCurve, line: float | None) -> tuple[float | None, float | None]:
    """(engine over probability, market over probability) at the sheet's line.

    The engine's read is interpolated between posted lines; the market's is only
    taken where a book actually priced that line.
    """
    if line is None or not curve:
        return None, None
    xs = sorted(curve)
    p = _interp(xs, [curve[x][0] for x in xs], float(line))
    mk = curve.get(float(line), (None, None))[1]
    return (round(p, 4) if p is not None else None), (round(mk, 4) if mk is not None else None)


def attach_engine(rows: list[LedgerRow], entries: list[LedgerEntry]) -> int:
    """Fill the engine's read into rows that lack one; returns how many were filled."""
    curves: dict[str, dict[str, OverCurve]] = {}
    n = 0
    for r in rows:
        if r.engine_p_over is not None:
            continue
        if r.date not in curves:
            curves[r.date] = over_curves(entries, Date.fromisoformat(r.date))
        curve = curves[r.date].get(r.game)
        if not curve:
            continue
        r.engine_total = engine_median(curve)
        r.engine_p_over, r.market_p_over = engine_at(curve, r.line)
        if r.engine_p_over is not None:
            n += 1
    return n


def engine_ledger_path(cfg: Config) -> Path:
    return cfg.audit_dir / ENGINE_LEDGER_NAME


def merge(ledger: list[LedgerRow], fresh: list[LedgerRow]) -> list[LedgerRow]:
    """Add sheet rows the ledger has not seen; a re-run of a day never regrades or duplicates."""
    seen = {(r.date, r.game, r.game_pk) for r in ledger}
    return ledger + [r for r in fresh if (r.date, r.game, r.game_pk) not in seen]


# --- finals -----------------------------------------------------------------------


def _schedule(day: Date) -> list[dict]:
    resp = http.get(
        f"{STATSAPI}/schedule",
        params={"sportId": 1, "date": day.isoformat(), "hydrate": "team,linescore"},
        timeout=30,
    )
    resp.raise_for_status()
    dates = resp.json().get("dates") or []
    return list(dates[0].get("games", [])) if dates else []


def _matchup(g: dict) -> str:
    return f"{g['teams']['away']['team'].get('abbreviation', '')} @ {g['teams']['home']['team'].get('abbreviation', '')}"


Final = tuple[str, int, int]  # (matchup label, away runs, home runs)


def finals(day: Date) -> dict[int, Final]:
    """game_pk -> (matchup, away runs, home runs) for every game that reached Final."""
    out: dict[int, Final] = {}
    for g in _schedule(day):
        if g.get("status", {}).get("abstractGameState") != "Final":
            continue
        a, h = g["teams"]["away"], g["teams"]["home"]
        if "score" not in a or "score" not in h:
            continue
        out[int(g["gamePk"])] = (_matchup(g), int(a["score"]), int(h["score"]))
    return out


def game_pks(day: Date) -> dict[str, int]:
    return {_matchup(g): g["gamePk"] for g in _schedule(day)}


def grade(rows: list[LedgerRow], results: dict[int, Final]) -> int:
    """Fill finals into ungraded rows in place; returns how many were graded.

    Matched on game_pk when the sheet stored one, else on the ``A @ H`` label.
    A doubleheader's two games share a label, so the label match is only trusted
    when the day has one final between those clubs.
    """
    by_label: dict[str, list[tuple[int, int]]] = {}
    for label, a, h in results.values():
        by_label.setdefault(label, []).append((a, h))
    n = 0
    for r in rows:
        if r.graded:
            continue
        hit: tuple[int, int] | None = None
        if r.game_pk and r.game_pk in results:
            _, a, h = results[r.game_pk]
            hit = (a, h)
        elif not r.game_pk and len(by_label.get(r.game, [])) == 1:
            hit = by_label[r.game][0]
        if hit is None or r.line is None:
            continue
        r.away_runs, r.home_runs = hit
        total = hit[0] + hit[1]
        r.result = OVER if total > r.line else UNDER if total < r.line else PUSH
        n += 1
    return n


# --- summary ------------------------------------------------------------------------


@dataclass
class Tally:
    hits: int = 0
    misses: int = 0
    pushes: int = 0

    def add(self, hit: bool | None, pushed: bool = False) -> None:
        if pushed:
            self.pushes += 1
        elif hit is True:
            self.hits += 1
        elif hit is False:
            self.misses += 1

    @property
    def n(self) -> int:
        return self.hits + self.misses

    @property
    def rate(self) -> float | None:
        return self.hits / self.n if self.n else None

    def text(self) -> str:
        if not self.n:
            return "0-0"
        s = f"{self.hits}-{self.misses}"
        if self.pushes:
            s += f"-{self.pushes}"
        return f"{s} ({100 * self.hits / self.n:.0f}%)"


_BUCKETS: tuple[tuple[str, int, int | None], ...] = (
    ("<= -5", -10**6, -5),
    ("-4..-1", -4, -1),
    ("0", 0, 0),
    ("1..9", 1, 9),
    ("10..19", 10, 19),
    (">= 20", 20, None),
)


def bucket(sum_pts: int) -> str:
    for name, lo, hi in _BUCKETS:
        if sum_pts >= lo and (hi is None or sum_pts <= hi):
            return name
    return "?"


# |SUM| bands: the sheet's confidence scale, sign-agnostic (+10 and -10 are the same strength).
_MAG_BANDS: tuple[tuple[str, int, int | None], ...] = (
    ("1..4", 1, 4),
    ("5..9", 5, 9),
    ("10..14", 10, 14),
    (">= 15", 15, None),
)

BREAK_EVEN = 0.524  # -110 both ways
MIN_GRADED = 30  # no verdict on a band before this many decided games


def mag_band(sum_pts: int) -> str | None:
    m = abs(sum_pts)
    for name, lo, hi in _MAG_BANDS:
        if m >= lo and (hi is None or m <= hi):
            return name
    return None


def wilson(hits: int, n: int, z: float = 1.96) -> tuple[float, float]:
    """95% Wilson score interval for a hit rate; (0, 1) when nothing is graded."""
    if n == 0:
        return 0.0, 1.0
    p = hits / n
    denom = 1 + z * z / n
    centre = (p + z * z / (2 * n)) / denom
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / denom
    return max(0.0, centre - half), min(1.0, centre + half)


def verdict(t: Tally) -> str:
    """What the band's sample supports: 'edge' when its whole 95% range clears break-even,
    'no edge' when the whole range sits below it, else 'unproven' (always, under MIN_GRADED)."""
    if t.n < MIN_GRADED:
        return f"unproven (n<{MIN_GRADED})"
    lo, hi = wilson(t.hits, t.n)
    if lo > BREAK_EVEN:
        return "edge"
    if hi < BREAK_EVEN:
        return "no edge"
    return "unproven"


@dataclass
class Summary:
    sign: Tally
    rank_over: Tally
    rank_under: Tally
    over_rate: Tally  # hits = overs, misses = unders: the slate's own base rate
    by_bucket: dict[str, Tally]
    by_mag: dict[str, Tally]  # sign hit rate by |SUM| band: does a bigger number win more often?
    days: int
    games: int
    legacy: int = 0  # graded rows scored by older bands, kept but not counted
    # The engine's side of the line, the side it likes more than the market, and
    # the sheet's call split by whether the engine agreed with it.
    engine: Tally = field(default_factory=Tally)
    edge: Tally = field(default_factory=Tally)
    agree: Tally = field(default_factory=Tally)
    split_sheet: Tally = field(default_factory=Tally)  # the sheet's call where the engine leaned the other way

    @property
    def split_engine(self) -> Tally:
        """The engine's call on the split games: the sheet's misses are its hits."""
        s = self.split_sheet
        return Tally(hits=s.misses, misses=s.hits, pushes=s.pushes)


def summarize(rows: list[LedgerRow]) -> Summary:
    """Tallies over the graded rows scored by the current bands only."""
    legacy = sum(1 for r in rows if r.graded and not r.current)
    graded = [r for r in rows if r.graded and r.current]
    sign, rank_o, rank_u, base = Tally(), Tally(), Tally(), Tally()
    by_bucket = {name: Tally() for name, _, _ in _BUCKETS}
    by_mag = {name: Tally() for name, _, _ in _MAG_BANDS}
    engine, edge, agree, split = Tally(), Tally(), Tally(), Tally()
    for r in graded:
        pushed = r.result == PUSH
        base.add(r.result == OVER if not pushed else None, pushed)
        if r.lean:
            sign.add(r.hit, pushed)
            band = mag_band(r.sum_pts)
            if band is not None:
                by_mag[band].add(r.hit, pushed)
        by_bucket[bucket(r.sum_pts)].add(r.result == OVER if not pushed else None, pushed)
        if r.engine_lean:
            engine.add(r.engine_hit, pushed)
        if r.edge_lean:
            edge.add(r.edge_hit, pushed)
        if r.agreement == "agree":
            agree.add(r.hit, pushed)
        elif r.agreement == "split":
            split.add(r.hit, pushed)
    for day in sorted({r.date for r in graded}):
        todays = sorted((r for r in graded if r.date == day), key=lambda r: -r.sum_pts)
        if len(todays) < RANK_MIN_GAMES:
            continue
        for r in todays[:RANK_N]:
            if r.sum_pts > 0:
                rank_o.add(r.result == OVER if r.result != PUSH else None, r.result == PUSH)
        for r in todays[-RANK_N:]:
            if r.sum_pts < 0:
                rank_u.add(r.result == UNDER if r.result != PUSH else None, r.result == PUSH)
    return Summary(
        sign, rank_o, rank_u, base, by_bucket, by_mag, len({r.date for r in graded}), len(graded), legacy,
        engine=engine, edge=edge, agree=agree, split_sheet=split,
    )


def engine_lines(total: Summary) -> list[str]:
    """The engine's record on the same games, and who was right when the two disagreed."""
    return [
        f"Engine on the same lines: side of line {total.engine.text()} | vs market {total.edge.text()}",
        f"  sheet & engine agree {total.agree.text()} | split: sheet {total.split_sheet.text()}, engine {total.split_engine.text()}",
    ]


def mag_lines(total: Summary) -> list[str]:
    """One line per |SUM| band: record, win rate, 95% range, verdict vs break-even."""
    out = []
    for name, t in total.by_mag.items():
        lo, hi = wilson(t.hits, t.n)
        rng = f" [{100 * lo:.0f}-{100 * hi:.0f}%]" if t.n else ""
        out.append(f"  |SUM| {name:>6}: {t.text()}{rng} {verdict(t)}")
    return out


def summary_text(day: Date, yesterday: list[LedgerRow], total: Summary) -> str:
    """The lines the morning email prints; ``day`` is the sheet date being graded."""
    y = summarize(yesterday)
    lines = [f"Totals sheet {day.isoformat()}: {y.games} graded, slate went {y.over_rate.text()} over."]
    if y.legacy:
        lines.append(f"  {y.legacy} row(s) scored by older bands ({', '.join(sorted({r.bands for r in yesterday if not r.current}))}) are on record but not counted.")
    if y.games:
        lines.append(
            f"  sign {y.sign.text()} | top-{RANK_N} overs {y.rank_over.text()} | bottom-{RANK_N} unders {y.rank_under.text()}"
        )
        for r in sorted(yesterday, key=lambda r: -r.sum_pts):
            mark = "" if r.hit is None else "hit" if r.hit else "miss"
            runs = "" if r.runs is None else f" {r.away_runs}-{r.home_runs} ({r.runs})"
            eng = "" if r.engine_total is None else f" eng {r.engine_total:.1f}"
            lines.append(f"  {r.sum_pts:+d} {r.game} {r.line if r.line is not None else '?'}{eng}{runs} {r.result} {mark}".rstrip())
    lines.append(
        f"Ledger ({total.days} days, {total.games} games): sign {total.sign.text()} | "
        f"top-{RANK_N} overs {total.rank_over.text()} | bottom-{RANK_N} unders {total.rank_under.text()} | "
        f"slate over rate {total.over_rate.text()}"
        + (f" | {total.legacy} older-band rows not counted" if total.legacy else "")
    )
    lines.append(f"Win rate by |SUM| band (sign as the call; break-even {100 * BREAK_EVEN:.1f}%):")
    lines.extend(mag_lines(total))
    lines.extend(engine_lines(total))
    return "\n".join(lines)


# --- workbook -------------------------------------------------------------------------


def write_workbook(path: Path, sheet_day: Date, yesterday: list[LedgerRow], ledger: list[LedgerRow]) -> Path:
    bold = Font(bold=True)
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    cols = [
        "Date", "Game", "Line", "SUM", "Lean", "Away", "Home", "Runs", "Result", "Hit", "Bands",
        "Engine", "Eng O%", "Mkt O%", "Eng lean", "Eng hit", "Agree",
    ]

    def mark(hit: bool | None) -> str:
        return "" if hit is None else "hit" if hit else "miss"

    def fill(ws, rows: list[LedgerRow]) -> None:
        ws.append(cols)
        for c in ws[1]:
            c.font = bold
            c.alignment = Alignment(horizontal="center")
        for r in sorted(rows, key=lambda r: (r.date, -r.sum_pts)):
            ws.append([
                r.date, r.game, r.line, r.sum_pts, r.lean, r.away_runs, r.home_runs, r.runs, r.result,
                mark(r.hit), r.bands,
                r.engine_total, r.engine_p_over, r.market_p_over, r.engine_lean, mark(r.engine_hit), r.agreement,
            ])
            for name, hit in (("Hit", r.hit), ("Eng hit", r.engine_hit)):
                cell = ws.cell(row=ws.max_row, column=cols.index(name) + 1)
                if hit is True:
                    cell.fill = green
                elif hit is False:
                    cell.fill = red
        for i, name in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = 12 if name != "Game" else 16
        ws.freeze_panes = "A2"

    wb = Workbook()
    ws = wb.active
    ws.title = f"Graded {sheet_day.isoformat()}"
    fill(ws, yesterday)

    total = summarize(ledger)
    s = wb.create_sheet("Summary")
    s.append(["Measure", "Record", "Rate"])
    for c in s[1]:
        c.font = bold

    def rec(label: str, t: Tally) -> None:
        s.append([label, t.text().split(" (")[0], None if t.rate is None else round(t.rate, 3)])

    s.append([f"Ledger: {total.days} days, {total.games} graded games on bands {BANDS}", None, None])
    if total.legacy:
        s.append([f"{total.legacy} graded rows scored by older bands are in the Ledger tab and not counted here", None, None])
    rec("Sign of SUM (+ Over / - Under)", total.sign)
    rec(f"Top-{RANK_N} SUM each day as Overs", total.rank_over)
    rec(f"Bottom-{RANK_N} SUM each day as Unders", total.rank_under)
    rec("Slate over rate (base rate, not the sheet)", total.over_rate)
    s.append([])
    s.append(["Over rate by SUM band", "overs-unders(-pushes)", "over rate"])
    for c in s[s.max_row]:
        c.font = bold
    for name, t in total.by_bucket.items():
        rec(name, t)
    s.append([])
    s.append(["Win rate by |SUM| band (sign as the call)", "hits-misses(-pushes)", "win rate", "95% low", "95% high", "verdict"])
    for c in s[s.max_row]:
        c.font = bold
    for name, t in total.by_mag.items():
        lo, hi = wilson(t.hits, t.n)
        s.append([
            name, t.text().split(" (")[0], None if t.rate is None else round(t.rate, 3),
            round(lo, 3) if t.n else None, round(hi, 3) if t.n else None, verdict(t),
        ])
    s.append([])
    s.append(["Engine on the same lines", "hits-misses(-pushes)", "win rate"])
    for c in s[s.max_row]:
        c.font = bold
    rec("Engine's side of the sheet's line", total.engine)
    rec("Side the engine likes more than the market (priced lines only)", total.edge)
    rec("Sheet's call when the engine agreed", total.agree)
    rec("Sheet's call when the engine leaned the other way", total.split_sheet)
    rec("Engine's call on those same split games", total.split_engine)
    s.append([])
    s.append([
        "Break-even at -110 is 52.4%. A band shows 'edge' only when its whole 95% range clears that, on 30+ decided games. "
        "Sign and rank rates are graded against the sheet's own line at write time; "
        "a top-3 row counts as an Over only when its SUM is positive, a bottom-3 row as an Under only when negative. "
        "Engine = the simulator's median total read off its game_total rows in the engine ledger (calibrated, run-environment "
        "corrected, before the market anchor); Eng O% its over probability at the sheet's line, Mkt O% the devigged over where "
        "a book priced that exact line."
    ])
    s.column_dimensions["A"].width = 44
    s.column_dimensions["B"].width = 20
    s.column_dimensions["C"].width = 10
    for col in "DEF":
        s.column_dimensions[col].width = 10

    fill(wb.create_sheet("Ledger"), ledger)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --- entry point -----------------------------------------------------------------------


def audit_path(cfg: Config, day: Date) -> Path:
    return cfg.output_dir / f"totals_audit_{day.isoformat()}.xlsx"


def run_audit(cfg: Config, day: Date, sheet_day: Date | None = None) -> tuple[Path | None, str]:
    """Grade the sheet from ``sheet_day`` (default: the day before ``day``) and refresh the ledger.

    Writes ``totals_audit_<day>.xlsx`` and its ``.txt`` summary beside the day's
    package; returns (workbook path or None when the ledger is empty, summary text).
    """
    sheet_day = sheet_day or (day - timedelta(days=1))
    path = ledger_path(cfg)
    ledger = read_ledger(path)

    # Yesterday's sheet is the record of what was delivered: an ungraded day is
    # re-read from it (a sheet rewritten after its rows were filed, or one a box
    # generated without the audit) before the finals go in.
    sheet = cfg.output_dir / f"totals_sheet_{sheet_day.isoformat()}.xlsx"
    if sheet.exists() and not any(r.date == sheet_day.isoformat() and r.graded for r in ledger):
        try:
            fresh = rows_from_sheet(sheet, sheet_day, game_pks(sheet_day))
            ledger = merge([r for r in ledger if r.date != sheet_day.isoformat()], fresh)
        except Exception as exc:
            log.warning("totals audit: could not read %s: %s", sheet.name, exc)

    if any(r.engine_p_over is None for r in ledger):
        try:
            n = attach_engine(ledger, load_ledger(engine_ledger_path(cfg)))
            if n:
                log.info("totals audit: engine read attached to %d rows", n)
        except Exception as exc:
            log.warning("totals audit: engine ledger unreadable: %s", exc)

    for d in sorted({r.date for r in ledger if not r.graded}):
        if Date.fromisoformat(d) >= day:
            continue
        try:
            n = grade([r for r in ledger if r.date == d], finals(Date.fromisoformat(d)))
        except Exception as exc:
            log.warning("totals audit: finals for %s unavailable: %s", d, exc)
            continue
        log.info("totals audit: graded %d games on %s", n, d)

    write_ledger(path, ledger)
    yesterday = [r for r in ledger if r.date == sheet_day.isoformat()]
    text = summary_text(sheet_day, yesterday, summarize(ledger))
    if not ledger:
        return None, text
    out = write_workbook(audit_path(cfg, day), sheet_day, yesterday, ledger)
    out.with_suffix(".txt").write_text(text + "\n")
    return out, text


def record_sheet(cfg: Config, day: Date, sheet: Path, pks: dict[str, int]) -> None:
    """Called by the sheet writer: file today's rows so the audit can grade them tomorrow.

    A sheet rewritten before its games are graded replaces the day's ungraded
    rows, so the ledger carries the scores that were actually delivered; rows
    already graded are never touched.
    """
    path = ledger_path(cfg)
    ledger = read_ledger(path)
    if any(r.date == day.isoformat() and r.graded for r in ledger):
        return
    fresh = rows_from_sheet(sheet, day, pks)
    try:
        attach_engine(fresh, load_ledger(engine_ledger_path(cfg)))
    except Exception as exc:
        log.warning("totals sheet: engine ledger unreadable: %s", exc)
    ledger = [r for r in ledger if r.date != day.isoformat()]
    write_ledger(path, merge(ledger, fresh))
