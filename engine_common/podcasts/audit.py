"""The handicapper audit: every show and host, by league, market and kind.

Read from each league's graded podcast ledger in the store. Official bets and
leans are never pooled. A record is descriptive: "underpowered" means it is
within two standard errors of the 52.4% break-even at -110, which is most of
them for a long time.
"""

from __future__ import annotations

import io
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from engine_common.podcasts.picks import (
    LEDGER_FIELDS,
    MARKET_LABEL,
    Record,
    load_ledger,
    record,
    show_rows,
)

LEDGERS = "ledgers"


def all_ledgers(root: Path) -> list[dict[str, str]]:
    """Every league's graded podcast picks in the store."""
    rows: list[dict[str, str]] = []
    for path in sorted((root / LEDGERS).glob("*.csv")):
        for r in load_ledger(path):
            r.setdefault("league", path.stem)
            r["league"] = r["league"] or path.stem
            rows.append(r)
    return rows


@dataclass(frozen=True)
class AuditRow:
    league: str
    show: str
    host: str  # "" for the show as a whole
    kind: str
    market: str  # "all" or a MARKET_LABEL value
    rec: Record


def audit(rows: Iterable[dict[str, str]]) -> list[AuditRow]:
    """One row per league x show (and named host) x kind x market, plus each 'all'."""
    graded = [r for r in rows if r.get("result")]
    out: list[AuditRow] = []
    for league in sorted({r["league"] for r in graded}):
        lg = [r for r in graded if r["league"] == league]
        for show in sorted({r["show_name"] for r in lg}):
            mine = [r for r in lg if r["show_name"] == show]
            groups: list[tuple[str, list[dict[str, str]]]] = [("", show_rows(mine))]
            groups += [
                (h, [r for r in mine if r["host"] == h])
                for h in sorted({r["host"] for r in mine if r["host"]})
            ]
            for host, hr in groups:
                for kind in ("official", "lean"):
                    kr = [r for r in hr if r["kind"] == kind]
                    if not kr:
                        continue
                    label = f"{show} · {host}" if host else show
                    out.append(AuditRow(league, show, host, kind, "all", record(label, kr)))
                    for market, short in MARKET_LABEL.items():
                        mr = [r for r in kr if r["market"] == market]
                        if mr:
                            out.append(AuditRow(league, show, host, kind, short, record(label, mr)))
    return out


_HEAD = (
    "League",
    "Show",
    "Host",
    "Kind",
    "Market",
    "W",
    "L",
    "P",
    "Win%",
    "SE",
    "Units",
    "Priced",
    "ROI (priced)",
    "CLV pts",
    "CLV n",
    "Underpowered",
)


def _cells(a: AuditRow) -> list[object]:
    r = a.rec
    return [
        a.league.upper(),
        a.show,
        a.host,
        a.kind,
        a.market,
        r.wins,
        r.losses,
        r.pushes,
        None if r.win_pct is None else round(r.win_pct, 4),
        None if r.se is None else round(r.se, 4),
        round(r.units, 3),
        r.priced,
        None if r.roi is None else round(r.roi, 4),
        None if r.clv_mean is None else round(r.clv_mean, 2),
        r.clv_n,
        "yes" if r.underpowered else "",
    ]


def workbook(rows: list[dict[str, str]]) -> bytes:
    """The cross-league handicapper audit: records, then every graded pick."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Handicapper audit"
    ws.append(list(_HEAD))
    for a in audit(rows):
        ws.append(_cells(a))
    picks = wb.create_sheet("Picks")
    picks.append(list(LEDGER_FIELDS))
    for r in sorted(rows, key=lambda r: (r["league"], r["date"], r["show_name"])):
        picks.append([r.get(f, "") for f in LEDGER_FIELDS])
    note = wb.create_sheet("Notes")
    for line in (
        "Display and audit only: no podcast pick moves an engine price, probability or tier.",
        "Graded at the number the host said; a spread or total with no number is not graded.",
        "Units and ROI only on priced picks: the host's stated price, else the board's at the "
        "host's number (price_source).",
        "Show rows count one opinion once however many hosts shared it; host rows only where "
        "the episode named the host.",
        "Underpowered: win% within 2 SE of the 52.4% break-even at -110.",
    ):
        note.append([line])
    for sheet in (ws, picks):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def table(rows: list[dict[str, str]], *, league: str | None = None) -> str:
    """The audit as a plain-text table, official bets and leans, all markets."""
    lines = [
        f"{'Lg':<4}{'Show / host':<52}{'Kind':<9}{'W-L-P':<9}{'Win% ± SE':<13}"
        f"{'Units':>7}{'ROI(priced)':>15}{'CLV':>12}"
    ]
    for a in audit(rows):
        if a.market != "all" or (league and a.league != league):
            continue
        r = a.rec
        pct = "—" if r.win_pct is None or r.se is None else f"{r.win_pct:.0%} ± {r.se:.0%}"
        roi = "—" if r.roi is None else f"{r.roi:+.0%} on {r.priced}"
        clv = "—" if r.clv_mean is None else f"{r.clv_mean:+.1f} on {r.clv_n}"
        flag = "  underpowered" if r.underpowered else ""
        lines.append(
            f"{a.league.upper():<4}{r.label[:51]:<52}{a.kind:<9}{r.wlp:<9}{pct:<13}"
            f"{r.units:>+7.1f}{roi:>15}{clv:>12}{flag}"
        )
    return "\n".join(lines)


__all__ = ["AuditRow", "all_ledgers", "audit", "table", "workbook"]
